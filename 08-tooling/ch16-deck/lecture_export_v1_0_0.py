#!/usr/bin/env python3
"""Exports the finished deck's slide titles and speaker notes to lecture_out_v1_0_0.json beside this file, in the shape of
chapter 13's lecture_v1_0_0.json (the Lecture tab of the chapter page). Each slide is linked to the concept it teaches; every
concept id is checked against the chapter's topics, read from the v1_0_2 RDODI document and ABox. The file is left here and not
written into ch16-page, whose files belong to the page's own build.
Usage: lecture_export_v1_0_0.py <deck.pptx> <python-version>   (needs python-pptx and rdflib)"""
__version__ = "1.0.0"
import importlib.util, json, os, sys
from pptx import Presentation
HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("diagrams16", os.path.join(HERE, "..", "sen0414_ch16_diagrams_v1_0_0.py"))
DG = importlib.util.module_from_spec(spec); spec.loader.exec_module(DG)
IDS = {n["id"] for n in DG.nodes_from_ttl(os.path.join(HERE, "..", ".."))}
CONCEPT = {4: "SqliteVersusServers", 5: "SingleFileServerless", 6: "TableIsTupleList", 7: "ConnectCreatesFile", 8: "StrictMode",
           9: "SchemaTable", 10: "PythonValueMapping", 11: "UpdateDeleteWhere", 12: "SqlInjection", 13: "QmarkPlaceholders",
           14: "NullComparison", 15: "LikeAndIndex", 16: "BeginCommitRollback", 17: "AutocommitAttribute", 18: "WriterLocked",
           19: "RowidReuse", 20: "ForeignKeyEnforcement", 21: "ForeignKeyMismatch", 22: "DumpShippedFails", 23: "VaccinationChecker"}
bad = [v for v in CONCEPT.values() if v not in IDS]
if bad:
    raise SystemExit("REFUSED: unknown concept ids %r" % bad)
prs = Presentation(sys.argv[1]); out = []
for i, sl in enumerate(prs.slides, 1):
    texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    title = "Chapter 16: SQLite Databases" if i == 1 else next(t for t in texts if t != ">>>")
    out.append({"n": i, "title": title, "say": sl.notes_slide.notes_text_frame.text.strip(), "concept": CONCEPT.get(i, ""), "visual": ""})
json.dump({"_version": __version__, "_deck": "1.0.0", "_python": sys.argv[2], "slides": out},
          open(os.path.join(HERE, "lecture_out_v1_0_0.json"), "w"), indent=1)
print("lecture_out_v1_0_0.json:", len(out), "slides,", sum(1 for s in out if s["concept"]), "linked to a concept")
