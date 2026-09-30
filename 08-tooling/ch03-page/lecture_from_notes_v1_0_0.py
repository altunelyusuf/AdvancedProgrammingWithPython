"""Builds the chapter 3 page's Lecture tab from the speaker notes of the released deck (SEN0414_Ch03_Loops_3e_v2_0_1.pptx), which are a talk track.
The deck was made before decks wrote lecture_out.json (chapters 5 and 6 do), so the talk track is read back from the notes. Sentences that address the lecturer or
point at the slide are dropped by the same rule as lecture_export_v1_0_0.py in ch05-deck. The slide title is the text of the shape named 'title' (or the first long
text of the slide); the concept is read from the name of the slide's visual shapes ('v:<Concept>:...'). Usage: lecture_from_notes_v1_0_0.py <deck.pptx> <page_data.json> <out.json>"""
__version__ = "1.0.0"
import json, re, sys
from pptx import Presentation
DROP = re.compile(r"\b(click|clicks|clicked|slide|slides|Ask|Give them|Answer aloud|Teach|students|the room|Program written for this course|send them|on the left|on the right|the code card|the box|the panel|the grid|the table|Give them a minute|Two real runs)\b", re.I)
ids = {n["id"] for n in json.load(open(sys.argv[2]))["nodes"]}
p = Presentation(sys.argv[1]); out = []; dropped = 0
for i, s in enumerate(p.slides, 1):
    texts = [(sh.name, sh.text_frame.text.strip()) for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() and sh.text_frame.text.strip() != ">>>"]
    title = next((t for n, t in texts if n == "title"), None) or next((t for n, t in texts if not n.startswith(("v:", "say")) and len(t) > 12), None) or "Slide %d" % i
    concept = ""
    for sh in s.shapes:
        m = re.match(r"v:([A-Za-z0-9_]+):", sh.name)
        if m and m.group(1) in ids: concept = m.group(1); break
    note = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
    sents = re.split(r"(?<=[.?!])\s+(?=[A-Z‘“'\"])", note)
    keep = [x for x in sents if x and not DROP.search(x)]; dropped += len(sents) - len(keep)
    out.append({"n": i, "title": title.split("\n")[0], "say": " ".join(keep) or title.split("\n")[0], "concept": concept, "visual": concept})
json.dump({"_version": __version__, "_deck": sys.argv[1].split("/")[-1], "_python": "3.14.4", "slides": out}, open(sys.argv[3], "w"), indent=1, ensure_ascii=False)
print("written", sys.argv[3], len(out), "slides;", dropped, "sentences dropped;", sum(1 for x in out if x["concept"]), "tied to a concept")
