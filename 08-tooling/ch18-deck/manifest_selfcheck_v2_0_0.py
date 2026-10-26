#!/usr/bin/env python3
"""BUILDER'S OWN CHECK (not the independent deck check): the chapter-code manifest of chapter 18 against the deck's code boxes, and every manifest example re-run
(a) each in a FRESH interpreter (python -I, a new working folder) and
(b) all in ONE interpreter, in manifest order, the way the page runs them (the driver is the page's own _run function, as in pyodide_probe_v2_0_0.js).
The interpreter that runs this script runs the examples; Python 3.13 is the controller's interpreter for them.
Usage: manifest_selfcheck_v2_0_0.py <deck.pptx> <manifest.json>      Prints one line per finding and PASS or FAIL; exit 0 only on PASS."""
__version__ = "2.0.0"
import json, os, shutil, subprocess, sys, tempfile
from pptx import Presentation
deck, man = sys.argv[1], json.load(open(sys.argv[2], encoding="utf-8"))
prs = Presentation(deck)
boxes = {}
for i, s in enumerate(prs.slides, 1):
    code = []
    for sh in s.shapes:
        if sh.has_text_frame and any(p.text.startswith(">>> ") for p in sh.text_frame.paragraphs):
            code += [p.text[4:] for p in sh.text_frame.paragraphs]
    if code:
        boxes[i] = "\n".join(code)
bad = []
ex = man["examples"]
if [e["order"] for e in ex] != list(range(1, len(ex) + 1)):
    bad.append("order numbers are not 1..n")
if [e["slide"] for e in ex] != sorted(boxes):
    bad.append("manifest slides %s differ from code slides %s" % ([e["slide"] for e in ex], sorted(boxes)))
for e in ex:
    if boxes.get(e["slide"]) != e["code"]:
        bad.append("example %d (%s): code differs from the code boxes of slide %d" % (e["order"], e["name"], e["slide"]))
names = [e["name"] for e in ex]
if len(set(names)) != len(names):
    bad.append("names are not unique")
# (a) a fresh interpreter for each example
for e in ex:
    work = tempfile.mkdtemp(prefix="ch18_run_")
    r = subprocess.run([sys.executable, "-I", "-c", e["code"]], capture_output=True, text=True, timeout=300, cwd=work)
    if r.returncode != 0 or r.stdout != e["expected"]:
        bad.append("FRESH example %d (%s): re-run printed %r, manifest says %r (rc %d %s)" % (e["order"], e["name"], r.stdout[:60], e["expected"][:60], r.returncode, r.stderr[-120:]))
    shutil.rmtree(work, ignore_errors=True)
# (b) all in one interpreter, in order, with the page's _run (stdout and stderr captured, an exception printed as 'Name: message', last expression printed)
driver = '''import sys, io, ast, json
def _run(src):
    _o = io.StringIO(); ns = {'__name__': '__main__'}
    old = sys.stdout; olde = sys.stderr; sys.stdout = _o; sys.stderr = _o
    try:
        t = ast.parse(src)
        if t.body and isinstance(t.body[-1], ast.Expr):
            last = ast.Expression(t.body.pop().value)
            exec(compile(t, '<code>', 'exec'), ns)
            v = eval(compile(last, '<code>', 'eval'), ns)
            if v is not None: print(repr(v))
        else:
            exec(compile(t, '<code>', 'exec'), ns)
    except Exception as e:
        print(type(e).__name__ + ': ' + str(e))
    finally:
        sys.stdout = old; sys.stderr = olde
    return _o.getvalue().rstrip()
codes = json.load(open(sys.argv[1]))
print(json.dumps([_run(c) for c in codes]))
'''
mod = tempfile.mkdtemp(prefix="ch18_mod_")
cj = os.path.join(mod, "codes.json")
json.dump([e["code"] for e in ex], open(cj, "w"))
work = tempfile.mkdtemp(prefix="ch18_run_")
r = subprocess.run([sys.executable, "-I", "-c", driver, cj], capture_output=True, text=True, timeout=900, cwd=work)
if r.returncode != 0:
    bad.append("ONE-INTERPRETER run crashed: %s" % r.stderr[-200:])
else:
    got = json.loads(r.stdout.strip().splitlines()[-1])
    for e, g in zip(ex, got):
        if g.rstrip() != e["expected"].rstrip():
            bad.append("ONE-INTERPRETER example %d (%s): printed %r, manifest says %r" % (e["order"], e["name"], g[:80], e["expected"][:80]))
shutil.rmtree(work, ignore_errors=True); shutil.rmtree(mod, ignore_errors=True)
print("manifest %s | %d examples | %d code slides in the deck | python %s" % (sys.argv[2], len(ex), len(boxes), sys.version.split()[0]))
for b in bad:
    print("FINDING", b)
print("FAIL" if bad else "PASS: manifest code equals the deck's code boxes in order; every example re-run in a fresh interpreter, and all in one interpreter in order, printed its expected result")
sys.exit(1 if bad else 0)
