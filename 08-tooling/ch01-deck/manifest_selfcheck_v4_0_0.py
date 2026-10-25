#!/usr/bin/env python3
"""BUILDER'S OWN CHECK (not the independent deck check): the chapter-code manifest of chapter 1 against the deck's code boxes,
and every manifest example re-run in a fresh interpreter. Usage: manifest_selfcheck_v4_0_0.py <deck.pptx> <manifest.json>
Prints one line per finding and PASS or FAIL; exit 0 only on PASS."""
__version__ = "4.0.0"
import json, subprocess, sys
from pptx import Presentation
deck, man = sys.argv[1], json.load(open(sys.argv[2], encoding="utf-8"))
prs = Presentation(deck)
boxes = {}
for i, s in enumerate(prs.slides, 1):
    code = []
    for sh in s.shapes:
        if sh.has_text_frame and any(p.text.startswith(">>> ") for p in sh.text_frame.paragraphs):
            code += [p.text[4:] for p in sh.text_frame.paragraphs]
    if code:
        boxes[i] = "\n".join(code)
bad = []
ex = man["examples"]
if [e["order"] for e in ex] != list(range(1, len(ex) + 1)):
    bad.append("order numbers are not 1..n")
if [e["slide"] for e in ex] != sorted(boxes):
    bad.append("manifest slides %s differ from code slides %s" % ([e["slide"] for e in ex], sorted(boxes)))
for e in ex:
    if boxes.get(e["slide"]) != e["code"]:
        bad.append("example %d (%s): code differs from the code boxes of slide %d" % (e["order"], e["name"], e["slide"]))
    r = subprocess.run([sys.executable, "-I", "-c", e["code"]], capture_output=True, text=True, timeout=60)
    if r.returncode != 0 or r.stdout != e["expected"]:
        bad.append("example %d (%s): re-run printed %r, manifest says %r (rc %d %s)" % (e["order"], e["name"], r.stdout[:60], e["expected"][:60], r.returncode, r.stderr[-80:]))
names = [e["name"] for e in ex]
if len(set(names)) != len(names):
    bad.append("names are not unique")
print("manifest %s | %d examples | %d code slides in the deck" % (sys.argv[2], len(ex), len(boxes)))
for b in bad:
    print("FINDING", b)
print("FAIL" if bad else "PASS: manifest code equals the deck's code boxes in order; every example re-run in a fresh interpreter printed its expected result")
sys.exit(1 if bad else 0)
