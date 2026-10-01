#!/usr/bin/env python3
"""Chapter 14 library run: executes the chapter's openpyxl claims under Python 3.14.4.
openpyxl is not part of the course interpreter, so it was installed into a scratch folder (uv pip install --target) and this
script is run with SEN0414_CH14_LIBS pointing at it. Its printed lines are the evidence cited by the chapter 14 data module for
every claim about openpyxl. The example workbooks are written here by the script itself (the book's downloadable files were not used;
the values of Table 14-1 are typed in from the chapter page). No Excel, LibreOffice or network is involved.
Usage: SEN0414_CH14_LIBS=<dir> python sen0414_ch14_library_run_v1_0_0.py <workdir>"""
__version__ = "1.0.0"
import os, sys, tempfile, zipfile, datetime, warnings, pprint, importlib
sys.path.insert(0, os.environ["SEN0414_CH14_LIBS"])
W = sys.argv[1]; os.makedirs(W, exist_ok=True); os.chdir(W)
def out(label, *v): print("%-30s" % label, *v)
def err(label, fn):
    try: fn(); out(label, "no error")
    except Exception as e: out(label, type(e).__name__ + ": " + str(e))
import openpyxl, importlib.metadata as md
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.styles import Font
out("python", sys.version.split()[0]); out("openpyxl", md.version("openpyxl"), "et_xmlfile", md.version("et_xmlfile"))
out("lxml/defusedxml/pillow", [importlib.util.find_spec(m) is not None for m in ("lxml", "defusedxml", "PIL")])
# 1. example3.xlsx stand-in (Table 14-1)
D = datetime.datetime
rows = [(D(2035,4,5,13,34,2),"Apples",73),(D(2035,4,5,3,41,23),"Cherries",85),(D(2035,4,6,12,46,51),"Pears",14),(D(2035,4,8,8,59,43),"Oranges",52),
        (D(2035,4,10,2,7,0),"Apples",152),(D(2035,4,10,18,10,37),"Bananas",23),(D(2035,4,10,2,40,46),"Strawberries",98)]
wb = openpyxl.Workbook(); s1 = wb.active; s1.title = "Sheet1"
for r in rows: s1.append(list(r))
wb.create_sheet("Sheet2"); wb.create_sheet("Sheet3"); wb.save("example3.xlsx")
wb = openpyxl.load_workbook("example3.xlsx"); out("type(wb)", type(wb)); out("sheetnames", wb.sheetnames)
sheet = wb["Sheet3"]; out("sheet / type / title", repr(sheet), type(sheet), sheet.title); out("active", repr(wb.active))
sheet = wb["Sheet1"]; out("sheet['A1']", repr(sheet["A1"])); out("A1.value", repr(sheet["A1"].value)); c = sheet["B1"]
out("B1 value, row, column, coordinate", repr(c.value), c.row, c.column, c.coordinate); out("C1.value", repr(sheet["C1"].value))
out("cell(row=1,column=2)", repr(sheet.cell(row=1, column=2)), repr(sheet.cell(row=1, column=2).value))
for i in range(1, 8, 2): out("odd row", i, sheet.cell(row=i, column=2).value)
out("max_row / max_column", sheet.max_row, sheet.max_column, type(sheet.max_row).__name__)
out("column letters", get_column_letter(1), get_column_letter(2), get_column_letter(27), get_column_letter(900), get_column_letter(sheet.max_column), column_index_from_string("A"), column_index_from_string("AA"), column_index_from_string("M"))
err("get_column_letter(0)", lambda: get_column_letter(0))
err("column_index_from_string('')", lambda: column_index_from_string(""))
sl = sheet["A1":"C3"]; out("slice A1:C3 shape", type(sl).__name__, [type(r).__name__ for r in sl], [len(r) for r in sl])
out("slice repr", repr(sl))
for row in sl: out("row", [(x.coordinate, str(x.value)) for x in row])
out("list(rows)/list(columns)", len(list(sheet.rows)), len(list(sheet.columns)), [len(r) for r in list(sheet.rows)][:2], [len(r) for r in list(sheet.columns)])
out("list(columns)[1] repr", repr(list(sheet.columns)[1])); out("column B values", [x.value for x in list(sheet.columns)[1]])
out("sheet['B'] / sheet[2]", type(sheet["B"]).__name__, len(sheet["B"]), type(sheet[2]).__name__, len(sheet[2]))
out("iter_rows values_only", list(sheet.iter_rows(min_row=1, max_row=2, max_col=3, values_only=True)))
out("ws.values first row", next(iter(sheet.values)))
out("data_type A1/B1/C1", sheet["A1"].data_type, sheet["B1"].data_type, sheet["C1"].data_type); out("number_format A1", sheet["A1"].number_format)
# reading creates cells
t = openpyxl.Workbook().active; out("empty max_row/max_column", t.max_row, t.max_column); t["A1"]; t["C5"]; out("after reading A1,C5", t.max_row, t.max_column, len(t._cells))
_ = t["A100"].value; out("after reading A100 (value only)", t.max_row, t.max_column, t["A100"].value)
# 2. writing
wb = openpyxl.Workbook(); out("new wb", wb.sheetnames, wb.active.title); sh = wb.active; sh.title = "Spam Bacon Eggs Sheet"; out("renamed", wb.sheetnames)
wb = openpyxl.load_workbook("example3.xlsx"); wb["Sheet1"].title = "Spam Spam Spam"; wb.save("example3_copy.xlsx")
out("copy sheetnames", openpyxl.load_workbook("example3_copy.xlsx").sheetnames, "original", openpyxl.load_workbook("example3.xlsx").sheetnames)
wb = openpyxl.Workbook(); out("create_sheet()", repr(wb.create_sheet()), wb.sheetnames); out("create_sheet(index=0,title)", repr(wb.create_sheet(index=0, title="First Sheet")), wb.sheetnames)
wb.create_sheet(index=2, title="Middle Sheet"); out("after index=2", wb.sheetnames); del wb["Middle Sheet"]; del wb["Sheet1"]; out("after del", wb.sheetnames)
wb.remove(wb["Sheet"]); out("after wb.remove", wb.sheetnames)
w2 = openpyxl.Workbook(); w2.create_sheet("Mysheet", -1); out("create_sheet(-1)", w2.sheetnames); w2.create_sheet("Mysheet", 0); out("duplicate title", w2.sheetnames)
with warnings.catch_warnings(record=True) as w32:
    warnings.simplefilter("always"); long_ws = openpyxl.Workbook().create_sheet("x" * 32); out("title of 32 characters", len(long_ws.title), [str(w.category.__name__) + ": " + str(w.message) for w in w32])
err("title with a slash", lambda: openpyxl.Workbook().create_sheet("a/b"))
err("del of missing sheet", lambda: openpyxl.Workbook().__delitem__("Nope")); err("missing sheet name", lambda: openpyxl.Workbook()["Nope"])
wb = openpyxl.Workbook(); sheet = wb["Sheet"]; sheet["A1"] = "Hello, world!"; out("A1 value", sheet["A1"].value)
sheet["A2"] = 3.5; sheet["A3"] = True; sheet["A4"] = D(2035, 1, 2, 3, 4, 5); sheet["A5"] = datetime.date(2035, 1, 2); sheet["A6"] = None
out("types written", [(c.coordinate, c.data_type, c.number_format) for c in sheet["A"]])
err("a list as a value", lambda: sheet.__setitem__("B1", [1, 2])); tzw = openpyxl.Workbook(); tzw.active["A1"] = datetime.datetime(2035, 1, 1, tzinfo=datetime.timezone.utc); out("aware datetime assigned", "ok, stored as", tzw.active["A1"].data_type); err("aware datetime saved", lambda: tzw.save("tz.xlsx"))
# save overwrites
wb.save("over.xlsx"); a = os.path.getsize("over.xlsx"); wb2 = openpyxl.Workbook(); wb2.save("over.xlsx"); out("overwrite without warning", a, os.path.getsize("over.xlsx"))
# 3. fonts
wb = openpyxl.Workbook(); sheet = wb["Sheet"]; f = Font(size=24, italic=True); sheet["A1"].font = f; sheet["A1"] = "Hello, world!"; wb.save("styles3.xlsx")
sheet["A2"] = "x"; out("default font", sheet["A2"].font.name, sheet["A2"].font.sz, sheet["A2"].font.b, sheet["A2"].font.i)
l = openpyxl.load_workbook("styles3.xlsx")["Sheet"]; x = zipfile.ZipFile("styles3.xlsx").read("xl/styles.xml").decode(); out("fonts xml", x[x.index("<fonts"):x.index("</fonts>") + 8]); out("A1 font after reload", l["A1"].font.name, l["A1"].font.sz, l["A1"].font.i, l["A1"].font.b, l["A1"].value)
b = Font(name="Times New Roman", bold=True); sheet["A3"].font = b; out("Font(name,bold)", b.name, b.sz, b.b, b.i)
ff = sheet["A3"].font; err("modify a Font attribute", lambda: setattr(ff, "bold", False)); out("Font is shared/immutable", sheet["A3"].font is b, sheet["A3"].font == b)
# 4. formulas
wb = openpyxl.Workbook(); sheet = wb["Sheet"]; sheet["A1"] = 200; sheet["A2"] = 300; sheet["A3"] = "=SUM(A1:A2)"; wb.save("writeFormula3.xlsx")
out("A3 value/data_type", sheet["A3"].value, sheet["A3"].data_type)
r = openpyxl.load_workbook("writeFormula3.xlsx"); out("reload formula", r.active["A3"].value, r.active["A3"].data_type)
r = openpyxl.load_workbook("writeFormula3.xlsx", data_only=True); out("data_only, never calculated", repr(r.active["A3"].value))
z = zipfile.ZipFile("writeFormula3.xlsx"); x = z.read("xl/worksheets/sheet1.xml").decode(); i = x.index("<c r=\"A3\""); out("A3 xml", x[i:x.index("</c>", i) + 4])
out("members", z.namelist())
# text that looks like a formula, and a text number
sheet["B1"] = "=not a formula("; sheet["B2"] = "123"; out("types", sheet["B1"].data_type, sheet["B2"].data_type, repr(sheet["B2"].value)); sheet["B3"] = "'=SUM(A1:A2)"; out("quoted", sheet["B3"].data_type)
sheet["B4"] = "=SUM(A1:A2)"; sheet["B4"].data_type = "s"; out("forced string", sheet["B4"].data_type, sheet["B4"].value)
# 5. dimensions, merge, freeze, hide
wb = openpyxl.Workbook(); sheet = wb["Sheet"]; sheet["A1"] = "Tall row"; sheet["B2"] = "Wide column"; sheet.row_dimensions[1].height = 70; sheet.column_dimensions["B"].width = 20
out("row_dimensions/column_dimensions", type(sheet.row_dimensions).__name__, type(sheet.row_dimensions[1]).__name__, type(sheet.column_dimensions["B"]).__name__, sheet.row_dimensions[1].height, sheet.column_dimensions["B"].width, sheet.column_dimensions["A"].width)
sheet.column_dimensions["C"].hidden = True; sheet.row_dimensions[5].height = 100; out("hidden C, height 5", sheet.column_dimensions["C"].hidden, sheet.row_dimensions[5].height); wb.save("dimensions3.xlsx")
x = zipfile.ZipFile("dimensions3.xlsx").read("xl/worksheets/sheet1.xml").decode(); out("cols xml", x[x.index("<cols>"):x.index("</cols>") + 7]); out("row 1 xml", x[x.index('<row r="1"'):x.index('<row r="1"') + 60])
r = openpyxl.load_workbook("dimensions3.xlsx")["Sheet"]; out("reload dims", r.row_dimensions[1].height, r.column_dimensions["B"].width, r.column_dimensions["C"].hidden)
wb = openpyxl.Workbook(); sheet = wb["Sheet"]; sheet.merge_cells("A1:D3"); sheet["A1"] = "Twelve cells merged together."; sheet.merge_cells("C5:D5"); sheet["C5"] = "Two merged cells."
out("merged ranges", sorted(str(m) for m in sheet.merged_cells.ranges)); out("B2 type", type(sheet["B2"]).__name__); err("write into merged B2", lambda: setattr(sheet["B2"], "value", "x")); out("B2 read", repr(sheet["B2"].value))
wb.save("merged3.xlsx"); wb = openpyxl.load_workbook("merged3.xlsx"); sheet = wb["Sheet"]; sheet.unmerge_cells("A1:D3"); sheet.unmerge_cells("C5:D5"); out("after unmerge", list(sheet.merged_cells.ranges), type(sheet["B2"]).__name__)
err("unmerge not merged", lambda: sheet.unmerge_cells("F1:G2"))
for fp in ("A2", "B1", "C1", "C2", "A1", None): 
    s = openpyxl.Workbook().active; s.freeze_panes = fp; out("freeze_panes", fp, "->", s.freeze_panes)
s = openpyxl.Workbook().active; s.freeze_panes = s["B3"]; out("freeze with Cell", s.freeze_panes, s.sheet_view.pane.xSplit, s.sheet_view.pane.ySplit)
# 6. charts
wb = openpyxl.Workbook(); sheet = wb.active
for i in range(1, 11): sheet["A" + str(i)] = i * i
out("openpyxl.chart without import", hasattr(openpyxl, "chart"))
ref_obj = openpyxl.chart.Reference(sheet, 1, 1, 1, 10); out("Reference(sheet,1,1,1,10)", ref_obj.min_col, ref_obj.min_row, ref_obj.max_col, ref_obj.max_row, str(ref_obj))
ref2 = openpyxl.chart.Reference(sheet, 2, 3, 4, 5); out("Reference(sheet,2,3,4,5)", ref2.min_col, ref2.min_row, ref2.max_col, ref2.max_row, str(ref2))
out("Series is", openpyxl.chart.Series.__module__, openpyxl.chart.Series.__name__)
series_obj = openpyxl.chart.Series(ref_obj, title="First series"); chart_obj = openpyxl.chart.BarChart(); chart_obj.title = "My Chart"; chart_obj.append(series_obj)
sheet.add_chart(chart_obj, "C5"); wb.save("sampleChart3.xlsx"); z = zipfile.ZipFile("sampleChart3.xlsx"); out("chart members", [n for n in z.namelist() if "chart" in n or "drawing" in n])
c = z.read("xl/charts/chart1.xml").decode(); out("chart xml has barChart/title/ref", "<c:barChart>" in c or "barChart" in c, "My Chart" in c, "First series" in c, c[c.index("<c:f>"):c.index("</c:f>") + 6] if "<c:f>" in c else c[c.index("<f>"):c.index("</f>") + 4])
r = openpyxl.load_workbook("sampleChart3.xlsx"); out("charts after reload", len(r.active._charts))
out("chart kinds", [n for n in ("BarChart", "LineChart", "ScatterChart", "PieChart") if hasattr(openpyxl.chart, n)])
from openpyxl.chart import BarChart, Reference, Series
wb = openpyxl.Workbook(); sheet = wb.active
for i in range(10): sheet.append([i])
ch = BarChart(); ch.add_data(Reference(sheet, min_col=1, min_row=1, max_col=1, max_row=10)); sheet.add_chart(ch, "E15"); out("default chart size", ch.width, ch.height, repr(ch.anchor))
out("add_data series count / title", len(ch.series), ch.title)
# 7. Project 9 (census) with a small stand-in workbook
wb = openpyxl.Workbook(); sheet = wb.active; sheet.title = "Population by Census Tract"; sheet.append(["CensusTract", "State", "County", "POP2010"])
data = [("01001020100", "AL", "Autauga", 1912), ("01001020200", "AL", "Autauga", 2170), ("01001020300", "AL", "Autauga", 3373), ("01003010100", "AL", "Baldwin", 1948), ("02020000100", "AK", "Anchorage", 4020), ("02020000200", "AK", "Anchorage", "5000")]
for d in data: sheet.append(list(d))
wb.save("censuspopdata.xlsx")
code = '''import openpyxl, pprint
wb = openpyxl.load_workbook('censuspopdata.xlsx')
sheet = wb['Population by Census Tract']
county_data = {}
for row in range(2, sheet.max_row + 1):
    state  = sheet['B' + str(row)].value
    county = sheet['C' + str(row)].value
    pop    = sheet['D' + str(row)].value
    county_data.setdefault(state, {})
    county_data[state].setdefault(county, {'tracts': 0, 'pop': 0})
    county_data[state][county]['tracts'] += 1
    county_data[state][county]['pop'] += int(pop)
result_file = open('census2010.py', 'w')
result_file.write('allData = ' + pprint.pformat(county_data))
result_file.close()
'''
exec(compile(code, "readCensusExcel.py", "exec")); print(open("census2010.py").read()); sys.path.insert(0, W); import census2010
out("census2010.allData['AK']['Anchorage']", census2010.allData["AK"]["Anchorage"])
out("pprint sorts keys", pprint.pformat({"b": 1, "a": 2}), pprint.pformat({"b": 1, "a": 2}, sort_dicts=False))
# 8. Project 10 (produce sales)
wb = openpyxl.Workbook(); sheet = wb.active; sheet.title = "Sheet"; sheet.append(["PRODUCE", "COST PER POUND", "POUNDS SOLD", "TOTAL"])
P = [("Potatoes", 0.86, 21.6), ("Okra", 2.26, 38.6), ("Fava beans", 2.69, 32.8), ("Watercress", 2.99, 89.4), ("Celery", 1.9, 4.3), ("Garlic", 2.25, 10.0), ("Lemon", 1.1, 11.5)]
for n, (a, b, c) in enumerate(P, 2): sheet.append([a, b, c, "=ROUND(B%d*C%d, 2)" % (n, n)])
wb.save("produceSales3.xlsx")
PRICE_UPDATES = {"Garlic": 3.07, "Celery": 1.19, "Lemon": 1.27}
wb = openpyxl.load_workbook("produceSales3.xlsx"); sheet = wb["Sheet"]
for row_num in range(2, sheet.max_row + 1):
    produce_name = sheet.cell(row=row_num, column=1).value
    if produce_name in PRICE_UPDATES: sheet.cell(row=row_num, column=2).value = PRICE_UPDATES[produce_name]
wb.save("updatedProduceSales3.xlsx"); u = openpyxl.load_workbook("updatedProduceSales3.xlsx")["Sheet"]
out("updated rows", [(u.cell(row=r, column=1).value, u.cell(row=r, column=2).value, u.cell(row=r, column=4).value) for r in range(2, 9)])
out("total cells still formulas", [u.cell(row=r, column=4).data_type for r in range(2, 9)]); out("data_only after update", [openpyxl.load_workbook("updatedProduceSales3.xlsx", data_only=True)["Sheet"].cell(row=r, column=4).value for r in (2, 6)])
out("original untouched", openpyxl.load_workbook("produceSales3.xlsx")["Sheet"]["B6"].value)
out("Lemon vs Lemons", "Lemons" in PRICE_UPDATES, "Lemon" in PRICE_UPDATES)
# 9. insert_rows (practice program 2) and delete_rows
wb = openpyxl.Workbook(); sheet = wb.active
for i in range(1, 6): sheet.append(["r%d" % i, i])
sheet["C5"] = "=SUM(B1:B4)"; sheet.insert_rows(3, 2); out("insert_rows(3,2) column A", [c.value for c in sheet["A"]]); out("formula moved, text unchanged", sheet["C7"].value, sheet["C5"].value)
sheet.delete_rows(3, 2); out("delete_rows(3,2)", [c.value for c in sheet["A"]])
# 10. optimised modes, errors
wb = openpyxl.Workbook(write_only=True); out("write_only sheetnames", wb.sheetnames); ws = wb.create_sheet(); ws.append([1, 2, 3]); err("write-only ws['A1']", lambda: ws["A1"]); wb.save("wo.xlsx")
_e = None
try: wb.save("wo2.xlsx")
except Exception as e: _e = type(e).__name__
out("write-only saved twice", _e)
r = openpyxl.load_workbook("example3.xlsx", read_only=True); ws = r["Sheet1"]; out("read_only types", type(ws).__name__, type(next(ws.rows)[0]).__name__, ws.max_row, ws.calculate_dimension()); err("read_only iter_cols", lambda: list(ws.iter_cols())); err("read_only save", lambda: r.save("ro.xlsx")); r.close()
err("not an xlsx", lambda: (open("bad.xlsx", "w").write("a,b"), openpyxl.load_workbook("bad.xlsx")))
err("old .xls", lambda: (open("old.xls", "wb").write(b"\xd0\xcf\x11\xe0"), openpyxl.load_workbook("old.xls")))
err("missing file", lambda: openpyxl.load_workbook("nothing.xlsx"))
with warnings.catch_warnings(record=True) as ws_:
    warnings.simplefilter("always"); openpyxl.load_workbook("example3.xlsx").save("w.xlsx"); out("warnings on load/save", [str(w.category.__name__) + ":" + str(w.message) for w in ws_])
# 11. dates and serial numbers
from openpyxl.utils.datetime import to_excel, from_excel
out("to_excel", to_excel(D(2035, 4, 5, 13, 34, 2)), to_excel(D(1900, 3, 1)), to_excel(D(1900, 1, 1)), "from", from_excel(45000), from_excel(1.5))
z = zipfile.ZipFile("example3.xlsx"); x = z.read("xl/worksheets/sheet1.xml").decode(); i = x.index('<c r="A1"'); out("A1 xml", x[i:x.index("</c>", i) + 4])
wb = openpyxl.Workbook(); out("epoch / iso_dates", wb.epoch, wb.iso_dates)
# 12. copy_worksheet, dict-like iteration
wb = openpyxl.load_workbook("example3.xlsx"); cp = wb.copy_worksheet(wb["Sheet1"]); out("copy_worksheet", cp.title, wb.sheetnames, [s.title for s in wb])
# 13. xlsx is a zip
z = zipfile.ZipFile("example3.xlsx"); out("xlsx members", z.namelist()); x = z.read("xl/worksheets/sheet1.xml").decode(); i = x.index('<c r="B1"'); out("B1 xml", x[i:x.index("</c>", i) + 4])

# 14. the standard-library workbook of the page's playground program opens in openpyxl
import io, json, contextlib
pg = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "ch14-page", "playground_v1_0_0.json")))
ns = {}; buf = io.StringIO()
with contextlib.redirect_stdout(buf): exec(compile(pg["code"], "playground", "exec"), ns)
for nm, data in (("hand-made original", ns["book_bytes"]), ("hand-made updated", ns["updated_bytes"])):
    w = openpyxl.load_workbook(io.BytesIO(data)); sh = w["Sheet"]
    out(nm, w.sheetnames, sh.max_row, sh.max_column, [sh.cell(row=r, column=2).value for r in range(2, 6)], [sh.cell(row=r, column=4).value for r in range(2, 4)], sh["D2"].data_type)
    out(nm + " data_only", [openpyxl.load_workbook(io.BytesIO(data), data_only=True)["Sheet"].cell(row=r, column=4).value for r in range(2, 4)])
# 15. limits
out("last columns", get_column_letter(16384), column_index_from_string("XFD"), get_column_letter(18278), column_index_from_string("ZZZ"))
err("column_index_from_string('AAAA')", lambda: column_index_from_string("AAAA")); z0 = openpyxl.Workbook().active
err("ws['A0']", lambda: z0["A0"]); err("ws.cell(row=0, column=1)", lambda: z0.cell(row=0, column=1)); err("ws['A1048577']", lambda: z0["A1048577"])
out("cell(row,column,value) returns the cell", z0.cell(row=1, column=1, value=5).value)
# 16. practice questions and the two practice programs
out("Q10 column 14 / row 14", get_column_letter(14), hasattr(openpyxl.utils, "get_row_letter"))
from openpyxl.formula.translate import Translator
out("Translator", Translator("=SUM(B1:B4)", origin="C5").translate_formula("C7"), Translator("=ROUND(B2*C2, 2)", origin="D2").translate_formula("D3"))
from openpyxl.formula import Tokenizer
out("Tokenizer", [(t.value, t.type, t.subtype) for t in Tokenizer("=ROUND(B2*C2, 2)").items])
def multiplication_table(n, path):
    wb = openpyxl.Workbook(); ws = wb.active; bold = Font(bold=True)
    for i in range(1, n + 1):
        ws.cell(row=1, column=i + 1, value=i).font = bold; ws.cell(row=i + 1, column=1, value=i).font = bold
        for j in range(1, n + 1): ws.cell(row=i + 1, column=j + 1, value=i * j)
    wb.save(path)
multiplication_table(6, "multiplicationTable.xlsx"); t6 = openpyxl.load_workbook("multiplicationTable.xlsx").active
out("multiplication table 6", t6.max_row, t6.max_column, t6["G7"].value, t6["D5"].value, t6["A1"].value, t6["B1"].font.b, t6["A4"].font.b, t6["C3"].font.b)
def blank_row_inserter(n, m, src, dst):
    sheet = openpyxl.load_workbook(src).active; wb2 = openpyxl.Workbook(); out_sheet = wb2.active
    for row in sheet.iter_rows():
        for cell in row:
            r = cell.row if cell.row < n else cell.row + m
            out_sheet.cell(row=r, column=cell.column, value=cell.value)
    wb2.save(dst)
src = openpyxl.Workbook(); s = src.active
for i in range(1, 6): s.append(["item%d" % i, i * 10])
src.save("myProduce.xlsx"); blank_row_inserter(3, 2, "myProduce.xlsx", "myProduce_after.xlsx")
a = openpyxl.load_workbook("myProduce_after.xlsx").active; out("blank row inserter 3 2", [a.cell(row=r, column=1).value for r in range(1, 8)])
b = openpyxl.load_workbook("myProduce.xlsx").active; b.insert_rows(3, 2); out("same by insert_rows", [b.cell(row=r, column=1).value for r in range(1, 8)])
