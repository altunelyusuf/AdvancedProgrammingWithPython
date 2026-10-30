#!/usr/bin/env python3
"""sen0414_deck_teaching_check v1.0.0 (python-pptx) - the teaching elements T1 to T6 of a SEN0414 lecture deck, measured. A NEW tool; it edits nothing.
Written for the Owner's review remarks of 2026-10-30 on the chapter decks (indented code, no prompts, no lambda, no 5N1K label words, useful speaker notes).
Usage:   sen0414_deck_teaching_check_v1_0_0.py DECK.pptx            |      sen0414_deck_teaching_check_v1_0_0.py --selftest
Result words as printed: PASS, FAIL, NOT_MEASURED (no code box at all, so T1 and T6 cannot be measured). Exit 0 only when T1..T6 all PASS; 1 when any FAIL; 2 when only NOT_MEASURED.
DEFINITIONS (mechanical)
  code box      : a text shape whose runs are all in a monospace font (Courier, Consolas, Mono) and whose text is not the brand mark '>>>'. Its lines are its paragraphs.
  brand mark    : a shape whose whole text is '>>>' lying in the top 12 percent of the slide (the deck's logo; not code).
  T1 indentation: every code box has no tab, every indent is a multiple of four spaces, and on the syntax tree every compound statement (def, class, if, elif, else, for, while, with, try,
                  except, finally) has its body on its own lines, exactly four spaces deeper than its header. A flattened one-liner such as 'if x: y' fails. A box that does not parse is judged
                  by the rule 'a line ending with a colon is followed by a deeper line' (its parse failure is T6's finding).
  T2 no lambda  : the word lambda occurs in no text shape, table cell, code box or speaker note.
  T3 labels     : no line of any text shape, cell or note starts with WHO, WHAT, WHEN, WHERE, WHY or HOW in capitals; none of the six words (any case) is followed by a colon, en dash,
                  em dash or a spaced hyphen; and no single shape names four or more of the six words (a heading that spells out the six questions).
  T4 no prompts : no line of any text shape, cell, code box or note starts with '>>>' or '...' (the brand mark is exempt).
  T5 notes      : every slide's speaker notes hold 40 to 60 words; the first slide and every divider (a slide without the brand mark) 25 to 60.
  T6 compiles   : every code box passes ast.parse, unless the box contains the text '# does not run on purpose'.
"""
__version__ = "1.0.0"
import ast, os, re, sys

MONO = ("courier", "consolas", "mono")
HEAD = r"(?:WHO|WHAT|WHEN|WHERE|WHY|HOW)"
CAPS = re.compile(r"^\s*%s\b" % HEAD, re.M)
COLON = re.compile(r"\b%s\b\s*(?::|–|—|\s-\s)" % HEAD, re.I)
SIX = re.compile(r"\b(%s)\b" % HEAD, re.I)
PROMPT = re.compile(r"^\s*(>>>|\.\.\.)(\s|$)")
COMPOUND = (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef, ast.For, ast.AsyncFor, ast.While, ast.With, ast.AsyncWith, ast.If, ast.Try)


def is_mono(sh):
    if not sh.has_text_frame:
        return False
    names = [(r.font.name or "").lower() for p in sh.text_frame.paragraphs for r in p.runs]
    return bool(names) and all(any(m in n for m in MONO) for n in names)


def facts(path):
    from pptx import Presentation
    P = Presentation(path)
    H = P.slide_height
    out = []
    for i, s in enumerate(P.slides, 1):
        codes, texts, cells, brand = [], [], [], False
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text
                if t.strip() == ">>>" and sh.top is not None and sh.top < 0.12 * H:
                    brand = True
                    continue
                if is_mono(sh):
                    codes.append("\n".join(p.text.rstrip() for p in sh.text_frame.paragraphs).rstrip("\n"))
                else:
                    texts.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        if c.text.strip():
                            cells.append(c.text)
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        out.append(dict(n=i, code=codes, texts=texts, cells=cells, notes=notes, brand=brand))
    return out


def indent_findings(code):
    probs = []
    for i, l in enumerate(code.split("\n"), 1):
        if "\t" in l:
            probs.append("line %d has a tab" % i)
        k = len(l) - len(l.lstrip(" "))
        if l.strip() and k % 4:
            probs.append("line %d indented %d spaces, not a multiple of four" % (i, k))
    try:
        tree = ast.parse(code)
    except SyntaxError:
        lines = code.split("\n")
        for i, l in enumerate(lines[:-1]):
            if l.rstrip().endswith(":") and not l.lstrip().startswith("#"):
                nx = next((x for x in lines[i + 1:] if x.strip()), None)
                if nx is not None and len(nx) - len(nx.lstrip()) <= len(l) - len(l.lstrip()):
                    probs.append("line %d ends with a colon and the next line is not deeper" % (i + 1))
        return probs

    def body(node, stmts):
        if not stmts:
            return
        f = stmts[0]
        if f.lineno == node.lineno:
            probs.append("line %d: the body is on the header line" % node.lineno)
        elif f.col_offset != node.col_offset + 4:
            probs.append("line %d: body at column %d under a header at column %d (needs +4)" % (f.lineno, f.col_offset, node.col_offset))
    for node in ast.walk(tree):
        if isinstance(node, COMPOUND):
            body(node, node.body)
            if isinstance(node, ast.If):
                oe = node.orelse
                if oe and not (len(oe) == 1 and isinstance(oe[0], ast.If)):
                    h = oe[0]
                    if h.lineno == node.lineno:
                        probs.append("line %d: the else body is on the header line" % node.lineno)
                    elif h.col_offset != node.col_offset + 4:
                        probs.append("line %d: else body at column %d under an if at column %d (needs +4)" % (h.lineno, h.col_offset, node.col_offset))
            if isinstance(node, (ast.For, ast.AsyncFor, ast.While)) and node.orelse:
                body(node, node.orelse)
            if isinstance(node, ast.Try):
                for h in node.handlers:
                    body(h, h.body)
                if node.orelse:
                    body(node, node.orelse)
                if node.finalbody:
                    body(node, node.finalbody)
    return probs


def run(path):
    F = facts(path)
    N = len(F)
    R = {}
    ncode = sum(len(f["code"]) for f in F)

    def put(k, name, bad, ok):
        R[k] = (name, "PASS" if not bad else "FAIL", ok if not bad else "; ".join(bad[:6]) + (" (+%d more)" % (len(bad) - 6) if len(bad) > 6 else ""))
    if ncode == 0:
        R["T1"] = ("code is indented four spaces per level", "NOT_MEASURED", "no code box in the deck")
        R["T6"] = ("every code box compiles", "NOT_MEASURED", "no code box in the deck")
    else:
        bad = ["slide %d: %s" % (f["n"], p) for f in F for c in f["code"] for p in indent_findings(c)]
        put("T1", "code is indented four spaces per level", bad, "%d code boxes: no tab, multiples of four, every body four deeper on its own lines" % ncode)
        bad = []
        for f in F:
            for c in f["code"]:
                if "# does not run on purpose" in c:
                    continue
                try:
                    ast.parse(c)
                except SyntaxError as e:
                    bad.append("slide %d: %s (line %s)" % (f["n"], e.msg, e.lineno))
        put("T6", "every code box compiles", bad, "%d code boxes parse with ast.parse" % ncode)
    allt = lambda f: f["texts"] + f["cells"] + f["code"] + [f["notes"]]
    put("T2", "no lambda", ["slide %d" % f["n"] for f in F if any(re.search(r"\blambda\b", t, re.I) for t in allt(f))], "no lambda on %d slides" % N)
    bad = []
    for f in F:
        for t in f["texts"] + f["cells"] + [f["notes"]]:
            m = CAPS.search(t) or COLON.search(t)
            if m:
                bad.append("slide %d: label word %r" % (f["n"], m.group(0).strip()[:12]))
        for t in f["texts"]:
            if len({m.lower() for m in SIX.findall(t)}) >= 4:
                bad.append("slide %d: a shape names four of the six question words" % f["n"])
    put("T3", "no 5N1K label words", bad, "no label word on %d slides" % N)
    bad = []
    for f in F:
        for t in allt(f):
            for l in t.split("\n"):
                if PROMPT.match(l):
                    bad.append("slide %d: prompt line %r" % (f["n"], l[:24]))
    put("T4", "no prompts", bad, "no '>>>' or '...' line on %d slides" % N)
    bad = []
    for f in F:
        w = len(f["notes"].split())
        lo = 25 if (f["n"] == 1 or not f["brand"]) else 40
        if not lo <= w <= 60:
            bad.append("slide %d: %d words (needs %d to 60)" % (f["n"], w, lo))
    put("T5", "speaker notes 40 to 60 words (title and dividers 25 to 60)", bad, "all %d notes in range" % N)
    return R


def report(path):
    R = run(path)
    print("sen0414_deck_teaching_check v%s | deck %s" % (__version__, path))
    for k in ("T1", "T2", "T3", "T4", "T5", "T6"):
        name, st, meas = R[k]
        print("%s %s %s | %s" % (k, st, name, meas))
    fails = [k for k in R if R[k][1] == "FAIL"]
    nm = [k for k in R if R[k][1] == "NOT_MEASURED"]
    print("SUMMARY PASS=%d FAIL=%d NOT_MEASURED=%d | failing: %s" % (sum(1 for v in R.values() if v[1] == "PASS"), len(fails), len(nm), ",".join(sorted(fails)) or "none"))
    if fails:
        print("RESULT FAILING (exit 1)")
        return 1
    if nm:
        print("RESULT NOT_MEASURED (exit 2)")
        return 2
    print("RESULT PASS (exit 0)")
    return 0


def selftest():
    import json, subprocess
    here = os.path.dirname(os.path.abspath(__file__))
    fx = os.path.join(here, "fixtures", "deck_teaching_v1_0_0")
    exp = json.load(open(os.path.join(fx, "expected.json"), encoding="utf-8"))
    bad = 0
    for name, want in exp.items():
        R = run(os.path.join(fx, name))
        got = {k: R[k][1] for k in ("T1", "T2", "T3", "T4", "T5", "T6")}
        ok = got == want
        bad += not ok
        print("%s %s | expected %s | got %s" % ("OK  " if ok else "WRONG", name, ",".join("%s=%s" % kv for kv in want.items() if kv[1] != "PASS") or "all PASS", ",".join("%s=%s" % kv for kv in got.items() if kv[1] != "PASS") or "all PASS"))
    print("SELFTEST %s: %d fixtures, %d wrong" % ("PASS" if not bad else "FAIL", len(exp), bad))
    return 1 if bad else 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(selftest() if sys.argv[1] == "--selftest" else report(sys.argv[1]))
