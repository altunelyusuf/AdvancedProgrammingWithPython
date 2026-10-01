#!/usr/bin/env python3
"""Chapter 17 library run: executes the chapter's pypdf and python-docx claims under Python 3.14.4.
pypdf, python-docx and pdfminer.six are not part of the course interpreter, so they were installed into a scratch folder
(uv pip install --target) and this script is run with SEN0414_CH17_LIBS pointing at it. Its printed lines are the evidence
cited by the chapter 17 data module for every claim about those three packages. The sample PDFs and Word files are written here
by the script itself (the book's downloadable files were not used). Usage: SEN0414_CH17_LIBS=<dir> python sen0414_ch17_library_run_v1_0_0.py <workdir>"""
__version__ = "1.0.0"
import os, sys, io, zipfile, struct, zlib, tempfile, subprocess, warnings
sys.path.insert(0, os.environ["SEN0414_CH17_LIBS"])
W = sys.argv[1]; os.makedirs(W, exist_ok=True); os.chdir(W)
import pypdf, docx, pdfminer, importlib.metadata as md
from docx.enum.text import WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
def out(label, *v): print("%-34s" % label, *v)
def make_pdf(pages, image=False):
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [%s] /Count %d >>" % (" ".join("%d 0 R" % (4 + 2 * i) for i in range(len(pages))), len(pages)), "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, t in enumerate(pages):
        s = "BT /F1 18 Tf 72 720 Td (%s) Tj ET" % t
        objs.append("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % (5 + 2 * i))
        objs.append("<< /Length %d >>\nstream\n%s\nendstream" % (len(s), s))
    b = b"%PDF-1.4\n"; offs = []
    for i, o in enumerate(objs, 1):
        offs.append(len(b)); b += ("%d 0 obj\n%s\nendobj\n" % (i, o)).encode("latin-1")
    x = len(b); b += ("xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)).encode()
    for o in offs: b += ("%010d 00000 n \n" % o).encode()
    return b + ("trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, x)).encode()
def wr(name, data):
    with open(name, "wb") as f: f.write(data)
out("versions", "python", sys.version.split()[0], "pypdf", pypdf.__version__, "python-docx", docx.__version__, "pdfminer.six", md.version("pdfminer.six"))
wr("a.pdf", make_pdf(["Page one text", "Page two text", "Page three", "Page four", "Page five", "Page six"]))
r = pypdf.PdfReader("a.pdf"); out("pages", len(r.pages), [p.extract_text() for p in r.pages][:2])
out("text joined without separator", repr("".join(p.extract_text() for p in r.pages)[:40]))
w = pypdf.PdfWriter(); w.append("a.pdf", (0, 5)); out("append (0,5)", len(w.pages))
with open("first_five_pages.pdf", "wb") as f: out("write returns", w.write(f))
w = pypdf.PdfWriter(); w.append("a.pdf", (0, 5, 2)); out("append (0,5,2)", len(w.pages))
w = pypdf.PdfWriter(); w.append("a.pdf", [0, 5]); out("append [0,5]", len(w.pages), [p.extract_text() for p in w.pages])
w = pypdf.PdfWriter(); w.append("a.pdf", pages=[0, 1, 2, 3, 4]); out("append pages=[0..4]", len(w.pages))
import inspect; out("append parameters", list(inspect.signature(pypdf.PdfWriter.append).parameters)[:4])
warnings.simplefilter("error")
w.merge(2, "a.pdf", (0, 5)); out("merge(2, file, (0,5))", len(w.pages), "warnings raised: none")
w.insert_blank_page(index=2); out("insert_blank_page(index=2)", len(w.pages))
out("add_blank_page mediabox", list(w.add_blank_page().mediabox), len(w.pages))
try: pypdf.PdfWriter().add_blank_page()
except Exception as e: out("blank page on empty writer", type(e).__name__)
warnings.simplefilter("default")
w = pypdf.PdfWriter(); w.append("a.pdf")
for p in w.pages: p.rotate(90)
out("rotate(90) rotation", w.pages[0].rotation)
w.pages[1].rotate(-90); out("rotate(90) then rotate(-90)", w.pages[1].rotation)
w.pages[2].rotate(270); out("rotate(90) then rotate(270)", w.pages[2].rotation)
try: w.pages[0].rotate(45)
except Exception as e: out("rotate(45)", type(e).__name__, e)
out("page index 4 is page", w.pages[4].extract_text())
wr("wm.pdf", make_pdf(["DRAFT"])); wm = pypdf.PdfReader("wm.pdf").pages[0]
w = pypdf.PdfWriter(); w.append("a.pdf"); [p.merge_page(wm, over=False) for p in w.pages]; w.write("under.pdf")
out("watermark over=False", repr(pypdf.PdfReader("under.pdf").pages[0].extract_text()))
w = pypdf.PdfWriter(); w.append("a.pdf"); [p.merge_page(wm, over=True) for p in w.pages]; w.write("over.pdf")
out("stamp over=True", repr(pypdf.PdfReader("over.pdf").pages[0].extract_text()))
for algo in ("AES-256", None):
    w = pypdf.PdfWriter(); w.append("a.pdf"); (w.encrypt("swordfish", algorithm=algo) if algo else w.encrypt("swordfish")); w.write("enc_%s.pdf" % algo)
    t = pypdf.PdfReader("enc_%s.pdf" % algo).trailer["/Encrypt"]; out("encrypt algorithm=%s" % algo, "V =", t["/V"], "R =", t["/R"])
r = pypdf.PdfReader("enc_AES-256.pdf"); out("is_encrypted", r.is_encrypted)
try: r.pages[0]
except Exception as e: out("pages[0] before decrypt", type(e).__name__, e)
out("decrypt wrong / right", r.decrypt("an incorrect password").name, r.decrypt("swordfish").name)
w = pypdf.PdfWriter(); w.append(r); w.write("decrypted.pdf"); out("decrypted copy", pypdf.PdfReader("decrypted.pdf").is_encrypted, len(pypdf.PdfReader("decrypted.pdf").pages))
w = pypdf.PdfWriter(); w.append("a.pdf"); w.encrypt("user", "owner", algorithm="AES-256"); w.write("two.pdf"); r = pypdf.PdfReader("two.pdf")
out("user then owner password", r.decrypt("user").name, pypdf.PdfReader("two.pdf").decrypt("owner").name)
def img_pdf():
    raw = bytes([255, 0, 0, 0, 255, 0, 0, 0, 255, 255, 255, 0]); cs = "q 100 0 0 100 10 10 cm /Im0 Do Q"
    bodies = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>", b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Resources << /XObject << /Im0 5 0 R >> >> /Contents 4 0 R >>",
              ("<< /Length %d >>\nstream\n%s\nendstream" % (len(cs), cs)).encode(), b"<< /Type /XObject /Subtype /Image /Width 2 /Height 2 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Length 12 >>\nstream\n" + raw + b"\nendstream"]
    b = b"%PDF-1.4\n"; offs = []
    for i, o in enumerate(bodies, 1): offs.append(len(b)); b += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    x = len(b); b += b"xref\n0 %d\n0000000000 65535 f \n" % (len(bodies) + 1)
    for o in offs: b += b"%010d 00000 n \n" % o
    return b + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(bodies) + 1, x)
wr("img.pdf", img_pdf())
try: pypdf.PdfReader("img.pdf").pages[0].images[0]
except Exception as e: out("page.images without Pillow", type(e).__name__, e)
from pdfminer.high_level import extract_text
out("pdfminer extract_text", repr(extract_text("a.pdf")[:30]))
# project 12, run as the book writes it, in a folder of its own
os.makedirs("proj", exist_ok=True); os.chdir("proj")
for name, n in (("b_report.pdf", 3), ("A_cover.pdf", 2), ("c_one.pdf", 1)): wr(name, make_pdf(["%s p%d" % (name, i + 1) for i in range(n)]))
book = ("import pypdf, os\npdf_filenames = []\nfor filename in os.listdir('.'):\n    if filename.endswith('.pdf'):\n        pdf_files.append(filename)\npdf_filenames.sort(key=str.lower)\n"
        "writer = pypdf.PdfWriter()\nfor pdf_filename in pdf_filenames:\n    reader = pypdf.PdfReader(pdf_filename)\n    writer.append(pdf_filename, (1, len(reader.pages)))\nwith open('combined.pdf', 'wb') as file:\n    writer.write(file)\n")
env = {**os.environ, "PYTHONPATH": os.environ["SEN0414_CH17_LIBS"]}
def runprog(src, tag):
    wr(tag + ".py", src.encode()); r = subprocess.run([sys.executable, tag + ".py"], capture_output=True, text=True, env=env)
    return r.stderr.strip().splitlines()[-1] if r.returncode else "exit 0"
out("project 12 as printed", runprog(book, "bad"))
fixed = book.replace("pdf_files.append", "pdf_filenames.append")
for k in (1, 2):
    out("project 12 fixed, run %d" % k, runprog(fixed, "good"), [p.extract_text() for p in pypdf.PdfReader("combined.pdf").pages])
skip = fixed.replace("if filename.endswith('.pdf'):", "if filename.endswith('.pdf') and filename != 'combined.pdf':")
out("project 12 skipping own output", runprog(skip, "skip"), [p.extract_text() for p in pypdf.PdfReader("combined.pdf").pages])
os.chdir(W)
# python-docx
d = docx.Document(); p = d.add_paragraph("A plain paragraph with some ")
r1 = p.add_run("bold"); r1.bold = True; p.add_run(" and some "); r2 = p.add_run("italic"); r2.italic = True
out("runs", len(p.runs), [x.text for x in p.runs], "bold flags", [x.bold for x in p.runs])
out("get_text join", repr("\n".join(x.text for x in d.paragraphs)))
out("add_paragraph style Title", d.add_paragraph("Hello, world!", "Title").style.name)
out("add_heading levels", {l: d.add_heading("H", l).style.name for l in (0, 1, 4, 9)})
try: d.add_heading("x", 10)
except Exception as e: out("add_heading(10)", type(e).__name__, e)
try: d.add_paragraph("z", "NoSuchStyle")
except Exception as e: out("unknown style name", type(e).__name__, e)
q = d.add_paragraph("q", "Quote"); rr = q.add_run("x"); rr.style = "Quote Char"; out("run style", rr.style.name, rr.style.type)
try: rr.style = "Quote"
except Exception as e: out("paragraph style on a run", type(e).__name__, e)
names = ["Normal", "Body Text", "Caption", "Heading 1", "Intense Quote", "List", "List Bullet", "List Continue", "List Number ", "List Paragraph", "MacroText", "No Spacing", "Quote", "Subtitle", "TOC Heading", "Title"]
have = {s.name for s in d.styles}; out("book style names not in the default template", [n for n in names if n not in have])
out("styles with macro or List Number", sorted(n for n in have if "acro" in n or n.startswith("List Number")))
out("styles by type", {t.name: sum(1 for s in d.styles if s.type == t) for t in WD_STYLE_TYPE if t.name != "LIST"})
f = d.paragraphs[0].runs[0]
out("Run attributes present", {a: hasattr(f, a) for a in ("bold", "italic", "underline", "strike", "double_strike", "outline", "rtl")})
out("Font attributes present", {a: hasattr(f.font, a) for a in ("bold", "italic", "underline", "strike", "double_strike", "all_caps", "small_caps", "shadow", "outline", "rtl", "imprint", "emboss")})
f.strike = True; out("run.strike = True on a Run", "no error; w:strike in the XML:", "w:strike" in f._r.xml)
f.font.strike = True; out("run.font.strike = True", "w:strike in the XML:", "w:strike" in f._r.xml)
seen = []
f.bold = True; seen.append(f.bold); f.bold = False; seen.append(f.bold); f.bold = None; seen.append(f.bold); out("bold True, False, None", seen)
st = d.styles.add_style("Citation", WD_STYLE_TYPE.PARAGRAPH); out("add_style", st.name, st.type)
d.add_page_break(); t = d.add_table(rows=2, cols=2); t.cell(0, 1).text = "parrot"; out("table cell", t.cell(0, 1).text)
d = docx.Document(); d.add_paragraph("This is on the first page!"); d.paragraphs[0].runs[0].add_break(WD_BREAK.PAGE); d.add_paragraph("This is on the second page!")
out("page break xml", d.paragraphs[0].runs[0]._r.xml.count('w:type="page"'), [x.text for x in d.paragraphs])
guests = ["Prof. Plum", "Miss Scarlet", "Col. Mustard", "Al Sweigart", "RoboCop"]; d = docx.Document()
for g in guests:
    pp = d.add_paragraph("It would be a pleasure to have the company of"); d.add_paragraph(g); pp.runs[0].add_break(WD_BREAK.PAGE)
out("invitations: guests, page breaks", len(guests), d.element.xml.count('w:type="page"'))
wr("p.png", (lambda w, h, rgb: b"\x89PNG\r\n\x1a\n" + b"".join(struct.pack(">I", len(dd)) + t + dd + struct.pack(">I", zlib.crc32(t + dd)) for t, dd in ((b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)), (b"IDAT", zlib.compress(b"".join(b"\x00" + bytes(rgb) * w for _ in range(h)))), (b"IEND", b""))))(4, 4, (200, 30, 30)))
sh = d.add_picture("p.png", width=docx.shared.Inches(1), height=docx.shared.Cm(4)); out("picture size in EMU", type(sh).__name__, sh.width, sh.height)
out("Inches(1), Cm(4)", int(docx.shared.Inches(1)), int(docx.shared.Cm(4)))
buf = io.BytesIO(); d.save(buf); z = zipfile.ZipFile(buf); out("docx is a zip", z.namelist()[:6])
try: docx.Document("a.pdf")
except Exception as e: out("Document() on a PDF", type(e).__name__)
