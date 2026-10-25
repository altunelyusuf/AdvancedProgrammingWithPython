#!/usr/bin/env python3
"""Builds the chapter 3 lecture deck 5.0.0 (Loops) on the deck standard, and its chapter-code manifest.

A NEW tool (the decks 1.0.0 to 4.0.0 and their builders stay untouched). Every code example of examples_v5_0_0.py is RUN here; the
printed result and every number in a result visual come from that run. One light background scheme, a number on every
slide, a chip '>>>' on every slide that is not a divider, native tables and charts, drawn diagrams with arrows.
Structure and drawing helpers follow the chapter 9 generator (deck_build_v2_0_0.py, deck_kit_v2_0_0.py), the content is chapter 3's.
Writes: 03-materials/ch03/SEN0414_Ch03_Loops_3e_v5_0_0.pptx
        03-materials/ch03/page/chapter_code_ch03_v1_0_0.json   (order = slide order; the code exactly as the code boxes show it)
Usage:  python3 deck_build_v5_0_0.py
"""
__version__ = "5.0.0"
import contextlib
import io
import json
import math
import os
import platform
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
import importlib.util


def load(name, fn):
    sp = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(sp)
    sys.modules[name] = m
    sp.loader.exec_module(m)
    return m


from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches
import deck_kit_v5_0_0 as K
from deck_kit_v5_0_0 import C, MONO, arrow, chart, code_box, label_shape, new_deck, new_slide, notes, picture, rect, table, tb
X = load("examples_v5_0_0", "examples_v5_0_0.py")
V = load("deck_visuals_v5_0_0", "deck_visuals_v5_0_0.py")
T = load("deck_content_v5_0_0", "deck_content_v5_0_0.py")

OUT_DECK = os.path.join(ROOT, "03-materials", "ch03", "SEN0414_Ch03_Loops_3e_v5_0_0.pptx")
OUT_MAN = os.path.join(ROOT, "03-materials", "ch03", "page", "chapter_code_ch03_v1_0_0.json")
ASSETS = os.path.join(ROOT, "03-materials", "ch03", "assets")
PYV = platform.python_version()


# ------------------------------------------------------------------ run every example
def run_example(ex):
    code = X.code_text(ex)
    buf = io.StringIO()
    ns = {}
    with contextlib.redirect_stdout(buf):
        exec(compile(code, "<%s>" % ex["id"], "exec"), ns)
    return buf.getvalue(), ns


RUN = {ex["id"]: run_example(ex) for ex in X.EXAMPLES}

# ------------------------------------------------------------------ the plan: slide kinds in order
PARTS = {
    1: ("Part 1: Repeating with while", "A test is checked before every pass; break and continue change the flow", [
        ("while", "Repeats a block as long as its test is true."),
        ("Accumulator", "A variable that collects a total pass by pass."),
        ("break and continue", "Leave the loop, or jump back to its test."),
        ("Walrus :=", "Names a value inside the test of a loop."),
    ]),
    2: ("Part 2: Counting with for and range", "A for loop takes the values one at a time; range describes them", [
        ("for", "Takes the items of a sequence one at a time."),
        ("range()", "Describes a row of numbers without storing it."),
        ("Half-open interval", "Includes the start and leaves out the stop."),
        ("Loop variable", "Holds the current item and keeps its last value."),
    ]),
    3: ("Part 3: What a for loop visits", "Containers, pairs and the iterator that hands out the items", [
        ("List and tuple", "Ordered containers, one changeable and one fixed."),
        ("Dictionary", "Keys with values, visited as pairs with items()."),
        ("enumerate and zip", "A counter and pairs without index arithmetic."),
        ("Iterator", "The object that hands out one item at a time."),
    ]),
    4: ("Part 4: Modules and leaving early", "Borrow code by name, and end a program with a status", [
        ("import", "Opens another file of code under its own name."),
        ("Namespace", "Each module keeps its own names apart."),
        ("sys.exit()", "Ends a program by raising the exception SystemExit."),
        ("Exit status", "A number that tells the caller how it ended."),
    ]),
    5: ("Part 5: Two short programs", "The book's games, with scripted input so that they run on the page", [
        ("Guess the Number", "A loop with hints and a limit of six guesses."),
        ("Rock, Paper, Scissors", "A loop that keeps the score of every outcome."),
        ("Scripted input", "iter and next stand in for a typing person."),
        ("Seeded random", "A fixed seed makes every run repeat itself."),
    ]),
}

PLAN = [("title",), ("route",), ("questions",),
        ("divider", 1), ("analogy", "A1"), ("code", "while"), ("code", "accumulator"), ("story", "gauss"), ("code", "breakloop"), ("code", "continue"), ("code", "walrus"), ("story", "pep572"),
        ("divider", 2), ("analogy", "A2"), ("code", "forrange"), ("code", "rangeargs"), ("story", "dijkstra"), ("code", "lazyrange"),
        ("divider", 3), ("analogy", "A3"), ("code", "listtuple"), ("code", "dictloop"), ("code", "enumzip"), ("code", "zipstrict"), ("code", "iterator"), ("code", "loopelse"),
        ("divider", 4), ("analogy", "A4"), ("code", "importrandom"), ("code", "importforms"), ("code", "importerror"), ("analogy", "A5"), ("code", "sysexit"),
        ("divider", 5), ("analogy", "A6"), ("code", "guess1"), ("code", "guess2"), ("code", "rps1"), ("code", "rps2"),
        ("current",), ("summary",), ("practice",), ("resources",)]
N = len(PLAN)
NCONCEPTS = len({r[0].lower() for a_ in T.ANALOGIES for r in a_[3]})
SLIDE_OF = {}      # ("code", id) -> slide number, ("divider", k) -> slide number
for i, p in enumerate(PLAN, 1):
    SLIDE_OF[p] = i
CODE_ORDER = [p[1] for p in PLAN if p[0] == "code"]
assert CODE_ORDER == [e["id"] for e in X.EXAMPLES], "the examples file must list the examples in slide order"
for e in X.EXAMPLES:
    for g, _ in e["groups"]:
        assert len((">>> " + "\n>>> ".join(g)).split()) <= 30, "code box over 30 words in %s" % e["id"]


def part_slides(k):
    """slide numbers from divider k up to the slide before divider k+1 (or the changes slide)"""
    a = SLIDE_OF[("divider", k)]
    b = SLIDE_OF[("divider", k + 1)] - 1 if k < 5 else SLIDE_OF[("current",)] - 1
    return a, b


# ------------------------------------------------------------------ slide frames
def base(prs, title, sub=None, chip=True, nt=None):
    s = new_slide(prs)
    if chip:
        label_shape(s, 0.4, 0.28, 0.6, 0.36, ">>>", C["teal"], "FFFFFF", size=13, bold=True, font=MONO, radius=0.2, name="chip")
    tb(s, 1.15 if chip else 0.5, 0.18, 8.4 if chip else 9.0, 0.55, title, 22, C["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, name="title")
    if sub:
        tb(s, 1.15 if chip else 0.5, 0.72, 8.4, 0.3, sub, 12, C["muted"], name="subtitle")
    if nt:
        notes(s, nt)
    return s


def takeaway(s, text):
    rect(s, 0.4, 5.22, 0.07, 0.3, C["teal"], name="accent")
    tb(s, 0.55, 5.18, 7.9, 0.36, text, 12, C["muted"], italic=True, anchor=MSO_ANCHOR.MIDDLE, name="takeaway")


def add_numbers(prs):
    n = len(prs.slides)
    for i, s in enumerate(prs.slides, 1):
        tb(s, 8.6, 5.2, 1.0, 0.32, "%d / %d" % (i, n), 11, C["muted"], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, name="slide number")


# ------------------------------------------------------------------ slide makers
def make_title(prs):
    s = base(prs, "Loops", None, chip=True, nt=T.NOTES["title"])
    s.shapes[-1].left, s.shapes[-1].top, s.shapes[-1].width, s.shapes[-1].height = Inches(0.6), Inches(1.7), Inches(8.8), Inches(0.9)
    for r in s.shapes[-1].text_frame.paragraphs[0].runs:
        r.font.size = K.Pt(38)
    tb(s, 0.6, 2.65, 8.8, 0.7, "Chapter 3 of Automate the Boring Stuff with Python, 3rd edition, for an advanced course", 16, C["muted"], name="subtitle")
    for i, (a, b) in enumerate([(str(len(X.EXAMPLES)), "code examples that run on the chapter page"), (str(len(T.STORIES)), "true stories with who, what, when, where, why and how"), (str(NCONCEPTS), "distinct chapter concepts, each with an everyday picture")]):
        xx = 0.6 + i * 3.0
        rect(s, xx, 3.5, 2.8, 1.15, C["card"], line=C["line"], radius=0.08, name="card")
        tb(s, xx + 0.1, 3.55, 2.6, 0.5, a, 28, C["teal"], bold=True, name="stat")
        tb(s, xx + 0.1, 4.05, 2.6, 0.55, b, 11.5, C["ink"], name="stat label")
    tb(s, 0.6, 4.85, 8.0, 0.3, "SEN0414 Advanced Programming, Fall 2026. Yusuf Altunel, İstanbul Kültür Üniversitesi", 11, C["muted"], name="author")
    return s


def make_route(prs):
    s = base(prs, "Today's route", "Five parts; every range is a run of slides you can jump to", nt=T.NOTES["route"])
    rows = [["Part", "Slides", "Code examples", "You will run"]]
    runs = {1: "while, accumulator, break, continue, :=", 2: "for, range, half-open ranges", 3: "lists, dictionaries, enumerate, zip, iter",
            4: "import, namespaces, sys.exit", 5: "Guess the Number, Rock, Paper, Scissors"}
    for k in range(1, 6):
        a, b = part_slides(k)
        nex = sum(1 for p in PLAN[a - 1:b] if p[0] == "code")
        rows.append([PARTS[k][0].split(": ", 1)[1], "%d - %d" % (a, b), str(nex), runs[k]])
    rows.append(["Changes, summary, practice, resources", "%d - %d" % (SLIDE_OF[("current",)], N), "-", "what current Python adds"])
    table(s, 0.5, 1.2, 9.0, rows, [3.0, 1.1, 1.3, 3.8], fs=13, rowh=0.55, bold_cols=(0,))
    takeaway(s, "Open the chapter-code tab of the page: the same %d examples, in this order." % len(X.EXAMPLES))
    return s


def make_questions(prs):
    s = base(prs, "Two questions this chapter answers", "Carry them through every part; the answers come back at the end", nt=T.NOTES["questions"])
    q = [("1", "What decides when a loop repeats and when it stops, and which statements leave or skip a pass?",
          "Answer on slide %d: the test, break, continue, and else on a loop." % SLIDE_OF[("summary",)]),
         ("2", "Which parts of the book's chapter differ in current Python, and on what source?",
          "Answer on slide %d: the walrus, strict zip and the new warning." % SLIDE_OF[("current",)])]
    for i, (n, qq, aa) in enumerate(q):
        yy = 1.3 + i * 1.75
        rect(s, 0.5, yy, 9.0, 1.55, C["card"], line=C["line"], radius=0.05, name="card")
        label_shape(s, 0.7, yy + 0.2, 0.5, 0.5, n, C["teal"], "FFFFFF", size=18, bold=True)
        tb(s, 1.4, yy + 0.12, 7.9, 0.75, qq, 16, C["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, name="question")
        tb(s, 1.4, yy + 0.95, 7.9, 0.45, aa, 12.5, C["teal"], anchor=MSO_ANCHOR.MIDDLE, name="answer")
    takeaway(s, "A loop is a test, a block and a way out.")
    return s


def make_divider(prs, k):
    title, sub, concepts = PARTS[k]
    s = new_slide(prs)
    notes(s, T.NOTES["divider%d" % k])
    tb(s, 0.5, 0.2, 9.0, 0.7, title, 30, C["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, name="title")
    tb(s, 0.5, 0.92, 9.0, 0.34, sub, 14, C["muted"], name="subtitle")
    rect(s, 0.5, 1.33, 1.2, 0.05, C["teal"], name="rule")
    yy = 1.55
    for name, mean in concepts:
        label_shape(s, 0.5, yy, 2.2, 0.5, name, C["tealLight"], C["ink"], size=14, bold=True, radius=0.3)
        tb(s, 2.85, yy, 2.8, 0.6, mean, 12.5, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="meaning")
        yy += 0.72
    a, b = part_slides(k)
    items = []
    for i in range(a + 1, b + 1):
        p = PLAN[i - 1]
        if p[0] == "code":
            items.append((i, X.BY_ID[p[1]]["title"]))
        elif p[0] == "story":
            items.append((i, T.STORIES[p[1]]["title"]))
        elif p[0] == "analogy":
            items.append((i, [a_[1] for a_ in T.ANALOGIES if a_[0] == p[1]][0]))
    rows = [["Slide", "In this part"]] + [[str(i), t if len(t) < 41 else t[:38] + "..."] for i, t in items]
    rh = min(0.4, 3.5 / len(rows))
    table(s, 5.8, 1.55, 3.8, rows, [0.65, 3.8], fs=10.5, rowh=rh, name="table")
    return s


def make_analogy(prs, aid):
    _, title, head, rows = [a for a in T.ANALOGIES if a[0] == aid][0]
    s = base(prs, title, None, nt=T.NOTES[aid])
    rect(s, 0.4, 0.82, 9.2, 0.62, C["amberLight"], radius=0.08, name="card")
    tb(s, 0.55, 0.84, 8.9, 0.58, [[("ANALOGY  ", {"bold": True, "color": C["amber"]}), (head, {})]], 14, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="analogy")
    rh = min(0.5, 3.4 / (len(rows) + 1))
    table(s, 0.4, 1.6, 9.2, [["Chapter concept", "Everyday picture"]] + [[a, b] for a, b in rows], [3.2, 6.0], fs=11.5 if len(rows) > 7 else None, rowh=rh, bold_cols=(0,))
    return s


def make_story(prs, sid):
    st = T.STORIES[sid]
    s = base(prs, st["title"], None, nt=T.STORY_NOTES[sid])
    picture(s, os.path.join(ASSETS, st["img"]), 0.4, 0.95, 2.9, 3.2, "Picture for the story: " + st["title"].replace("Story: ", ""), "photo " + sid)
    tb(s, 0.4, 4.2, 2.9, 0.7, st["credit"], 9.5, C["muted"], name="credit")
    for fs in (13, 12.5, 12, 11.5, 11):
        cpl = (6.0 - 0.1) * 72.0 / (0.5 * fs)
        hs = []
        for lab in ("who", "what", "when", "where", "why", "how"):
            nlines = math.ceil((len(st[lab]) + 7) / (cpl * 0.93))
            hs.append(max(0.42, nlines * fs * 1.2 / 72.0 + 0.16))
        total = sum(hs) + 0.07 * 6 + 0.36
        if total <= 4.2:
            break
    yy = 0.92
    for lab, h in zip(("who", "what", "when", "where", "why", "how"), hs):
        rect(s, 3.5, yy, 6.1, h, C["card"], line=C["line"], radius=0.1, name="card")
        tb(s, 3.55, yy, 6.0, h, [[(lab.upper() + "  ", {"bold": True, "color": C["teal"]}), (st[lab], {})]], fs, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name=lab)
        yy += h + 0.07
    tb(s, 3.5, yy + 0.02, 6.1, 0.36, [[("READ  ", {"bold": True, "color": C["amber"]}), (st["linktext"], {"link": st["link"], "color": C["teal"]})]], fs, C["ink"], name="read")
    return s


def plan_code(ex, visual):
    """Choose layout and font: the largest code font (14 down to 9 pt) for which the code, its explanations and the result visual all fit.
    Order of preference: the explanation beside the code (side) at 11 pt or more, then below the code at 11 pt or more, then smaller."""
    maxlen = max(len(l) for l in X.code_lines(ex)) + 4
    forced_below = ex.get("layout") == "below"
    cands = []
    for lay in (["below"] if forced_below else ["side", "below"]):
        for fs in (14, 13, 12, 11, 10.5, 10, 9.5, 9):
            cands.append((lay, fs))
    cands.sort(key=lambda c: 0 if c[1] >= 11 else 1)
    top, gapv = 0.92, 0.06
    last = None
    for lay, fs in cands:
        need_w = (maxlen * 0.6 * fs + 14.4) / 72.0 + 0.08
        if lay == "side":
            box_w = max(3.2, need_w)
            if box_w > 6.3:
                continue
            expl_w = 9.2 - box_w + 0.02
        else:
            box_w, expl_w = 9.2, 9.2
            if need_w > 9.2:
                continue
        nfs = 12.5 if fs >= 12 else (12 if fs >= 10.5 else 11)
        lh = fs * 1.22 / 72.0
        cpl = max(10, (expl_w - 0.3) * 72.0 / (0.5 * nfs))
        rows = []
        for g, note in ex["groups"]:
            code_h = len(g) * lh + 0.2
            n_h = math.ceil(len(note) / cpl) * nfs * 1.25 / 72.0 + 0.14
            rows.append((code_h, n_h))
        heights = [max(a, b, 0.42) for a, b in rows] if lay == "side" else [a + b for a, b in rows]
        total = sum(heights) + gapv * (len(heights) - 1)
        ry = top + total + 0.08
        avail = 5.04 - (ry + 0.3)
        cand = (lay, fs, nfs, box_w, expl_w, rows, heights, ry, avail)
        if avail >= visual.need:
            return cand
        if last is None or (avail > last[-1] and fs >= 10):
            last = cand
    if last is None:
        raise SystemExit("REFUSED: %s has no layout in which its longest code line fits the slide" % ex["id"])
    print("WARN %s: no layout leaves %.2f in for the result (wants %.2f)" % (ex["id"], last[-1], visual.need))
    return last


def make_code(prs, exid):
    ex = X.BY_ID[exid]
    out, ns = RUN[exid]
    s = base(prs, ex["title"], None, nt=T.NOTES[exid])
    visual = V.VISUALS[ex["visual"]]
    lay, fs, nfs, box_w, expl_w, rows, heights, ry, avail = plan_code(ex, visual)
    y, gapv = 0.92, 0.06
    for (g, note), (ch_, nh_), hh in zip(ex["groups"], rows, heights):
        if lay == "side":
            code_box(s, 0.4, y, box_w, hh, g, fs)
            tb(s, 0.4 + box_w - 0.02, y, expl_w, hh, note, nfs, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="explanation", margin=(0.14, 0.02))
        else:
            code_box(s, 0.4, y, box_w, ch_, g, fs)
            tb(s, 0.4, y + ch_ + 0.01, box_w, nh_, note, nfs, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="explanation", margin=(0.12, 0.02))
        y += hh + gapv
    tb(s, 0.4, ry, 6.0, 0.26, "What the run printed (Python %s)" % PYV, 11, C["teal"], bold=True, name="result label")
    visual(s, 0.4, ry + 0.3, 9.2, max(avail, 1.0), ns, out)
    takeaway(s, T.TAKEAWAY[exid])
    return s


def make_current(prs):
    s = base(prs, "Current Python adds more", "Features newer than the book, from the chapter research", nt=T.NOTES["current"])
    rows = [["Feature", "Since", "What it gives you"],
            ["Assignment expression x := value", "3.8", "Names a value inside the test of a loop"],
            ["zip(a, b, strict=True)", "3.10", "Raises ValueError when the lengths differ"],
            ["return, break or continue leaving a finally block", "3.14", "The compiler emits a SyntaxWarning"],
            ["range() is a lazy sequence", "all", "Values are computed when they are asked for"],
            ["sys.exit() raises SystemExit", "all", "finally clauses still run on the way out"],
            ["Control-C raises KeyboardInterrupt", "all", "Stops a loop that never ends"]]
    table(s, 0.5, 1.25, 9.0, rows, [4.1, 0.8, 4.1], rowh=0.52, bold_cols=(0,), mono_cols=(), align_cols={1: PP_ALIGN.CENTER})
    takeaway(s, "Run what the book printed; then check what changed.")
    return s


def make_summary(prs):
    s = base(prs, "Summary", None, nt=T.NOTES["summary"])
    rect(s, 0.4, 0.82, 9.2, 0.5, C["amberLight"], radius=0.08, name="card")
    tb(s, 0.55, 0.83, 8.9, 0.48, [[("ANALOGY  ", {"bold": True, "color": C["amber"]}), ("A summary is a packing list: only what you really need goes into the bag.", {})]], 13, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="analogy")
    pts = ["A while loop tests first and repeats its block while the test is True.",
           "A for loop takes the items of any iterable; range stops before its stop value.",
           "break leaves, continue returns to the test, else runs when no break happened.",
           "enumerate and zip replace index arithmetic; iter and next sit under every for.",
           "import opens a namespace; sys.exit() raises SystemExit with an exit status."]
    for i, t_ in enumerate(pts):
        yy = 1.5 + i * 0.62
        label_shape(s, 0.4, yy + 0.04, 0.4, 0.4, str(i + 1), C["teal"], "FFFFFF", size=13, bold=True)
        tb(s, 0.9, yy, 4.6, 0.52, t_, 11.5, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="point")
    per = {k: sum(1 for e in X.EXAMPLES if e["part"] == k) for k in range(1, 6)}
    tb(s, 5.7, 1.4, 3.9, 0.5, "code examples in each part of this chapter", 10.5, C["teal"], bold=True, name="chart label")
    chart(s, "col", 5.6, 1.95, 4.0, 3.0, ["Part %d" % k for k in per], [("examples", list(per.values()))], fs=10, legend=False, gap=50, vmax=8, major=2)
    takeaway(s, "A loop is a test, a block and a way out.")
    return s


def make_practice(prs):
    s = base(prs, "Practice Questions", "Predict first, then run it on the chapter page", nt=T.NOTES["practice"])
    rect(s, 0.4, 1.05, 9.2, 0.5, C["amberLight"], radius=0.08, name="card")
    tb(s, 0.55, 1.06, 8.9, 0.48, [[("ANALOGY  ", {"bold": True, "color": C["amber"]}), ("Practice questions are finger scales for a pianist: short drills before the performance.", {})]], 13, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="analogy")
    qs = ["Change the while example so that it prints only the even values of spam. Which line do you add?",
          "What does list(range(10, 0, -3)) give? Predict it, then run it.",
          "Why does zip([1, 2, 3], 'ab') stop after two pairs, and how do you make it raise an error?",
          "Rewrite the Guess the Number loop to allow ten guesses and print how many are left."]
    for i, q in enumerate(qs):
        yy = 1.8 + i * 0.8
        rect(s, 0.4, yy, 9.2, 0.68, C["card"], line=C["line"], radius=0.08, name="card")
        label_shape(s, 0.55, yy + 0.14, 0.4, 0.4, str(i + 1), C["teal"], "FFFFFF", size=13, bold=True)
        tb(s, 1.1, yy + 0.04, 8.4, 0.6, q, 13, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="question")
    takeaway(s, "Each answer fits in a few lines you can run.")
    return s


def make_resources(prs):
    s = base(prs, "Resources: where to go next", "Reading and watching, chosen for this chapter; every link is live", nt=T.NOTES["resources"])
    for i, (t, url, host) in enumerate(T.RESOURCES):
        col, row = divmod(i, 7)
        xx = 0.4 + col * 4.7
        yy = 1.1 + row * 0.57
        rect(s, xx, yy, 4.5, 0.5, C["card"], line=C["line"], radius=0.1, name="card")
        tb(s, xx + 0.08, yy + 0.01, 4.35, 0.48, [[(t, {"bold": True})], [(url if len(url) < 58 else host, {"link": url, "color": C["teal"], "size": 9})]], 10.5, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="resource")
    takeaway(s, "Pick one link and read it before the next class.")
    return s


# ------------------------------------------------------------------ build
def build():
    prs = new_deck()
    for p in PLAN:
        kind = p[0]
        if kind == "title":
            make_title(prs)
        elif kind == "route":
            make_route(prs)
        elif kind == "questions":
            make_questions(prs)
        elif kind == "divider":
            make_divider(prs, p[1])
        elif kind == "analogy":
            make_analogy(prs, p[1])
        elif kind == "story":
            make_story(prs, p[1])
        elif kind == "code":
            make_code(prs, p[1])
        elif kind == "current":
            make_current(prs)
        elif kind == "summary":
            make_summary(prs)
        elif kind == "practice":
            make_practice(prs)
        elif kind == "resources":
            make_resources(prs)
    add_numbers(prs)
    prs.core_properties.title = "SEN0414 Chapter 3: Loops (lecture deck 5.0.0)"
    prs.core_properties.author = "Yusuf Altunel"
    os.makedirs(os.path.dirname(OUT_DECK), exist_ok=True)
    prs.save(OUT_DECK)
    return len(prs.slides)


def write_manifest():
    exs = []
    for k, ex in enumerate(X.EXAMPLES, 1):
        out, _ = RUN[ex["id"]]
        exs.append({"order": k, "id": ex["id"], "name": ex["name"], "slide": SLIDE_OF[("code", ex["id"])], "concepts": ex["concepts"],
                    "code": X.code_text(ex), "expected": out})
    names = [e["name"] for e in exs]
    assert len(set(names)) == len(names)
    man = {"chapter": 3, "version": "1.0.0", "kind": "chapter-code manifest",
           "deck": os.path.relpath(OUT_DECK, ROOT).replace(os.sep, "/"),
           "about": "The code examples of the chapter 3 deck, in the order of the code slides (order = slide order). The page's chapter-code tab lists them by NAME in this order; "
                    "every example is self-contained and runs with the Python standard library only and without a keyboard (typed answers are scripted with iter() and next()); "
                    "the expected result is what the run printed. The code is exactly the text of the deck's code boxes with the '>>> ' prompt removed.",
           "executed_with": {"python": PYV, "library": None},
           "library": {"file": None, "sha256": None, "data_files": [], "page_requirements": [
               "The Python standard library modules random, math, sys and io (imported by the examples that use them)",
               "The built-ins print, repr, len, range, list, sum, iter, next, enumerate, zip, bool and type; the assignment expression (Python 3.8 or later) and zip(strict=True) (Python 3.10 or later)",
               "No keyboard input: typed answers are scripted with iter() and next()"]},
           "prelude": "",
           "examples": exs}
    os.makedirs(os.path.dirname(OUT_MAN), exist_ok=True)
    with open(OUT_MAN, "w", encoding="utf-8") as f:
        json.dump(man, f, ensure_ascii=False, indent=1)
        f.write("\n")


if __name__ == "__main__":
    n = build()
    write_manifest()
    print("deck %s | %d slides | %d code examples | manifest %s" % (os.path.relpath(OUT_DECK, ROOT), n, len(X.EXAMPLES), os.path.relpath(OUT_MAN, ROOT)))
