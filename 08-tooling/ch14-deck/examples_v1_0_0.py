"""Chapter 14 deck examples: every console session, whole program, chart and table the lecture deck shows. Nothing here carries a
result; examples_run_v1_0_0.py runs it all under the interpreter given and records the results. A console group is one session:
its rows run in order in one namespace (as the deck check re-runs them per slide, one namespace per slide, so a group that shares a
slide with another begins with its own imports). No row touches the network or the file system. A row that needs a line break
inside a string writes it as an escaped \\n, so the stored row text never holds a newline character.

Programs come from sen0414_ch14_corpus_base_v1_0_0.py (the chapter's programs, each making stand-ins of the book's three files in a
temporary folder first), all twelve, and two written for the deck, fontsAndSizes and mergeAndFreeze, which are the font, size,
merge and freeze lines of the corpus program layoutFontMergeFreeze with their prints cut to lines that fit a slide (the corpus
program prints lines of 119 and 137 characters). A program that needs a package beyond openpyxl names it: liveExample3 needs requests
and reaches automatetheboringstuff.com (LIVE), safeLoadEntityBomb needs defusedxml. recalcWithLibreOffice runs the soffice of the
machine. Charts and tables are a setup and one expression, run again by the deck check."""
__version__ = "1.0.0"
import importlib.util
import os
import re

HERE = os.path.dirname(os.path.abspath(__file__))
_s = importlib.util.spec_from_file_location("base14", os.path.join(HERE, "..", "sen0414_ch14_corpus_base_v1_0_0.py"))
B = importlib.util.module_from_spec(_s)
_s.loader.exec_module(B)

CONSOLES = {
    "title": ["import openpyxl", "openpyxl.__version__"],
    "library": ["import openpyxl, io, zipfile, importlib.metadata as m",
                "openpyxl.__version__",
                "[r.split()[0] for r in m.requires('openpyxl')]",
                "b = io.BytesIO(); openpyxl.Workbook().save(b)",
                "len(zipfile.ZipFile(b).namelist())",
                "[n for n in zipfile.ZipFile(b).namelist() if n.startswith('xl/w')]"],
    "sheets": ["import openpyxl, io",
               "wb = openpyxl.Workbook()",
               "wb.create_sheet('Sheet2')",
               "b = io.BytesIO(); wb.save(b)",
               "wb = openpyxl.load_workbook(b)",
               "wb.sheetnames",
               "wb.active",
               "wb['Sheet2'].title",
               "wb['Nope']"],
    "cells": ["import openpyxl, datetime",
              "ws = openpyxl.Workbook().active",
              "data = [('Apples', 73), ('Cherries', 85), ('Pears', 14)]",
              "for row in data: ws.append(row)",
              "ws['A1'].value",
              "ws.cell(row=2, column=2).value",
              "ws['B3'].coordinate, ws.max_row, ws.max_column",
              "[[c.value for c in r] for r in ws['A1':'B2']]",
              "openpyxl.utils.get_column_letter(27)",
              "ws['D1'] = datetime.datetime(2035, 4, 5, 13, 34, 2)",
              "ws['D1'].value, ws['D1'].data_type"],
    "touch": ["import openpyxl",
              "ws = openpyxl.Workbook().active",
              "ws.max_row",
              "ws.cell(row=40, column=1).value is None",
              "ws.max_row, ws.dimensions",
              "abs(ws['A2'].value)"],
}

# ---- the programs ------------------------------------------------------------------------------------------------------------
FONTS = r"""import openpyxl, io, zipfile, re
from openpyxl.styles import Font
sheet = openpyxl.Workbook().active
sheet['A1'] = 'Hello, world!'
sheet['A1'].font = Font(size=24, italic=True)
sheet['B3'] = 'Bold Times New Roman'
sheet['B3'].font = Font(name='Times New Roman', bold=True)
for cell in (sheet['A1'], sheet['B3'], sheet['C5']):
    print(cell.coordinate, cell.font.name, cell.font.sz)
try: sheet['A1'].font.bold = False
except AttributeError as e: print('AttributeError', str(e)[:36])
sheet['A5'] = 'Tall row'; sheet['B5'] = 'Wide column'
sheet.row_dimensions[5].height = 70
sheet.column_dimensions['B'].width = 20
sheet.column_dimensions['D'].hidden = True
b = io.BytesIO(); sheet.parent.save(b)
xml = zipfile.ZipFile(b).read('xl/worksheets/sheet1.xml').decode()
print(*re.findall(r'<col [^>]*>', xml), sep='\n')
print(*re.findall(r'<row r="5"[^>]*>', xml))
"""

MERGE = r"""import openpyxl
sheet = openpyxl.Workbook().active
sheet.merge_cells('F1:I3')
sheet['F1'] = 'Twelve cells merged together.'
print(*sheet.merged_cells.ranges, type(sheet['G2']).__name__)
try: sheet['G2'] = 'x'
except AttributeError as e: print(e)
sheet.unmerge_cells('F1:I3')
print(len(sheet.merged_cells.ranges), type(sheet['G2']).__name__)
for setting in ('A2', 'B1', 'C1', 'C2'):
    ws = openpyxl.Workbook().active
    ws.freeze_panes = setting
    pane = ws.sheet_view.pane
    print(setting, pane.xSplit, pane.ySplit)
"""

# name -> (code, packages it needs beyond openpyxl, live?)
PROGRAMS = {name: (code, needs, live) for name, (code, needs, live) in B.PROGRAMS.items() if name != "recalcWithLibreOffice_v1_0_0.py"}
PROGRAMS["recalcWithLibreOffice_v1_0_0.py"] = B.PROGRAMS["recalcWithLibreOffice_v1_0_0.py"]
PROGRAMS["fontsAndSizes_v1_0_0.py"] = (FONTS, (), False)
PROGRAMS["mergeAndFreeze_v1_0_0.py"] = (MERGE, (), False)

# ---- a chart and a table, each a setup and one expression ------------------------------------------------------------------------
_CENSUS = re.search(r"for t in (\[.*?\]): ws\.append\(list\(t\)\)", B.PRE).group(1)
TABLES = {
    "countyTally": {
        "setup": "import openpyxl\nwb = openpyxl.Workbook(); sheet = wb.active\nsheet.append(['CensusTract', 'State', 'County', 'POP2010'])\n"
                 "for t in " + _CENSUS + ": sheet.append(list(t))\n"
                 "county_data = {}\nfor row in range(2, sheet.max_row + 1):\n"
                 "    state, county, pop = sheet['B' + str(row)].value, sheet['C' + str(row)].value, sheet['D' + str(row)].value\n"
                 "    county_data.setdefault(state, {})\n    county_data[state].setdefault(county, {'tracts': 0, 'pop': 0})\n"
                 "    county_data[state][county]['tracts'] += 1\n    county_data[state][county]['pop'] += int(pop)\n",
        "expr": "[['State', 'County', 'Tracts', 'Population']] + [[s, c, str(d['tracts']), '{:,}'.format(d['pop'])] "
                "for s in sorted(county_data) for c, d in sorted(county_data[s].items())]",
    }
}
CHARTS = {
    "squares": {
        "setup": "import openpyxl, io\nwb = openpyxl.Workbook(); sheet = wb.active\nfor i in range(1, 11): sheet['A' + str(i)] = i * i\n"
                 "ref_obj = openpyxl.chart.Reference(sheet, 1, 1, 1, 10)\nchart_obj = openpyxl.chart.BarChart()\n"
                 "chart_obj.append(openpyxl.chart.Series(ref_obj, title='First series'))\nsheet.add_chart(chart_obj, 'C5')\n"
                 "b = io.BytesIO(); wb.save(b); again = openpyxl.load_workbook(b).active\n"
                 "assert again._charts[0].series[0].val.numRef.f == \"'Sheet'!$A$1:$A$10\"\n",
        "expr": "[[str(r), again.cell(row=r, column=1).value] for r in range(1, 11)]",
    }
}
