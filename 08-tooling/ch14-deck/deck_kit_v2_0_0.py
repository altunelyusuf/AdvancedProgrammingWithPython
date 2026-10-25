"""Chapter 14 deck kit, version 2.0.0 (the drawing helpers of the chapter 13 deck kit v2.0.0, copied unchanged in behaviour): drawing helpers on python-pptx 1.0.2 (one light background scheme, native charts and
tables, drawn diagrams with connectors, code boxes whose every line starts with the '>>> ' prompt).
Used by deck_build_v2_0_0.py of chapter 14. Nothing here knows chapter content."""
__version__ = "2.0.0"
import re
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

W_IN, H_IN = 10.0, 5.625
FONT, MONO = "Calibri", "Courier New"
C = dict(bg="FBF8F2", ink="1F2933", muted="5B6773", teal="0E7C86", tealLight="D7EEF0", amber="B7791F", amberLight="F8E7C6",
         coral="B83A1B", green="2F855A", greenLight="DCF1E4", card="FFFFFF", line="D9D2C3", code="1E2A38", codeTxt="E6EDF3",
         pale="8FA3B8", kw="C9A0FF", str="9BE3A8", num="7CC4FF", cmt="8FA3B8", grey="ECE6D8", greyTxt="8A8F98", redLight="F7DAD2")


def rgb(h):
    return RGBColor.from_string(h)


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
    return prs


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(C["bg"])
    return s


def _style_run(r, size, color, bold=False, italic=False, font=None, link=None):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font or FONT
    r.font.color.rgb = rgb(color)
    if link:
        r.hyperlink.address = link


def fill_text(tf, paras, size, color, bold=False, italic=False, font=None, align=None, line_spacing=None):
    """paras: list of paragraphs; a paragraph is a str or a list of (text, opts) runs."""
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if align is not None:
            para.alignment = align
        if line_spacing:
            para.line_spacing = line_spacing
        runs = [(p, {})] if isinstance(p, str) else p
        for text, o in runs:
            r = para.add_run()
            r.text = text
            _style_run(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", italic), o.get("font", font), o.get("link"))


def tb(s, x, y, w, h, paras, size=12, color=None, bold=False, italic=False, font=None, align=None, anchor=MSO_ANCHOR.TOP,
       name=None, wrap=True, margin=(0.05, 0.03), line_spacing=None):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Inches(margin[0])
    tf.margin_top = tf.margin_bottom = Inches(margin[1])
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    fill_text(tf, paras, size, color or C["ink"], bold, italic, font, align, line_spacing)
    if name:
        sh.name = name
    return sh


def rect(s, x, y, w, h, fill, line=None, radius=None, name=None, shape=None, lw=0.75):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if name:
        sh.name = name
    return sh


def label_shape(s, x, y, w, h, text, fill, color, size=11, bold=False, line=None, font=None, radius=0.12, align=PP_ALIGN.CENTER, name=None):
    """An auto shape that carries its own text (a chip, a cell, a pill)."""
    sh = rect(s, x, y, w, h, fill, line=line, radius=radius, name=name)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_text(tf, [text], size, color, bold, font=font, align=align)
    return sh


def arrow(s, x1, y1, x2, y2, color=None, width=1.5, name="arrow"):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color or C["muted"])
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    te = etree.SubElement(ln, qn("a:tailEnd"))
    te.set("type", "triangle")
    te.set("w", "med")
    te.set("len", "med")
    c.name = name
    return c


def line_seg(s, x1, y1, x2, y2, color=None, width=1.0):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color or C["line"])
    c.line.width = Pt(width)
    return c


# ---- code
KW = {"for", "in", "def", "return", "lambda", "if", "else", "not", "and", "or", "import", "next", "print"}
TOK = re.compile(r"(f?'''(?:.|\n)*?'''|f?'(?:[^'\\]|\\.)*'|f?\"(?:[^\"\\]|\\.)*\"|#.*$|\b\d[\w.]*\b|\b[A-Za-z_]\w*\b|\s+|.)")


def code_runs(line):
    runs = []
    for t in TOK.findall(line):
        col = C["codeTxt"]
        if re.match(r"(f|r|b)?['\"]", t) and not t.isidentifier():
            col = C["str"]
        elif t.startswith("#"):
            col = C["cmt"]
        elif re.match(r"\d", t):
            col = C["num"]
        elif t in KW:
            col = C["kw"]
        runs.append((t, {"color": col, "font": MONO}))
    return runs


def code_box(s, x, y, w, h, lines, fs):
    """A dark card whose every paragraph is '>>> ' + one line of code; no wrapping; the line width is checked against the card."""
    longest = max(len(l) for l in lines) + 4
    need = longest * 0.6 * fs + 0.2 * 72
    if need > w * 72:
        raise SystemExit("REFUSED: a code line of %d characters at %.1f pt does not fit a %.2f in box" % (longest, fs, w))
    sh = rect(s, x, y, w, h, C["code"], radius=0.06, name="code")
    tf = sh.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    paras = [[(">>> ", {"color": C["pale"], "font": MONO})] + code_runs(l) for l in lines]
    fill_text(tf, paras, fs, C["codeTxt"], font=MONO, align=PP_ALIGN.LEFT)
    return sh


# ---- native table
def table(s, x, y, w, rows, colw, fs=None, rowh=0.27, header=True, mono_cols=(), cell_fill=None, cell_color=None, bold_cols=(), align_cols=None, name="table"):
    nr, nc = len(rows), len(rows[0])
    if fs is None:
        fs = max(10.0, min(13.0, round(rowh * 72 * 0.44 * 2) / 2.0))
    gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rowh * nr))
    gf.name = name
    t = gf.table
    tblPr = t._tbl.tblPr
    for a in ("bandRow", "firstRow"):
        tblPr.set(a, "0")
    # no style id: plain table, every cell filled by hand
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        tblPr.remove(sid)
    tot = float(sum(colw))
    for j, cw in enumerate(colw):
        t.columns[j].width = Inches(w * cw / tot)
    for i in range(nr):
        t.rows[i].height = Inches(rowh)
        for j in range(nc):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            hd = header and i == 0
            fill = C["teal"] if hd else (C["card"] if i % 2 else "F5F0E4")
            if cell_fill and (i, j) in cell_fill:
                fill = cell_fill[(i, j)]
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            col = "FFFFFF" if hd else (cell_color[(i, j)] if cell_color and (i, j) in cell_color else C["ink"])
            tf = cell.text_frame
            tf.word_wrap = True
            al = None
            if align_cols and j in align_cols:
                al = align_cols[j]
            fill_text(tf, [str(rows[i][j])], fs, col, bold=hd or j in bold_cols, font=MONO if (j in mono_cols and not hd) else None, align=al)
            _cell_borders(cell)
    return gf


def _cell_borders(cell, color=None):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", "6350")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr"))
        c.set("val", color or C["line"])
    # borders must come before the fill in tcPr
    fills = [e for e in tcPr if e.tag in (qn("a:solidFill"), qn("a:noFill"))]
    for f in fills:
        tcPr.remove(f)
        tcPr.append(f)


# ---- native charts
def chart(s, kind, x, y, w, h, cats, series, fs=10, colors=None, legend=None, number_format='General', gap=60, name="chart", axis_title=None, vmax=None, major=None):
    cd = CategoryChartData()
    cd.categories = cats
    for nm, vals in series:
        cd.add_series(nm, vals)
    ct = {"col": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED}[kind]
    gf = s.shapes.add_chart(ct, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = name
    ch = gf.chart
    ch.font.size = Pt(fs)
    ch.font.name = FONT
    ch.font.color.rgb = rgb(C["ink"])
    ch.has_title = False
    ch.has_legend = bool(legend if legend is not None else len(series) > 1)
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(fs)
    pl = ch.plots[0]
    pl.gap_width = gap
    pl.has_data_labels = True
    pl.data_labels.font.size = Pt(fs)
    pl.data_labels.number_format = number_format
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    cols = colors or [C["teal"], C["amber"], C["coral"]]
    for i, se in enumerate(pl.series):
        se.format.fill.solid()
        se.format.fill.fore_color.rgb = rgb(cols[i % len(cols)])
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(C["line"])
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(fs)
    va.minimum_scale = 0
    if vmax:
        va.maximum_scale = vmax
    if major:
        va.major_unit = major
    ch.category_axis.tick_labels.font.size = Pt(fs)
    ch.category_axis.format.line.color.rgb = rgb(C["line"])
    if axis_title:
        va.has_title = True
        va.axis_title.text_frame.text = axis_title
        va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(fs)
        va.axis_title.text_frame.paragraphs[0].runs[0].font.bold = False
    return gf


def picture(s, path, x, y, w, h, descr, name):
    """Fit the picture into the box without distortion (centred)."""
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    sc = min(w / iw, h / ih)
    pw, ph = iw * sc, ih * sc
    p = s.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
    p.name = name
    p._element.xpath("./p:nvPicPr/p:cNvPr")[0].set("descr", descr)
    return p


def notes(s, text):
    n = len(text.split())
    if n > 60:
        raise SystemExit("REFUSED: speaker note of %d words" % n)
    s.notes_slide.notes_text_frame.text = text
