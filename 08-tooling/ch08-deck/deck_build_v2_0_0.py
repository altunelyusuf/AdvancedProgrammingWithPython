#!/usr/bin/env python3
"""Builds the chapter 8 lecture deck 2.0.0 (Strings and Text Editing) on the deck standard, and its chapter-code manifest.

A NEW tool (the deck of 1.0.0 and its builder stay untouched). Every code example of examples_v2_0_0.py is RUN here; the
printed result and every number in a result visual come from that run. One light background scheme, a number on every
slide, a chip '>>>' on every slide that is not a divider, native tables and charts, drawn diagrams with arrows.
Writes: 03-materials/ch08/SEN0414_Ch08_Strings_3e_v2_0_0.pptx
        03-materials/ch08/page/chapter_code_ch08_v1_0_0.json   (order = slide order; the code exactly as the code boxes show it)
Usage:  python3 deck_build_v2_0_0.py
"""
__version__ = "2.0.0"
import contextlib
import io
import json
import math
import os
import platform
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
import importlib.util


def load(name, fn):
    sp = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(sp)
    sys.modules[name] = m
    sp.loader.exec_module(m)
    return m


from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches
import deck_kit_v2_0_0 as K
from deck_kit_v2_0_0 import C, MONO, arrow, chart, code_box, label_shape, new_deck, new_slide, notes, picture, rect, table, tb
X = load("examples_v2_0_0", "examples_v2_0_0.py")
V = load("deck_visuals_v2_0_0", "deck_visuals_v2_0_0.py")
T = load("deck_content_v2_0_0", "deck_content_v2_0_0.py")

OUT_DECK = os.path.join(ROOT, "03-materials", "ch08", "SEN0414_Ch08_Strings_3e_v2_0_0.pptx")
OUT_MAN = os.path.join(ROOT, "03-materials", "ch08", "page", "chapter_code_ch08_v1_0_0.json")
ASSETS = os.path.join(ROOT, "03-materials", "ch08", "assets")
PYV = platform.python_version()


# ------------------------------------------------------------------ run every example
def run_example(ex):
    code = X.code_text(ex)
    buf = io.StringIO()
    ns = {}
    with contextlib.redirect_stdout(buf):
        exec(compile(code, "<%s>" % ex["id"], "exec"), ns)
    return buf.getvalue(), ns


RUN = {ex["id"]: run_example(ex) for ex in X.EXAMPLES}

# ------------------------------------------------------------------ the plan: slide kinds in order
PARTS = {
    1: ("Part 1: Literals and reading a string", "How to write text in code, and how to read pieces of it", [
        ("String literal", "Text written between quotes in the code."),
        ("Escape sequence", "A backslash code for a character you cannot type."),
        ("Raw string", "A literal that keeps every backslash as typed."),
        ("Index and slice", "A number for one character, and a range of them."),
        ("The in operator", "A yes-or-no test for text inside text."),
    ]),
    2: ("Part 2: Formatting text", "Putting values into text: f-strings and the older styles", [
        ("F-string", "Text with holes that fill themselves from variables."),
        ("Format specifier", "A code after a colon for decimals, width and alignment."),
        ("Debug specifier", "An equals sign that prints the name and the value."),
        ("Older styles", "The percent operator and format() still fill slots."),
    ]),
    3: ("Part 3: String methods", "Methods that make a new string, test the string, or cut it", [
        ("Make a new string", "upper, lower, strip and rjust return a changed copy."),
        ("Test the string", "isalpha, isdecimal and startswith answer True or False."),
        ("Cut and join", "split, partition and join take text apart and back."),
        ("Find and replace", "find, count and replace search inside the text."),
    ]),
    4: ("Part 4: Characters and bytes", "Every character is a number, and bytes depend on the encoding", [
        ("Code point", "The number Unicode gives to one character."),
        ("ord and chr", "Two functions that convert between character and number."),
        ("Encoding", "The rule that turns code points into bytes."),
        ("UTF-8", "One to four bytes per character, ASCII unchanged."),
    ]),
    5: ("Part 5: Programs and the clipboard", "Small programs that use everything from parts 1 to 4", [
        ("Alternating case", "Position decides whether a letter is upper or lower."),
        ("Pig Latin", "Move the opening consonants to the end and add ay."),
        ("Table printer", "Pad every cell to the width of its column."),
        ("Clipboard program", "Paste the text, change it, copy it back."),
    ]),
}

PLAN = [("title",), ("route",), ("questions",),
        ("divider", 1), ("analogy", "A1"), ("code", "quotes"), ("code", "escapes"), ("code", "raw_multiline"), ("analogy", "A2"), ("code", "index"), ("code", "slice"), ("code", "in_immutable"),
        ("divider", 2), ("analogy", "A3"), ("story", "pep498"), ("code", "fstring"), ("code", "fspec"), ("code", "older"),
        ("divider", 3), ("analogy", "A4"), ("code", "case"), ("code", "chars"), ("code", "casetests"), ("analogy", "A5"), ("code", "affixes"), ("code", "joinsplit"),
        ("analogy", "A6"), ("code", "justify"), ("code", "strip"), ("code", "search"),
        ("divider", 4), ("analogy", "A7"), ("story", "unicode"), ("code", "codepoints"), ("story", "utf8"), ("code", "encoding"), ("code", "book_rjust"),
        ("divider", 5), ("analogy", "A8"), ("code", "alternate"), ("code", "piglatin"), ("code", "tableprinter"), ("story", "tesler"), ("code", "clipboard"),
        ("current",), ("summary",), ("practice",), ("resources",)]
N = len(PLAN)
SLIDE_OF = {}      # ("code", id) -> slide number, ("divider", k) -> slide number
for i, p in enumerate(PLAN, 1):
    SLIDE_OF[p] = i
CODE_ORDER = [p[1] for p in PLAN if p[0] == "code"]
assert CODE_ORDER == [e["id"] for e in X.EXAMPLES], "the examples file must list the examples in slide order"


def part_slides(k):
    """slide numbers from divider k up to the slide before divider k+1 (or the changes slide)"""
    a = SLIDE_OF[("divider", k)]
    b = SLIDE_OF[("divider", k + 1)] - 1 if k < 5 else SLIDE_OF[("current",)] - 1
    return a, b


# ------------------------------------------------------------------ slide frames
def base(prs, title, sub=None, chip=True, nt=None):
    s = new_slide(prs)
    if chip:
        label_shape(s, 0.4, 0.28, 0.6, 0.36, ">>>", C["teal"], "FFFFFF", size=13, bold=True, font=MONO, radius=0.2, name="chip")
    tb(s, 1.15 if chip else 0.5, 0.18, 8.4 if chip else 9.0, 0.55, title, 22, C["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, name="title")
    if sub:
        tb(s, 1.15 if chip else 0.5, 0.72, 8.4, 0.3, sub, 12, C["muted"], name="subtitle")
    if nt:
        notes(s, nt)
    return s


def takeaway(s, text):
    rect(s, 0.4, 5.22, 0.07, 0.3, C["teal"], name="accent")
    tb(s, 0.55, 5.18, 7.9, 0.36, text, 12, C["muted"], italic=True, anchor=MSO_ANCHOR.MIDDLE, name="takeaway")


def add_numbers(prs):
    n = len(prs.slides)
    for i, s in enumerate(prs.slides, 1):
        tb(s, 8.6, 5.2, 1.0, 0.32, "%d / %d" % (i, n), 11, C["muted"], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, name="slide number")


# ------------------------------------------------------------------ slide makers
def make_title(prs):
    s = base(prs, "Strings and Text Editing", None, chip=True, nt=T.NOTES["title"])
    s.shapes[-1].left, s.shapes[-1].top, s.shapes[-1].width, s.shapes[-1].height = Inches(0.6), Inches(1.7), Inches(8.8), Inches(0.9)
    for r in s.shapes[-1].text_frame.paragraphs[0].runs:
        r.font.size = K.Pt(38)
    tb(s, 0.6, 2.65, 8.8, 0.7, "Chapter 8 of Automate the Boring Stuff with Python, 3rd edition, for an advanced course", 16, C["muted"], name="subtitle")
    for i, (a, b) in enumerate([("24", "code examples that run on the chapter page"), ("4", "true stories with who, what, when, where, why and how"), ("57", "distinct chapter concepts, each with an everyday picture")]):
        xx = 0.6 + i * 3.0
        rect(s, xx, 3.5, 2.8, 1.15, C["card"], line=C["line"], radius=0.08, name="card")
        tb(s, xx + 0.1, 3.55, 2.6, 0.5, a, 28, C["teal"], bold=True, name="stat")
        tb(s, xx + 0.1, 4.05, 2.6, 0.55, b, 11.5, C["ink"], name="stat label")
    tb(s, 0.6, 4.85, 8.0, 0.3, "SEN0414 Advanced Programming, Fall 2026. Yusuf Altunel, İstanbul Kültür Üniversitesi", 11, C["muted"], name="author")
    return s


def make_route(prs):
    s = base(prs, "Today's route", "Five parts; every range is a run of slides you can jump to", nt=T.NOTES["route"])
    rows = [["Part", "Slides", "Code examples", "You will run"]]
    runs = {1: "quotes, escapes, slices, in", 2: "f-strings, format specs, older styles", 3: "case, tests, split, join, padding, strip, find",
            4: "ord, chr, UTF-8 bytes", 5: "alternating case, Pig Latin, table printer, clipboard"}
    for k in range(1, 6):
        a, b = part_slides(k)
        nex = sum(1 for p in PLAN[a - 1:b] if p[0] == "code")
        rows.append([PARTS[k][0].split(": ", 1)[1], "%d - %d" % (a, b), str(nex), runs[k]])
    rows.append(["Changes, summary, practice, resources", "%d - %d" % (SLIDE_OF[("current",)], N), "-", "what current Python adds"])
    table(s, 0.5, 1.2, 9.0, rows, [3.0, 1.1, 1.3, 3.8], fs=13.5, rowh=0.55, bold_cols=(0,))
    takeaway(s, "Open the chapter-code tab of the page: the same %d examples, in this order." % len(X.EXAMPLES))
    return s


def make_questions(prs):
    s = base(prs, "Two questions this chapter answers", "Carry them through every part; the answers come back at the end", nt=T.NOTES["questions"])
    q = [("1", "Which operations on a string make a new string, which only test it, and why can none of them change it?",
          "Answer on slide %d: none changes it; each makes a new string or answers." % SLIDE_OF[("summary",)]),
         ("2", "Which printed outputs and concepts of the book differ in current Python, and on what source?",
          "Answer on slides %d and %d: the rjust lines and the newer features." % (SLIDE_OF[("code", "book_rjust")], SLIDE_OF[("current",)]))]
    for i, (n, qq, aa) in enumerate(q):
        yy = 1.3 + i * 1.75
        rect(s, 0.5, yy, 9.0, 1.55, C["card"], line=C["line"], radius=0.05, name="card")
        label_shape(s, 0.7, yy + 0.2, 0.5, 0.5, n, C["teal"], "FFFFFF", size=18, bold=True)
        tb(s, 1.4, yy + 0.12, 7.9, 0.75, qq, 16, C["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, name="question")
        tb(s, 1.4, yy + 0.95, 7.9, 0.45, aa, 12.5, C["teal"], anchor=MSO_ANCHOR.MIDDLE, name="answer")
    takeaway(s, "Strings never change; operations make new strings or answers.")
    return s


def make_divider(prs, k):
    title, sub, concepts = PARTS[k]
    s = new_slide(prs)
    notes(s, T.NOTES["divider%d" % k])
    tb(s, 0.5, 0.2, 9.0, 0.7, title, 30, C["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, name="title")
    tb(s, 0.5, 0.92, 9.0, 0.34, sub, 14, C["muted"], name="subtitle")
    rect(s, 0.5, 1.33, 1.2, 0.05, C["teal"], name="rule")
    yy = 1.55
    for name, mean in concepts:
        label_shape(s, 0.5, yy, 2.2, 0.5, name, C["tealLight"], C["ink"], size=14, bold=True, radius=0.3)
        tb(s, 2.85, yy, 3.3, 0.6, mean, 13, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="meaning")
        yy += 0.72
    a, b = part_slides(k)
    items = []
    for i in range(a + 1, b + 1):
        p = PLAN[i - 1]
        if p[0] == "code":
            items.append((i, X.BY_ID[p[1]]["title"]))
        elif p[0] == "story":
            items.append((i, T.STORIES[p[1]]["title"].replace("Story: ", "Story: ")))
        elif p[0] == "analogy":
            items.append((i, [a_[1] for a_ in T.ANALOGIES if a_[0] == p[1]][0]))
    rows = [["Slide", "In this part"]] + [[str(i), t if len(t) < 46 else t[:43] + "..."] for i, t in items]
    rh = min(0.4, 3.55 / len(rows))
    table(s, 6.3, 1.55, 3.3, rows, [0.65, 3.4], rowh=rh, name="table")
    return s


def make_analogy(prs, aid):
    _, title, head, rows = [a for a in T.ANALOGIES if a[0] == aid][0]
    s = base(prs, title, None, nt=T.NOTES[aid])
    rect(s, 0.4, 0.82, 9.2, 0.62, C["amberLight"], radius=0.08, name="card")
    tb(s, 0.55, 0.84, 8.9, 0.58, [[("ANALOGY  ", {"bold": True, "color": C["amber"]}), (head, {})]], 14, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="analogy")
    rh = min(0.5, 3.55 / (len(rows) + 1))
    table(s, 0.4, 1.6, 9.2, [["Chapter concept", "Everyday picture"]] + [[a, b] for a, b in rows], [3.0, 6.2], rowh=rh, bold_cols=(0,))
    return s


def make_story(prs, sid):
    st = T.STORIES[sid]
    s = base(prs, st["title"], None, nt=T.STORY_NOTES[sid])
    picture(s, os.path.join(ASSETS, st["img"]), 0.4, 0.95, 2.9, 3.2, "Picture for the story: " + st["title"].replace("Story: ", ""), "photo " + sid)
    tb(s, 0.4, 4.2, 2.9, 0.7, st["credit"], 9.5, C["muted"], name="credit")
    yy = 0.92
    for lab in ("who", "what", "when", "where", "why", "how"):
        txt = st[lab]
        nlines = math.ceil((len(txt) + 7) / 64.0)
        h = max(0.44, nlines * 0.235 + 0.14)
        rect(s, 3.5, yy, 6.1, h, C["card"], line=C["line"], radius=0.1, name="card")
        tb(s, 3.55, yy, 6.0, h, [[(lab.upper() + "  ", {"bold": True, "color": C["teal"]}), (txt, {})]], 13, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name=lab)
        yy += h + 0.07
    tb(s, 3.5, yy + 0.02, 6.1, 0.36, [[("READ  ", {"bold": True, "color": C["amber"]}), (st["linktext"], {"link": st["link"], "color": C["teal"]})]], 13, C["ink"], name="read")
    return s


def plan_code(ex, visual):
    """Choose layout and font: the largest code font (14 down to 9 pt) for which the code, its explanations and the result visual all fit.
    Order of preference: the explanation beside the code (side) at 11 pt or more, then below the code at 11 pt or more, then smaller."""
    maxlen = max(len(l) for l in X.code_lines(ex)) + 4
    forced_below = ex.get("layout") == "below"
    cands = []
    for lay in (["below"] if forced_below else ["side", "below"]):
        for fs in (14, 13, 12, 11, 10.5, 10, 9.5, 9):
            cands.append((lay, fs))
    cands.sort(key=lambda c: 0 if c[1] >= 11 else 1)
    top, gapv = 0.92, 0.06
    last = None
    for lay, fs in cands:
        need_w = (maxlen * 0.6 * fs + 14.4) / 72.0 + 0.08
        if lay == "side":
            box_w = max(3.2, need_w)
            if box_w > 5.7:
                continue
            expl_w = 9.2 - box_w + 0.02
        else:
            box_w, expl_w = 9.2, 9.2
            if need_w > 9.2:
                continue
        nfs = 12.5 if fs >= 12 else (12 if fs >= 10.5 else 11)
        lh = fs * 1.22 / 72.0
        cpl = max(10, (expl_w - 0.3) * 72.0 / (0.5 * nfs))
        rows = []
        for g, note in ex["groups"]:
            code_h = len(g) * lh + 0.2
            n_h = math.ceil(len(note) / cpl) * nfs * 1.25 / 72.0 + 0.14
            rows.append((code_h, n_h))
        heights = [max(a, b, 0.42) for a, b in rows] if lay == "side" else [a + b for a, b in rows]
        total = sum(heights) + gapv * (len(heights) - 1)
        ry = top + total + 0.08
        avail = 5.04 - (ry + 0.3)
        last = (lay, fs, nfs, box_w, expl_w, rows, heights, ry, avail)
        if avail >= visual.need:
            return last
    print("WARN %s: no layout leaves %.2f in for the result (wants %.2f)" % (ex["id"], last[-1], visual.need))
    return last


def make_code(prs, exid):
    ex = X.BY_ID[exid]
    out, ns = RUN[exid]
    s = base(prs, ex["title"], None, nt=T.NOTES[exid])
    visual = V.VISUALS[ex["visual"]]
    lay, fs, nfs, box_w, expl_w, rows, heights, ry, avail = plan_code(ex, visual)
    y, gapv = 0.92, 0.06
    for (g, note), (ch_, nh_), hh in zip(ex["groups"], rows, heights):
        if lay == "side":
            code_box(s, 0.4, y, box_w, hh, g, fs)
            tb(s, 0.4 + box_w - 0.02, y, expl_w, hh, note, nfs, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="explanation", margin=(0.14, 0.02))
        else:
            code_box(s, 0.4, y, box_w, ch_, g, fs)
            tb(s, 0.4, y + ch_ + 0.01, box_w, nh_, note, nfs, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="explanation", margin=(0.12, 0.02))
        y += hh + gapv
    tb(s, 0.4, ry, 6.0, 0.26, "What the run printed (Python %s)" % PYV, 11, C["teal"], bold=True, name="result label")
    visual(s, 0.4, ry + 0.3, 9.2, max(avail, 1.0), ns, out)
    takeaway(s, T.TAKEAWAY[exid])
    return s


def make_current(prs):
    s = base(prs, "Current Python adds more", "Features newer than the book, from the chapter research", nt=T.NOTES["current"])
    rows = [["Feature", "Since", "What it gives you"],
            ["f'{x=}' debug specifier", "3.8", "Prints the name and the value together"],
            ["removeprefix, removesuffix", "3.9", "Cut exact text, not a set of characters"],
            ["f-string restrictions lifted (PEP 701)", "3.12", "A field may reuse the quote and hold a backslash"],
            ["Unknown escape such as \\d warns", "3.12", "The backslash is kept, and a SyntaxWarning appears"],
            ["Template strings, t'...' (PEP 750)", "3.14", "A Template object to process; a SyntaxError in 3.13"],
            ["UTF-8 mode by default (PEP 686)", "3.15", "open() will no longer depend on the locale"]]
    table(s, 0.5, 1.25, 9.0, rows, [3.6, 0.8, 4.6], fs=13.5, rowh=0.56, bold_cols=(0,), mono_cols=(), align_cols={1: PP_ALIGN.CENTER})
    takeaway(s, "Run what the book printed; then check what changed.")
    return s


def make_summary(prs):
    s = base(prs, "Summary", None, nt=T.NOTES["summary"])
    rect(s, 0.4, 0.82, 9.2, 0.5, C["amberLight"], radius=0.08, name="card")
    tb(s, 0.55, 0.83, 8.9, 0.48, [[("ANALOGY  ", {"bold": True, "color": C["amber"]}), ("A summary is a packing list: only what you really need goes into the bag.", {})]], 13, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="analogy")
    pts = ["Strings never change in place; every operation makes a new string or an answer.",
           "Quotes, escapes, raw and triple quotes decide which characters the literal holds.",
           "Indexes and slices read pieces; in and not in test for text inside text.",
           "f-strings put values into text; the specifier after the colon shapes them.",
           "Methods make, test or cut; characters are numbers; bytes depend on the encoding."]
    for i, t in enumerate(pts):
        yy = 1.5 + i * 0.62
        label_shape(s, 0.4, yy + 0.04, 0.4, 0.4, str(i + 1), C["teal"], "FFFFFF", size=13, bold=True)
        tb(s, 0.9, yy, 4.6, 0.52, t, 11.5, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="point")
    kinds = {"New string": ["upper", "lower", "casefold", "strip", "lstrip", "rstrip", "rjust", "ljust", "center", "replace", "join", "removeprefix", "removesuffix"],
             "True or False": ["isupper", "islower", "isalpha", "isalnum", "isdecimal", "isspace", "istitle", "startswith", "endswith"],
             "Number or pieces": ["find", "count", "split", "splitlines", "partition"]}
    for names in kinds.values():
        for nm in names:
            assert hasattr(str, nm), nm
    tb(s, 5.7, 1.4, 3.9, 0.5, "String methods in this chapter, by what they give back", 10.5, C["teal"], bold=True, name="chart label")
    chart(s, "col", 5.6, 1.95, 4.0, 3.0, list(kinds), [("methods", [len(v) for v in kinds.values()])], fs=10, legend=False, gap=50, vmax=16)
    takeaway(s, "No string operation changes a string in place.")
    return s


def make_practice(prs):
    s = base(prs, "Practice questions", "Predict first, then run it on the chapter page", nt=T.NOTES["practice"])
    rect(s, 0.4, 1.05, 9.2, 0.5, C["amberLight"], radius=0.08, name="card")
    tb(s, 0.55, 1.06, 8.9, 0.48, [[("ANALOGY  ", {"bold": True, "color": C["amber"]}), ("Practice questions are finger scales for a pianist: short drills before the performance.", {})]], 13, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="analogy")
    qs = ["What does 'a-b-c'.split('-')[::-1] give? Name the slice rule you used.",
          "Write one line that centres your name in 30 characters, with * as filler.",
          "Why does 'Hello'.upper() leave 'Hello' unchanged? What must you write to keep the result?",
          "Change the Pig Latin lambda so a word that starts with a vowel gets way, not yay."]
    for i, q in enumerate(qs):
        yy = 1.8 + i * 0.8
        rect(s, 0.4, yy, 9.2, 0.68, C["card"], line=C["line"], radius=0.08, name="card")
        label_shape(s, 0.55, yy + 0.14, 0.4, 0.4, str(i + 1), C["teal"], "FFFFFF", size=13, bold=True)
        tb(s, 1.1, yy + 0.04, 8.4, 0.6, q, 13, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="question")
    takeaway(s, "Each answer fits in one line you can run.")
    return s


def make_resources(prs):
    s = base(prs, "Resources: where to go next", "Reading and watching, chosen for this chapter; every link is live", nt=T.NOTES["resources"])
    for i, (t, url, host) in enumerate(T.RESOURCES):
        col, row = divmod(i, 6)
        xx = 0.4 + col * 4.7
        yy = 1.2 + row * 0.66
        rect(s, xx, yy, 4.5, 0.58, C["card"], line=C["line"], radius=0.1, name="card")
        tb(s, xx + 0.08, yy + 0.02, 4.35, 0.54, [[(t, {"bold": True})], [(url if len(url) < 58 else host, {"link": url, "color": C["teal"], "size": 9.5})]], 11, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="resource")
    takeaway(s, "Pick one link and read it before the next class.")
    return s


# ------------------------------------------------------------------ build
def build():
    prs = new_deck()
    for p in PLAN:
        kind = p[0]
        if kind == "title":
            make_title(prs)
        elif kind == "route":
            make_route(prs)
        elif kind == "questions":
            make_questions(prs)
        elif kind == "divider":
            make_divider(prs, p[1])
        elif kind == "analogy":
            make_analogy(prs, p[1])
        elif kind == "story":
            make_story(prs, p[1])
        elif kind == "code":
            make_code(prs, p[1])
        elif kind == "current":
            make_current(prs)
        elif kind == "summary":
            make_summary(prs)
        elif kind == "practice":
            make_practice(prs)
        elif kind == "resources":
            make_resources(prs)
    add_numbers(prs)
    prs.core_properties.title = "SEN0414 Chapter 8: Strings and Text Editing (lecture deck 2.0.0)"
    prs.core_properties.author = "Yusuf Altunel"
    os.makedirs(os.path.dirname(OUT_DECK), exist_ok=True)
    prs.save(OUT_DECK)
    return len(prs.slides)


def write_manifest():
    exs = []
    for k, ex in enumerate(X.EXAMPLES, 1):
        out, _ = RUN[ex["id"]]
        exs.append({"order": k, "id": ex["id"], "name": ex["name"], "slide": SLIDE_OF[("code", ex["id"])], "concepts": ex["concepts"],
                    "code": X.code_text(ex), "expected": out})
    names = [e["name"] for e in exs]
    assert len(set(names)) == len(names)
    man = {"chapter": 8, "version": "1.0.0", "kind": "chapter-code manifest",
           "deck": os.path.relpath(OUT_DECK, ROOT).replace(os.sep, "/"),
           "about": "The code examples of the chapter 8 deck, in the order of the code slides (order = slide order). The page's chapter-code tab lists them by NAME in this order; "
                    "every example is self-contained and runs with the Python standard library only; the expected result is what the run printed. The code is exactly the text of the "
                    "deck's code boxes with the '>>> ' prompt removed; every statement is one line.",
           "executed_with": {"python": PYV, "library": None},
           "library": {"file": None, "sha256": None, "data_files": [], "page_requirements": ["Python str, bytes, list, tuple, dict and the built-ins print, len, ord, chr, enumerate, zip, max, next (no import)"]},
           "prelude": "",
           "examples": exs}
    os.makedirs(os.path.dirname(OUT_MAN), exist_ok=True)
    with open(OUT_MAN, "w", encoding="utf-8") as f:
        json.dump(man, f, ensure_ascii=False, indent=1)
        f.write("\n")


if __name__ == "__main__":
    n = build()
    write_manifest()
    print("deck %s | %d slides | %d code examples | manifest %s" % (os.path.relpath(OUT_DECK, ROOT), n, len(X.EXAMPLES), os.path.relpath(OUT_MAN, ROOT)))
