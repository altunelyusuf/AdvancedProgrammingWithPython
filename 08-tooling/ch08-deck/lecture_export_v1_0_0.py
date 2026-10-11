#!/usr/bin/env python3
"""Exports the finished deck's slide titles and speaker notes to ../ch08-page/lecture_v1_0_0.json, the Lecture tab of the
chapter page (each slide's notes are reader's prose, so they are the text as they stand). Each slide is linked to the
concept it teaches; every concept id is checked against the chapter's page_data nodes (chapter 8 has no
sen0414_ch08_corpus_v*.py - the document ontology TTL is the content source, read here from the newest page_data json).
Usage: lecture_export_v1_0_0.py <deck.pptx> <python-version>"""
__version__ = "1.0.0"
import glob, json, os, sys
from pptx import Presentation
HERE = os.path.dirname(os.path.abspath(__file__))
_pf = sorted(glob.glob(os.path.join(HERE, "..", "ch08-page", "page_data_v*.json")),
             key=lambda x: [int(v) for v in x.rsplit("_v", 1)[1][:-5].split("_")])
IDS = {n["id"] for n in json.load(open(_pf[-1]))["nodes"]} if _pf else set()
CONCEPT = {3: "StringLiterals", 4: "StringQuoting", 5: "MultilineStrings", 6: "StringIndexing", 7: "StringImmutability",
           8: "FStrings", 9: "FStrings", 10: "FormatInterpolation", 11: "ChangingCase", 12: "CharacterTests",
           13: "JoinMethod", 14: "StripMethods", 15: "OrdAndChr", 16: "FindAndCount", 17: "Utf8Encoding",
           18: "Utf8Encoding", 19: "PigLatin", 20: "TablePrinter", 21: "ClipboardModule", 22: "ClipboardModule"}
bad = [v for v in CONCEPT.values() if v not in IDS]
if bad:
    raise SystemExit("REFUSED: unknown concept ids %r" % bad)
prs = Presentation(sys.argv[1]); out = []
for i, sl in enumerate(prs.slides, 1):
    texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    title = texts[0] if i > 1 else "Chapter 8: Manipulating Strings"
    if i > 1:
        title = next(t for t in texts if t not in (">>>",))
    out.append({"n": i, "title": title, "say": sl.notes_slide.notes_text_frame.text.strip(),
                "concept": CONCEPT.get(i, ""), "visual": ""})
json.dump({"_version": __version__, "_deck": "1.0.0", "_python": sys.argv[2], "slides": out},
          open(os.path.join(HERE, "..", "ch08-page", "lecture_v1_0_0.json"), "w"), indent=1)
print("lecture_v1_0_0.json:", len(out), "slides,", sum(1 for s in out if s["concept"]), "linked to a concept")
