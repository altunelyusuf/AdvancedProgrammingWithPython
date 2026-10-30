#!/usr/bin/env python3
"""BUILDER'S OWN CHECK (not the independent checks): the chapter-code manifest of chapter 8 against the deck's code boxes (one monospace box per code slide, no prompt),
and every manifest example re-run in a fresh interpreter. Usage: manifest_selfcheck_v2_1_0.py <deck.pptx> <manifest.json>. Exit 0 only on PASS."""
__version__ = "2.1.0"
import json, subprocess, sys
from pptx import Presentation
deck, man = sys.argv[1], json.load(open(sys.argv[2], encoding="utf-8"))
boxes = {}
for i, s in enumerate(Presentation(deck).slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() and sh.text_frame.text.strip() != ">>>":
            names = [(r.font.name or "").lower() for p in sh.text_frame.paragraphs for r in p.runs]
            if names and all("courier" in n for n in names):
                boxes.setdefault(i, []).append("\n".join(p.text.rstrip() for p in sh.text_frame.paragraphs).rstrip())
bad = []
ex = man["examples"]
if [e["order"] for e in ex] != list(range(1, len(ex) + 1)):
    bad.append("order numbers are not 1..n")
if [e["slide"] for e in ex] != sorted(boxes) or any(len(v) != 1 for v in boxes.values()):
    bad.append("manifest slides %s differ from code slides %s (one box each expected)" % ([e["slide"] for e in ex], sorted(boxes)))
for e in ex:
    if boxes.get(e["slide"], [None])[0] != e["code"].rstrip():
        bad.append("example %d (%s): code differs from the code box of slide %d" % (e["order"], e["name"], e["slide"]))
    r = subprocess.run([sys.executable, "-I", "-c", e["code"]], capture_output=True, text=True, timeout=60)
    if r.returncode != 0 or r.stdout != e["expected"]:
        bad.append("example %d (%s): re-run printed %r, manifest says %r (rc %d %s)" % (e["order"], e["name"], r.stdout[:60], e["expected"][:60], r.returncode, r.stderr[-80:]))
names = [e["name"] for e in ex]
if len(set(names)) != len(names):
    bad.append("names are not unique")
print("manifest %s | %d examples | %d code slides in the deck" % (sys.argv[2], len(ex), len(boxes)))
for b in bad:
    print("FINDING", b)
print("FAIL" if bad else "PASS: manifest code equals the deck's code boxes in order; every example re-run in a fresh interpreter printed its expected result")
sys.exit(1 if bad else 0)
