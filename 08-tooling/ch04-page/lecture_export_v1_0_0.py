#!/usr/bin/env python3
"""Exports the chapter 4 deck 3.0.0's slide titles and speaker notes (a short talk track, as they stand) to ch04-page/lecture_v1_0_0.json, the Lecture tab of
the chapter page (new file; the old page carried none). Each teaching slide is tied to the concept it teaches; every concept id is checked against the nodes of the
page data given. Usage: lecture_export_v1_0_0.py <deck.pptx> <page_data.json> <python-version>"""
__version__ = "1.0.0"
import json, os, sys
from pptx import Presentation
HERE = os.path.dirname(os.path.abspath(__file__))
CONCEPT = {6: "DefStatement", 7: "Deduplication", 10: "PositionalArgument", 11: "ReturnStatement", 12: "ImplicitNone", 14: "NamedParameter",
           15: "KeywordOnlyParameter", 16: "DefaultArgument", 19: "CallStack", 20: "FrameObject", 22: "LocalScope", 23: "LocalIsolation",
           24: "SameNameVariables", 25: "GlobalStatement", 26: "NonlocalStatement", 27: "ScopeIdentification", 28: "UnboundLocalError",
           31: "TryExcept", 32: "TryExcept", 33: "TryExcept", 34: "MultipleExceptTypes", 35: "Traceback", 36: "MultipleExceptTypes",
           39: "ZigzagProgram", 40: "SpikeProgram", 41: "CollatzSequence", 42: "CollatzSequence", 43: "ModernPractice"}
IDS = {n["id"] for n in json.load(open(sys.argv[2], encoding="utf-8"))["nodes"]}
bad = [v for v in CONCEPT.values() if v not in IDS]
if bad:
    raise SystemExit("REFUSED: unknown concept ids %r" % bad)
prs = Presentation(sys.argv[1]); out = []
for i, sl in enumerate(prs.slides, 1):
    texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip() and sh.text_frame.text.strip() != ">>>"]
    out.append({"n": i, "title": texts[0].split("\n")[0], "say": sl.notes_slide.notes_text_frame.text.strip(), "concept": CONCEPT.get(i, ""), "visual": ""})
json.dump({"_version": "1.0.0", "_deck": os.path.basename(sys.argv[1]), "_python": sys.argv[3], "slides": out},
          open(os.path.join(HERE, "lecture_v1_0_0.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
print("lecture_v1_0_0.json:", len(out), "slides,", sum(1 for s in out if s["concept"]), "tied to a concept")
