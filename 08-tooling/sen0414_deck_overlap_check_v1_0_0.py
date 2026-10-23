#!/usr/bin/env python3
"""sen0414_deck_overlap_check v1.0.0 (python-pptx) - geometric scan of a deck for text that collides with other content.
A NEW tool; it edits nothing. Port of the sibling course's deck_overlap_check_v1_0_0.py (SEN0401): only the file name and the tool name changed; the measure is the sibling's. For every slide it estimates the rendered extent of each text frame (greedy word wrap, average glyph advance
0.50 em proportional / 0.60 em monospace, line height 1.2 em, the frame's anchor and insets) and of each table (rows grow to fit their wrapped
cells), then reports
  OVERLAP  two content items (text frame, table, chart, picture) whose extents intersect by more than 0.03 in in both directions
  SPILL    a text frame whose estimated text is taller than its own box by more than 0.06 in
  EDGE     content whose extent leaves the slide
Text frames that sit on a card are compared by their TEXT extent, not their card, so a card behind text is not a collision.
Usage: sen0414_deck_overlap_check_v1_0_0.py <deck.pptx>        Exit 0 when nothing is reported, 1 otherwise.  Result words: CLEAN / FINDINGS."""
__version__ = "1.0.0"
import sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.text import MSO_ANCHOR

IN = 914400.0
TOL = 0.03


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh


def adv(font_name, size):
    mono = font_name and ("courier" in font_name.lower() or "mono" in font_name.lower())
    return (0.60 if mono else 0.50) * size / 72.0


def para_lines(p, width_in, default_size=12.0):
    """number of wrapped lines and the line height (in) of one paragraph"""
    runs = p.runs
    size = max([r.font.size.pt for r in runs if r.font.size] or [default_size])
    name = next((r.font.name for r in runs if r.font.name), None)
    text = "".join(r.text for r in runs)
    a = adv(name, size)
    wrap = True
    if not text:
        return 1, size * 1.2 / 72.0
    words = text.split(" ")
    lines, cur = 1, 0.0
    for w in words:
        wl = len(w) * a
        sp = a
        if cur == 0:
            cur = wl
        elif cur + sp + wl <= width_in:
            cur += sp + wl
        else:
            lines += 1
            cur = wl
        while cur > width_in and width_in > 0 and wrap and len(w) * a > width_in * 1.0 and False:
            pass
    return lines, size * 1.2 / 72.0


def text_extent(sh):
    tf = sh.text_frame
    l, t, w, h = [v / IN for v in (sh.left or 0, sh.top or 0, sh.width or 0, sh.height or 0)]
    bp = tf._txBody.find("{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr")
    li = int(bp.get("lIns", 91440)) / IN
    ri = int(bp.get("rIns", 91440)) / IN
    ti = int(bp.get("tIns", 45720)) / IN
    bi = int(bp.get("bIns", 45720)) / IN
    wrap = bp.get("wrap", "square") != "none"
    inner = w - li - ri
    total, maxw = 0.0, 0.0
    for p in tf.paragraphs:
        runs = p.runs
        size = max([r.font.size.pt for r in runs if r.font.size] or [12.0])
        name = next((r.font.name for r in runs if r.font.name), None)
        text = "".join(r.text for r in runs)
        if wrap:
            n, lh = para_lines(p, inner)
        else:
            n, lh = 1, size * 1.2 / 72.0
        total += n * lh
        maxw = max(maxw, min(len(text) * adv(name, size), inner if wrap else 99))
    anchor = tf.vertical_anchor
    room = h - ti - bi
    if anchor == MSO_ANCHOR.MIDDLE:
        y0 = t + ti + (room - total) / 2
    elif anchor == MSO_ANCHOR.BOTTOM:
        y0 = t + h - bi - total
    else:
        y0 = t + ti
    al = tf.paragraphs[0].alignment
    x0 = l + li
    ww = maxw if wrap else maxw
    if al is not None and int(al) == 2:        # center
        x0 = l + li + (inner - ww) / 2
    elif al is not None and int(al) == 3:      # right
        x0 = l + w - ri - ww
    return (x0, y0, ww, total), (total > room + 0.06 and wrap and total - room), room


def table_extent(sh):
    tbl = sh.table
    l, t = (sh.left or 0) / IN, (sh.top or 0) / IN
    widths = [c.width / IN for c in tbl.columns]
    tot = 0.0
    for r in tbl.rows:
        need = r.height / IN
        for j, c in enumerate(r.cells):
            tf = c.text_frame
            inner = widths[j] - (c.margin_left + c.margin_right) / IN
            hh = (c.margin_top + c.margin_bottom) / IN
            for p in tf.paragraphs:
                n, lh = para_lines(p, inner, 11.0)
                hh += n * lh
            need = max(need, hh)
        tot += need
    return (l, t, sum(widths), tot)


def inter(a, b):
    ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return ox, oy


def scan(path):
    prs = Presentation(path)
    W, H = prs.slide_width / IN, prs.slide_height / IN
    out = []
    for i, sl in enumerate(prs.slides, 1):
        items = []
        for sh in walk(sl.shapes):
            nm = sh.name
            if getattr(sh, "has_table", False) and sh.has_table:
                items.append(("table " + nm, table_extent(sh)))
            elif getattr(sh, "has_chart", False) and sh.has_chart:
                items.append(("chart " + nm, ((sh.left or 0) / IN, (sh.top or 0) / IN, (sh.width or 0) / IN, (sh.height or 0) / IN)))
            elif sh.shape_type == 13:
                items.append(("picture " + nm, ((sh.left or 0) / IN, (sh.top or 0) / IN, (sh.width or 0) / IN, (sh.height or 0) / IN)))
            elif sh.has_text_frame and sh.text_frame.text.strip():
                ext, spill, room = text_extent(sh)
                items.append(("text '%s'" % sh.text_frame.text.strip().replace("\n", " ")[:28], ext))
                if spill:
                    out.append((i, "SPILL", "text '%s' is %.2f in taller than its box" % (sh.text_frame.text.strip()[:28], spill)))
        for k, (na, a) in enumerate(items):
            if a[0] < -TOL or a[1] < -TOL or a[0] + a[2] > W + TOL or a[1] + a[3] > H + TOL:
                out.append((i, "EDGE", "%s leaves the slide (%.2f..%.2f x %.2f..%.2f in)" % (na, a[0], a[0] + a[2], a[1], a[1] + a[3])))
            for nb, b in items[k + 1:]:
                ox, oy = inter(a, b)
                if ox > TOL and oy > TOL:
                    out.append((i, "OVERLAP", "%s  x  %s  (%.2f by %.2f in)" % (na, nb, ox, oy)))
    return out


if __name__ == "__main__":
    res = scan(sys.argv[1])
    for i, kind, msg in res:
        print("slide %d %s %s" % (i, kind, msg))
    print("RESULT %s (%d findings)" % ("CLEAN" if not res else "FINDINGS", len(res)))
    sys.exit(1 if res else 0)
