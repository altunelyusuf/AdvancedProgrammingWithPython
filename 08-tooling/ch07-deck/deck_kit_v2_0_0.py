#!/usr/bin/env python3
"""Chapter 7 deck v2.0.0 - the drawing kit (a NEW tool file; the v1.0.0 builders and the v1.0.0 deck are untouched): palette, slide chrome,
code boxes, native tables and charts, drawn diagrams. The content is in deck_build_v2_0_0.py.

Built against the deck standard D1-D12 of the owner's teaching-material rules:
 D1 numbered slides   D2 no text box over 30 words   D3 each code slide: explanation beside the code + the executed result as a visual
 D4 an everyday analogy for each key concept   D5 5N1K stories with picture and live link   D6 native charts only   D7 notes <= 60 words
 D8 resources slide   D9 one background   D10 dividers name their concepts, one line each, with one visual   D11 coverage   D12 no code line wraps.

The look is the one of the v1.0.0 deck and of the sibling course's decks (warm ivory background, orange accent, slate green, Cambria and Calibri)."""
__version__ = "2.0.0"
import math, os, re, sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TOOL = os.path.abspath(os.path.join(HERE, ".."))
ROOT = os.path.abspath(os.path.join(TOOL, ".."))
ASSETS = os.path.join(ROOT, "03-materials", "ch07", "assets")

C = dict(dark="2B2622", orange="D97A14", ink="2B2622", mute="6E655C", card="F2ECE1", code="23272E", codeTxt="E6EDF3", green="7EE787",
         white="FFFFFF", slate="2F6D62", rule="DFD6C7", bg="FAF6EF", band="EFE6D6", blue="79C0FF", kw="FFA657", pale="8A9B94")
HF, BF, MONO = "Cambria", "Calibri", "Courier New"
rgb = lambda h: RGBColor.from_string(h)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
LAYOUT = prs.slide_layouts[5]               # Title Only
SLIDES = []                                  # slides in order, for numbering
SRC_PREFIX = "Automate the Boring Stuff 3e, chapter 7: "


def nostyle(shape):
    from pptx.oxml.ns import qn
    el = shape._element.find(qn("p:style"))
    if el is not None:
        shape._element.remove(el)


def fill(shape, hexcol):
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(hexcol)


def _text(shape, text, size, color, bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=BF, margin=0.05, wrap=True):
    tf = shape.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = Inches(min(margin, 0.04))
    paras = text if isinstance(text, list) else [text]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para if isinstance(para, list) else [(para, {})]
        for rt, ro in runs:
            r = p.add_run(); r.text = rt
            f = r.font; f.name = ro.get("font", font); f.size = Pt(ro.get("size", size)); f.bold = ro.get("bold", bold); f.italic = ro.get("italic", italic)
            f.color.rgb = rgb(ro.get("color", color))
            if ro.get("link"):
                r.hyperlink.address = ro["link"]


def new_slide(title, sub=None, src=None, notes="", chip=True):
    s = prs.slides.add_slide(LAYOUT)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(C["bg"])
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(1.25), Inches(0.22), Inches(8.3), Inches(0.62)
    tf = t.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for a in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, a, 0)
    tf.text = title
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.runs[0]; r.font.name = HF; r.font.bold = True; r.font.color.rgb = rgb(C["dark"])
    r.font.size = Pt(26 if len(title) <= 34 else (23 if len(title) <= 46 else 20))
    if chip:
        c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(0.33), Inches(0.62), Inches(0.42))
        c.name = "chip"; fill(c, C["orange"]); c.line.color.rgb = rgb(C["orange"]); nostyle(c)
        _text(c, ">>>", 14, C["white"], bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=MONO, margin=0)
    if sub:
        tb(s, 1.25, 0.84, 8.3, 0.32, sub, 13, C["mute"], italic=True, name="subtitle")
    if src:
        tb(s, 0.5, 5.3, 6.9, 0.22, SRC_PREFIX + src, 9.5, C["mute"], name="source")
    s.notes_slide.notes_text_frame.text = notes
    SLIDES.append(s)
    return s


def tb(s, x, y, w, h, text, size=13, color=None, bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=BF, name=None, wrap=True, margin=0.05):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sh.name = name
    _text(sh, text, size, color or C["ink"], bold, italic, align, anchor, font, margin, wrap)
    return sh


def card(s, x, y, w, h, color=None, name=None, line=None, radius=0.08):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = min(0.5, radius / min(w, h))
    fill(sh, color or C["card"]); sh.line.color.rgb = rgb(line or color or C["card"])
    if name:
        sh.name = name
    nostyle(sh)
    return sh


def box_text(s, x, y, w, h, text, size=12, color=None, fillc=None, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, name=None, line=None, italic=False, font=BF):
    sh = card(s, x, y, w, h, fillc or C["card"], name, line)
    _text(sh, text, size, color or C["ink"], bold, italic, align, anchor, font, 0.08)
    return sh


def takeaway(s, text, y=4.9):
    tb(s, 0.5, y, 9.0, 0.34, text, 13, C["slate"], bold=True, italic=True, name="takeaway", anchor=MSO_ANCHOR.MIDDLE)


def picture(s, fname, x, y, w, h, cap=None, name="picture"):
    from PIL import Image
    path = fname if os.path.isabs(fname) else os.path.join(ASSETS, fname)
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih)
    pw, ph = iw * sc, ih * sc
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    pic = s.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph)); pic.name = name
    pic._element.xpath("./p:nvPicPr/p:cNvPr")[0].set("descr", cap or name)
    if cap:
        tb(s, x, y + h + 0.01, w, 0.36, cap, 9, C["mute"], italic=True, name="credit", align=PP_ALIGN.CENTER)
    return pic


# ---- tables (native), charts (native), diagrams (drawn)
def table(s, x, y, w, rows, colw, fs=11, rowh=0.3, header=True, name="result table", mono_cols=(), bold_first=False):
    nr, nc = len(rows), len(rows[0])
    gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rowh * nr)); gf.name = name
    t = gf.table
    tot = sum(colw)
    for j, cw in enumerate(colw):
        t.columns[j].width = Inches(w * cw / tot)
    for i in range(nr):
        t.rows[i].height = Inches(rowh)
        for j in range(nc):
            c = t.cell(i, j); c.margin_left = c.margin_right = Inches(0.07); c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            hdr = header and i == 0
            c.fill.fore_color.rgb = rgb(C["dark"] if hdr else (C["card"] if i % 2 else C["white"]))
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.text = ""
            r = p.add_run(); r.text = str(rows[i][j]) if str(rows[i][j]) != "" else " "
            ismono = (j in mono_cols and not hdr)
            r.font.size = Pt(fs); r.font.name = MONO if ismono else BF
            r.font.bold = hdr or (bold_first and j == 0)
            r.font.color.rgb = rgb(C["white"] if hdr else C["ink"])
    return gf


def chart(s, kind, cats, series, x, y, w, h, title=None, numfmt="0", colors=None, name="result", fs=11, legend=False, gap=60):
    cd = CategoryChartData(); cd.categories = cats
    for nm, vals in series:
        cd.add_series(nm, vals)
    gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd); gf.name = name
    ch = gf.chart
    ch.font.size = Pt(fs); ch.font.name = BF; ch.font.color.rgb = rgb(C["ink"])
    ch.has_legend = legend
    if legend:
        from pptx.enum.chart import XL_LEGEND_POSITION
        ch.legend.position = XL_LEGEND_POSITION.RIGHT; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(fs)
    if title:
        ch.has_title = True; ch.chart_title.text_frame.text = title
        r = ch.chart_title.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(fs + 1); r.font.bold = True; r.font.name = BF
    else:
        ch.has_title = False
    plot = ch.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels; dl.number_format = numfmt; dl.number_format_is_linked = False; dl.font.size = Pt(fs); dl.font.bold = True
    try:
        plot.gap_width = gap
    except Exception:
        pass
    cols = colors or [C["orange"], C["slate"], C["pale"], C["blue"]]
    for i, ser in enumerate(plot.series):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(cols[i % len(cols)])
    va = ch.value_axis
    va.minimum_scale = 0
    va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = rgb(C["rule"])
    va.tick_labels.font.size = Pt(fs - 1)
    ch.category_axis.tick_labels.font.size = Pt(fs)
    return gf


def group(s, name):
    g = s.shapes.add_group_shape(); g.name = name
    return g


def gbox(g, x, y, w, h, text, size=11, fillc=None, color=None, bold=False, line=None, align=PP_ALIGN.CENTER, font=BF):
    sh = g.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.12
    fill(sh, fillc or C["card"]); sh.line.color.rgb = rgb(line or fillc or C["rule"]); nostyle(sh)
    _text(sh, str(text).split("\n"), size, color or C["ink"], bold, False, align, MSO_ANCHOR.MIDDLE, font, 0.04)
    return sh


def garrow(g, x1, y1, x2, y2, color=None):
    cn = g.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = rgb(color or C["mute"]); cn.line.width = Pt(2); nostyle(cn)
    ln = cn.line._get_or_add_ln()
    from lxml import etree
    tail = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd"); tail.set("type", "triangle")
    return cn


def flow(s, x, y, w, h, labels, name="visual: flow", hot=(), size=11, mono=False):
    """one drawn diagram: boxes joined by arrows, left to right"""
    g = group(s, name); n = len(labels)
    bw = (w - (n - 1) * 0.3) / n
    for i, lab in enumerate(labels):
        gbox(g, x + i * (bw + 0.3), y, bw, h, lab, size, C["orange"] if i in hot else C["card"], C["white"] if i in hot else C["ink"], bold=i in hot,
             font=MONO if mono else BF)
        if i < n - 1:
            garrow(g, x + i * (bw + 0.3) + bw, y + h / 2, x + (i + 1) * (bw + 0.3), y + h / 2)
    return g


def pairs(s, x, y, w, h, items, name="visual: key to value", size=11, kw=None):
    """a drawn dictionary: a column of key boxes, an arrow from each to its value box"""
    g = group(s, name); n = len(items)
    gap = 0.07
    bh = min(0.42, (h - (n - 1) * gap) / n)
    top = y + (h - (n * bh + (n - 1) * gap)) / 2
    kw = kw or w * 0.38
    aw = 0.45
    vw = w - kw - aw
    for i, (k, v) in enumerate(items):
        yy = top + i * (bh + gap)
        gbox(g, x, yy, kw, bh, str(k), size, C["orange"], C["white"], bold=True, font=MONO)
        gbox(g, x + kw + aw, yy, vw, bh, str(v), size, C["card"], C["ink"], font=MONO)
        garrow(g, x + kw, yy + bh / 2, x + kw + aw, yy + bh / 2)
    return g


# ---- code blocks
KW = {"for", "in", "def", "return", "lambda", "if", "else", "while", "try", "except", "as", "import", "from", "class", "match", "case", "del",
      "not", "is", "print", "len", "list", "sorted", "reversed", "sum", "iter", "next", "str", "range", "dict", "type"}
TOK = re.compile(r"(\"(?:[^\"\\]|\\.)*\"|'(?:[^'\\]|\\.)*'|#.*$|\b\d[\w.]*\b|\b[A-Za-z_]\w*\b|\s+|.)")


def code_runs(line):
    runs = []
    for t in TOK.findall(line):
        col = C["codeTxt"]
        if t.startswith('"') or t.startswith("'"):
            col = C["green"]
        elif t.startswith("#"):
            col = C["pale"]
        elif re.match(r"\d", t):
            col = C["blue"]
        elif t in KW:
            col = C["kw"]
        runs.append((t, {"color": col, "font": MONO}))
    return runs


def code_box(s, x, y, w, h, lines, fs, name="code"):
    sh = card(s, x, y, w, h, C["code"], name, radius=0.08)
    paras = [[(">>> ", {"color": C["pale"], "font": MONO})] + code_runs(l) for l in lines]
    _text(sh, paras, fs, C["codeTxt"], font=MONO, anchor=MSO_ANCHOR.MIDDLE, margin=0.1, wrap=False)
    return sh
