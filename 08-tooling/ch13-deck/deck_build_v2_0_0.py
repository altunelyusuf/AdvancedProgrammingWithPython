#!/usr/bin/env python3
"""Builds the chapter 13 lecture deck 2.0.0 (Web Scraping) on the deck standard, and its chapter-code manifest.

A NEW tool (the deck of 1.0.0 and its builder stay untouched, in this folder). Every code example of examples_v2_0_0.py is RUN here; the
printed result and every number in a result visual come from that run. One light background scheme, a number on every slide, a chip '>>>'
on every slide that is not a divider, native tables and charts, drawn diagrams with arrows, story pictures embedded with their licences.
Writes: 03-materials/ch13/SEN0414_Ch13_WebScraping_3e_v2_0_0.pptx
        03-materials/ch13/page/chapter_code_ch13_v1_0_0.json   (order = slide order; the code exactly as the code boxes show it)
Usage:  python3 deck_build_v2_0_0.py
"""
__version__ = "2.0.0"
import contextlib
import glob
import io
import json
import math
import os
import platform
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
import importlib.util


def load(name, fn):
    sp = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(sp)
    sys.modules[name] = m
    sp.loader.exec_module(m)
    return m


from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches
import deck_kit_v2_0_0 as K
from deck_kit_v2_0_0 import C, MONO, arrow, chart, code_box, label_shape, new_deck, new_slide, notes, picture, rect, table, tb
X = load("examples_v2_0_0", "examples_v2_0_0.py")
V = load("deck_visuals_v2_0_0", "deck_visuals_v2_0_0.py")
T = load("deck_content_v2_0_0", "deck_content_v2_0_0.py")

OUT_DECK = os.path.join(ROOT, "03-materials", "ch13", "SEN0414_Ch13_WebScraping_3e_v2_0_0.pptx")
OUT_MAN = os.path.join(ROOT, "03-materials", "ch13", "page", "chapter_code_ch13_v1_0_0.json")
ASSETS = os.path.join(ROOT, "03-materials", "ch13", "assets")
PYV = platform.python_version()
IMG = {}
for _f in ("story_images_v2_0_0.json", "story_images_other_v2_0_0.json"):
    if os.path.exists(os.path.join(ASSETS, _f)):
        IMG.update(json.load(open(os.path.join(ASSETS, _f), encoding="utf-8")))


def research_labels():
    f = sorted(glob.glob(os.path.join(ROOT, "03-materials", "ch13", "rdodi", "sen0414_ch13_research_v*.ttl")), key=lambda p: [int(x) for x in re.findall(r"_v(\d+)_(\d+)_(\d+)", p)[-1]])[-1]
    return re.findall(r"a sen0414:PrimarySourceConcept\s*;\s*rdfs:label\s+\"((?:[^\"\\]|\\.)*)\"", open(f, encoding="utf-8").read())


RESEARCH = research_labels()


def _norm(t):
    t = re.sub(r"\([^)]*\)", " ", t.lower())
    t = re.sub(r"[^a-z0-9]+", " ", t).split()
    return " ".join(w[:-1] if len(w) > 3 and w.endswith("s") else w for w in t)


N_CONCEPTS = len({_norm(l) for l in RESEARCH})


# ------------------------------------------------------------------ run every example
def run_example(ex):
    code = X.code_text(ex)
    buf = io.StringIO()
    ns = {"__name__": "__main__"}
    cwd = os.getcwd()
    with contextlib.redirect_stdout(buf):
        exec(compile(code, "<%s>" % ex["id"], "exec"), ns)
    assert os.getcwd() == cwd, "example %s left the current folder" % ex["id"]
    return buf.getvalue(), ns


RUN = {ex["id"]: run_example(ex) for ex in X.EXAMPLES}

# ------------------------------------------------------------------ the plan: slide kinds in order
PARTS = {
    1: ("Part 1: Addresses and manners", "An address has parts, a reply has a code, and a polite program obeys the site's rules", [
        ("URL", "An address cut into scheme, host, path and query."),
        ("Status code", "A three-digit number whose first digit is the class."),
        ("User agent", "The name a program gives to the server."),
        ("robots.txt", "The paths a crawler is asked not to visit."),
    ]),
    2: ("Part 2: Opening and downloading", "Show a page to a person, or fetch it into the program and check the reply", [
        ("webbrowser", "Opens a page in your own browser and reads nothing."),
        ("requests", "Downloads a page or a file into the program."),
        ("raise_for_status", "Stops the program when the reply is a failure."),
        ("iter_content", "Saves a big download in small chunks."),
    ]),
    3: ("Part 3: A weather API", "A service answers in JSON; the program reads it and keeps its key secret", [
        ("API", "A service you call by address and read as data."),
        ("API key", "A secret that names your account to the service."),
        ("json.loads", "Turns JSON text into dictionaries and lists."),
        ("Kelvin", "The service's temperature scale; subtract 273.15."),
    ]),
    4: ("Part 4: Reading HTML", "A page is a tree of elements; a parser, not a pattern, finds what you need", [
        ("HTML tags", "Nested elements, each with its own attributes."),
        ("select", "Finds the elements that match a selector path."),
        ("get_text, get", "Read the text, or one attribute, of an element."),
        ("Regex trap", "A pattern misses links that a parser finds."),
    ]),
    5: ("Part 5: Driving a browser", "When a script builds the page, only a real browser can read it", [
        ("Dynamic page", "A script writes the content after the download."),
        ("Selenium", "Drives a real browser through a separate driver."),
        ("Playwright", "Drives a browser and waits by itself."),
        ("Headless", "A browser that runs with no window at all."),
    ]),
}

PLAN = [("title",), ("route",), ("questions",),
        ("divider", 1), ("analogy", "A1"), ("code", "url_anatomy"), ("code", "query_quote"), ("code", "status_codes"), ("code", "user_agent"), ("code", "robots_rules"),
        ("story", "koster"), ("code", "crawl_delay"), ("story", "gray"),
        ("divider", 2), ("analogy", "A2"), ("analogy", "A3"), ("code", "webbrowser_open"), ("code", "showmap"), ("code", "fetch_data"), ("story", "reitz"), ("code", "raise_status"),
        ("code", "chunks"),
        ("divider", 3), ("analogy", "A4"), ("code", "json_loads"), ("code", "kelvin"), ("code", "api_url"),
        ("divider", 4), ("analogy", "A5"), ("analogy", "A6"), ("code", "html_elements"), ("code", "css_select"), ("story", "richardson"), ("code", "regex_pitfall"),
        ("code", "links_verify"), ("code", "xkcd_chain"), ("code", "search_results"),
        ("divider", 5), ("analogy", "A7"), ("analogy", "A8"), ("code", "dynamic_page"), ("story", "huggins"), ("code", "selenium_shape"), ("code", "playwright_shape"),
        ("story", "playwright"), ("code", "auto_wait"),
        ("tools",), ("current",), ("summary",), ("practice",), ("resources",)]
N = len(PLAN)
SLIDE_OF = {p: i for i, p in enumerate(PLAN, 1)}
CODE_ORDER = [p[1] for p in PLAN if p[0] == "code"]
assert CODE_ORDER == [e["id"] for e in X.EXAMPLES], "the examples file must list the examples in slide order"
STORY_COUNT = sum(1 for p in PLAN if p[0] == "story")


def part_slides(k):
    a = SLIDE_OF[("divider", k)]
    b = SLIDE_OF[("divider", k + 1)] - 1 if k < 5 else SLIDE_OF[("tools",)] - 1
    return a, b


# ------------------------------------------------------------------ slide frames
def base(prs, title, sub=None, chip=True, nt=None):
    s = new_slide(prs)
    if chip:
        label_shape(s, 0.4, 0.28, 0.6, 0.36, ">>>", C["teal"], "FFFFFF", size=13, bold=True, font=MONO, radius=0.2, name="chip")
    tb(s, 1.15 if chip else 0.5, 0.18, 8.4 if chip else 9.0, 0.55, title, 22 if len(title) < 52 else 19, C["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, name="title")
    if sub:
        tb(s, 1.15 if chip else 0.5, 0.72, 8.4, 0.3, sub, 12, C["muted"], name="subtitle")
    if nt:
        notes(s, nt)
    return s


def takeaway(s, text):
    rect(s, 0.4, 5.22, 0.07, 0.3, C["teal"], name="accent")
    tb(s, 0.55, 5.18, 7.9, 0.36, text, 12, C["muted"], italic=True, anchor=MSO_ANCHOR.MIDDLE, name="takeaway")


def add_numbers(prs):
    n = len(prs.slides)
    for i, s in enumerate(prs.slides, 1):
        tb(s, 8.6, 5.2, 1.0, 0.32, "%d / %d" % (i, n), 11, C["muted"], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, name="slide number")


def analogy_box(s, y, text, h=0.55, size=13):
    rect(s, 0.4, y, 9.2, h, C["amberLight"], radius=0.08, name="card")
    tb(s, 0.55, y + 0.01, 8.9, h - 0.02, [[("ANALOGY  ", {"bold": True, "color": C["amber"]}), (text, {})]], size, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="analogy")


# ------------------------------------------------------------------ slide makers
def make_title(prs):
    s = base(prs, "Web Scraping", None, chip=True, nt=T.NOTES["title"])
    s.shapes[-1].left, s.shapes[-1].top, s.shapes[-1].width, s.shapes[-1].height = Inches(0.6), Inches(1.7), Inches(8.8), Inches(0.9)
    for r in s.shapes[-1].text_frame.paragraphs[0].runs:
        r.font.size = K.Pt(38)
    tb(s, 0.6, 2.65, 8.8, 0.7, "Chapter 13 of Automate the Boring Stuff with Python, 3rd edition, for an advanced course", 16, C["muted"], name="subtitle")
    for i, (a, b) in enumerate([(str(len(X.EXAMPLES)), "code examples that run on the chapter page"), (str(STORY_COUNT), "true stories with who, what, when, where, why and how"),
                                (str(N_CONCEPTS), "distinct chapter concepts, each with an everyday picture")]):
        xx = 0.6 + i * 3.0
        rect(s, xx, 3.5, 2.8, 1.15, C["card"], line=C["line"], radius=0.08, name="card")
        tb(s, xx + 0.1, 3.55, 2.6, 0.5, a, 28, C["teal"], bold=True, name="stat")
        tb(s, xx + 0.1, 4.05, 2.6, 0.55, b, 11.5, C["ink"], name="stat label")
    tb(s, 0.6, 4.85, 8.0, 0.3, "SEN0414 Advanced Programming, Fall 2026. Yusuf Altunel, İstanbul Kültür Üniversitesi", 11, C["muted"], name="author")
    return s


def make_route(prs):
    s = base(prs, "Today's route", "Five parts; every range is a run of slides you can jump to", nt=T.NOTES["route"])
    rows = [["Part", "Slides", "Code examples", "You will run"]]
    runs = {1: "urlsplit, quote, status codes, robots.txt", 2: "webbrowser, showmap, a download, chunks", 3: "json.loads, Kelvin, an endpoint and its key",
            4: "tags, select, links, comics, search results", 5: "a script-built page, Selenium and Playwright shapes"}
    for k in range(1, 6):
        a, b = part_slides(k)
        nex = sum(1 for p in PLAN[a - 1:b] if p[0] == "code")
        rows.append([PARTS[k][0].split(": ", 1)[1], "%d - %d" % (a, b), str(nex), runs[k]])
    rows.append(["Tools, changes, summary, practice, resources", "%d - %d" % (SLIDE_OF[("tools",)], N), "-", "which tool, and what current Python changes"])
    table(s, 0.5, 1.2, 9.0, rows, [3.0, 1.1, 1.3, 3.8], fs=13, rowh=0.55, bold_cols=(0,))
    takeaway(s, "Open the chapter-code tab of the page: the same %d examples, in this order." % len(X.EXAMPLES))
    return s


def make_questions(prs):
    s = base(prs, "Two questions this chapter answers", "Carry them through every part; the answers come back at the end", nt=T.NOTES["questions"])
    q = [("1", "Which tool reaches which kind of page, and what decides the choice?",
          "Answer on slide %d: webbrowser, requests, an API, a parser, or a browser driver." % SLIDE_OF[("tools",)]),
         ("2", "Which printed outputs and programs of the book differ in current Python and on today's web?",
          "Answer on slide %d: the changes, from the chapter research." % SLIDE_OF[("current",)])]
    for i, (n, qq, aa) in enumerate(q):
        yy = 1.3 + i * 1.75
        rect(s, 0.5, yy, 9.0, 1.55, C["card"], line=C["line"], radius=0.05, name="card")
        label_shape(s, 0.7, yy + 0.2, 0.5, 0.5, n, C["teal"], "FFFFFF", size=18, bold=True)
        tb(s, 1.4, yy + 0.12, 7.9, 0.75, qq, 16, C["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, name="question")
        tb(s, 1.4, yy + 0.95, 7.9, 0.45, aa, 12.5, C["teal"], anchor=MSO_ANCHOR.MIDDLE, name="answer")
    takeaway(s, "Choose the tool by the page in front of the program.")
    return s


def divider_item_title(p):
    if p[0] == "code":
        return X.BY_ID[p[1]]["title"]
    if p[0] == "story":
        return T.STORIES[p[1]]["title"]
    if p[0] == "analogy":
        return [a_[1] for a_ in T.ANALOGIES if a_[0] == p[1]][0]
    raise SystemExit("no divider title for %r" % (p,))


def make_divider(prs, k):
    title, sub, concepts = PARTS[k]
    s = new_slide(prs)
    notes(s, T.NOTES["divider%d" % k])
    tb(s, 0.5, 0.2, 9.0, 0.7, title, 30, C["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, name="title")
    tb(s, 0.5, 0.92, 9.0, 0.34, sub, 14, C["muted"], name="subtitle")
    rect(s, 0.5, 1.33, 1.2, 0.05, C["teal"], name="rule")
    yy = 1.55
    for name, mean in concepts:
        label_shape(s, 0.5, yy, 2.2, 0.5, name, C["tealLight"], C["ink"], size=14, bold=True, radius=0.3)
        tb(s, 2.85, yy, 2.8, 0.6, mean, 12.5, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="meaning")
        yy += 0.72
    a, b = part_slides(k)
    items = [(i, divider_item_title(PLAN[i - 1])) for i in range(a + 1, b + 1)]
    rows = [["Slide", "In this part"]] + [[str(i), t if len(t) < 44 else t[:41] + "..."] for i, t in items]
    rh = min(0.4, 3.5 / len(rows))
    table(s, 5.8, 1.55, 3.8, rows, [0.65, 3.8], fs=10 if len(rows) > 9 else 10.5, rowh=rh, name="table")
    return s


def make_analogy(prs, aid):
    _, title, head, rows = [a for a in T.ANALOGIES if a[0] == aid][0]
    s = base(prs, title, None, nt=T.NOTES[aid])
    analogy_box(s, 0.82, head, 0.62, 14)
    rh = min(0.5, 3.4 / (len(rows) + 1))
    table(s, 0.4, 1.55, 9.2, [["Chapter concept", "Everyday picture"]] + [[a, b] for a, b in rows], [3.9, 5.3], fs=11 if len(rows) > 5 else None, rowh=rh, bold_cols=(0,))
    return s


def framed(path, w, h):
    """Return a picture file of the frame's proportions: a wide or tall picture is centred, undistorted, on a white card so it fills the frame."""
    import tempfile
    from PIL import Image
    im = Image.open(path).convert("RGBA")
    k = max(im.width / w, im.height / h)
    cw, ch = int(round(w * k)), int(round(h * k))
    bg = Image.new("RGBA", (cw, ch), (255, 255, 255, 255))
    bg.alpha_composite(im, ((cw - im.width) // 2, (ch - im.height) // 2))
    fd, out = tempfile.mkstemp(suffix=".png", prefix="ch13_story_")
    os.close(fd)
    bg.convert("RGB").save(out)
    return out


def make_story(prs, sid):
    st = T.STORIES[sid]
    rec = IMG[st["img"]]
    desc = " ".join(re.sub(r"\s+", " ", rec.get("ImageDescription", "")).split()[:10])
    credit = "%s. Image: %s, %s, Wikimedia Commons." % (desc.rstrip(".,;"), (rec["Artist"] or "see file page").replace("\n", " "), rec["LicenseShortName"])
    if not rec.get("SlideCredit") and not rec["LicenseShortName"].startswith(("CC", "Public")):
        credit = "%s. Logo: %s, %s." % (desc.rstrip(".,;"), rec["Credit"], rec["LicenseShortName"].split(" (")[0])
    if st["img"] == "image_richardson_tortoise.jpg":
        credit = "Tenniel's Mock Turtle (1865, public domain), as shown on crummy.com."
    if rec.get("SlideCredit"):
        credit = rec["SlideCredit"]
    s = base(prs, st["title"], None, nt=T.STORY_NOTES[sid])
    picture(s, framed(os.path.join(ASSETS, st["img"]), 2.9, 3.2), 0.4, 0.95, 2.9, 3.2, "Picture for the story: " + st["title"].replace("Story: ", ""), "photo " + sid)
    tb(s, 0.4, 4.2, 2.9, 0.7, credit, 9.5, C["muted"], name="credit")
    yy = 0.92
    for lab in ("who", "what", "when", "where", "why", "how"):
        txt = st[lab]
        nlines = math.ceil((len(txt) + 7) / 64.0)
        h = max(0.44, nlines * 0.235 + 0.14)
        rect(s, 3.5, yy, 6.1, h, C["card"], line=C["line"], radius=0.1, name="card")
        tb(s, 3.55, yy, 6.0, h, [[(lab.upper() + "  ", {"bold": True, "color": C["teal"]}), (txt, {})]], 13, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name=lab)
        yy += h + 0.07
    if yy > 4.85:
        print("WARN story %s: the READ line starts at %.2f in" % (sid, yy))
    tb(s, 3.5, yy + 0.02, 6.1, 0.36, [[("READ  ", {"bold": True, "color": C["amber"]}), (st["linktext"], {"link": st["link"], "color": C["teal"]})]], 13, C["ink"], name="read")
    return s


def plan_code(ex, visual):
    """Choose layout and font: the largest code font (14 down to 9 pt) for which the code, its explanations and the result visual all fit.
    Order of preference: the explanation beside the code (side) at 11 pt or more, then below the code at 11 pt or more, then smaller."""
    maxlen = max(len(l) for l in X.code_lines(ex)) + 4
    forced_below = ex.get("layout") == "below"
    cands = []
    for lay in (["below"] if forced_below else ["side", "below"]):
        for fs in (14, 13, 12, 11, 10.5, 10, 9.5, 9):
            cands.append((lay, fs))
    # from 11 pt up the layout comes first (side before below, each at every size); below 11 pt the larger font wins, side before below at each size
    cands.sort(key=lambda c: (0, 0) if c[1] >= 11 else (1, -c[1]))
    top, gapv = 0.86, 0.05
    best = None
    for lay, fs in cands:
        need_w = (maxlen * 0.6 * fs + 14.4) / 72.0 + 0.08
        if lay == "side":
            box_w = max(3.2, need_w)
            if box_w > 6.2:
                continue
            expl_w = 9.2 - box_w + 0.02
        else:
            box_w, expl_w = 9.2, 9.2
            if need_w > 9.2:
                continue
        nfs = 12.5 if fs >= 12 else (12 if fs >= 10.5 else 11)
        lh = fs * 1.22 / 72.0
        cpl = max(10, (expl_w - 0.3) * 72.0 / (0.54 * nfs))
        rows = []
        for g, note in ex["groups"]:
            code_h = len(g) * lh + 0.15
            n_h = math.ceil(len(note) / cpl) * nfs * 1.25 / 72.0 + 0.14
            rows.append((code_h, n_h))
        heights = [max(a, b, 0.42) for a, b in rows] if lay == "side" else [a + b for a, b in rows]
        total = sum(heights) + gapv * (len(heights) - 1)
        ry = top + total + 0.08
        avail = 5.04 - (ry + 0.3)
        cur = (lay, fs, nfs, box_w, expl_w, rows, heights, ry, avail)
        if avail >= visual.need:
            return cur
        if best is None or avail > best[-1] + 1e-9:
            best = cur
    if best is None:
        raise SystemExit("REFUSED: %s has no layout in which its longest code line fits the slide" % ex["id"])
    print("WARN %s: no layout leaves %.2f in for the result (wants %.2f; fonts %.1f pt)" % (ex["id"], best[-1], visual.need, best[1]))
    return best


def make_code(prs, exid):
    ex = X.BY_ID[exid]
    out, ns = RUN[exid]
    s = base(prs, ex["title"], None, nt=T.NOTES[exid])
    visual = V.VISUALS[ex["visual"]]
    lay, fs, nfs, box_w, expl_w, rows, heights, ry, avail = plan_code(ex, visual)
    y, gapv = 0.86, 0.05
    for (g, note), (ch_, nh_), hh in zip(ex["groups"], rows, heights):
        if lay == "side":
            code_box(s, 0.4, y, box_w, hh, g, fs)
            tb(s, 0.4 + box_w - 0.02, y, expl_w, hh, note, nfs, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="explanation", margin=(0.14, 0.02))
        else:
            code_box(s, 0.4, y, box_w, ch_, g, fs)
            tb(s, 0.4, y + ch_ + 0.01, box_w, nh_, note, nfs, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="explanation", margin=(0.12, 0.02))
        y += hh + gapv
    tb(s, 0.4, ry, 6.0, 0.26, "What the run printed (Python %s)" % PYV, 11, C["teal"], bold=True, name="result label")
    visual(s, 0.4, ry + 0.3, 9.2, max(avail, 1.0), ns, out)
    takeaway(s, T.TAKEAWAY[exid])
    return s


def make_tools(prs):
    s = base(prs, "Which tool for which page?", "Choose by the page in front of the program", nt=T.NOTES["tools"])
    table(s, 0.5, 1.2, 9.0, T.TOOLS, [3.2, 2.4, 3.8], rowh=0.55, bold_cols=(0,), fs=12.5)
    tb(s, 0.5, 4.55, 9.0, 0.5, "Always: check the reply, pass a timeout, read robots.txt, wait between pages and keep the key secret.", 12, C["muted"], italic=True, name="note")
    takeaway(s, "The simplest tool that reaches the page is the right tool.")
    return s


def make_current(prs):
    s = base(prs, "Current Python and today's web change more", "What differs from the book, from the chapter research", nt=T.NOTES["current"])
    table(s, 0.5, 1.1, 9.0, T.CURRENT, [4.3, 1.3, 3.4], rowh=0.43, bold_cols=(0,), fs=10.5, align_cols={1: PP_ALIGN.CENTER})
    takeaway(s, "Run what the book printed; then check what changed.")
    return s


def make_summary(prs):
    s = base(prs, "Summary", None, nt=T.NOTES["summary"])
    analogy_box(s, 0.82, "A summary is a packing list: only what you really need goes into the bag.", 0.5, 13)
    for i, t_ in enumerate(T.SUMMARY_POINTS):
        yy = 1.5 + i * 0.62
        label_shape(s, 0.4, yy + 0.04, 0.4, 0.4, str(i + 1), C["teal"], "FFFFFF", size=13, bold=True)
        tb(s, 0.9, yy, 4.6, 0.52, t_, 11.5, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="point")
    per_part = [sum(1 for e in X.EXAMPLES if e["part"] == k) for k in range(1, 6)]
    tb(s, 5.7, 1.4, 3.9, 0.5, "code examples in each part of this deck", 10.5, C["teal"], bold=True, name="chart label")
    chart(s, "col", 5.6, 1.95, 4.0, 3.0, ["Part %d" % k for k in range(1, 6)], [("examples", per_part)], fs=10, legend=False, gap=50, vmax=8, major=2)
    takeaway(s, "Choose the tool by the page; behave politely.")
    return s


def make_practice(prs):
    s = base(prs, "Practice Questions", "Predict first, then run what you can on the chapter page", nt=T.NOTES["practice"])
    analogy_box(s, 1.05, "Practice questions are finger scales for a pianist; the practice programs are the small recital pieces.", 0.5, 13)
    for i, q in enumerate(T.PRACTICE):
        yy = 1.75 + i * 0.8
        rect(s, 0.4, yy, 9.2, 0.68, C["card"], line=C["line"], radius=0.08, name="card")
        label_shape(s, 0.55, yy + 0.14, 0.4, 0.4, str(i + 1), C["teal"], "FFFFFF", size=13, bold=True)
        tb(s, 1.1, yy + 0.04, 8.4, 0.6, q, 13, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="question")
    takeaway(s, "Each answer fits in a few lines you can run.")
    return s


def make_resources(prs):
    s = base(prs, "Resources: where to go next", "Reading and watching, chosen for this chapter; every link is live", nt=T.NOTES["resources"])
    for i, (t, url, host) in enumerate(T.RESOURCES):
        col, row = divmod(i, 7)
        xx = 0.4 + col * 4.7
        yy = 1.1 + row * 0.57
        rect(s, xx, yy, 4.5, 0.5, C["card"], line=C["line"], radius=0.1, name="card")
        tb(s, xx + 0.08, yy + 0.01, 4.35, 0.48, [[(t, {"bold": True})], [(url if len(url) < 58 else host, {"link": url, "color": C["teal"], "size": 9})]], 10.5, C["ink"], anchor=MSO_ANCHOR.MIDDLE, name="resource")
    takeaway(s, "Pick one link and read it before the next class.")
    return s


MAKERS = {"title": make_title, "route": make_route, "questions": make_questions, "tools": make_tools, "current": make_current, "summary": make_summary,
          "practice": make_practice, "resources": make_resources}


# ------------------------------------------------------------------ build
def audit(prs):
    """The builder's own look at the standard's measurable limits (the independent checks are run by their owners)."""
    bad = []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and len(sh.text_frame.text.split()) > 30:
                bad.append("slide %d: a text box of %d words" % (i, len(sh.text_frame.text.split())))
        if s.has_notes_slide and len(s.notes_slide.notes_text_frame.text.split()) > 60:
            bad.append("slide %d: a note of %d words" % (i, len(s.notes_slide.notes_text_frame.text.split())))
    return bad


def build():
    prs = new_deck()
    for p in PLAN:
        kind = p[0]
        if kind in MAKERS:
            MAKERS[kind](prs)
        elif kind == "divider":
            make_divider(prs, p[1])
        elif kind == "analogy":
            make_analogy(prs, p[1])
        elif kind == "story":
            make_story(prs, p[1])
        elif kind == "code":
            make_code(prs, p[1])
    add_numbers(prs)
    bad = audit(prs)
    if bad:
        raise SystemExit("REFUSED: " + "; ".join(bad))
    prs.core_properties.title = "SEN0414 Chapter 13: Web Scraping (lecture deck 2.0.0)"
    prs.core_properties.author = "Yusuf Altunel"
    os.makedirs(os.path.dirname(OUT_DECK), exist_ok=True)
    prs.save(OUT_DECK)
    return len(prs.slides)


def write_manifest():
    exs = []
    for k, ex in enumerate(X.EXAMPLES, 1):
        out, _ = RUN[ex["id"]]
        exs.append({"order": k, "id": ex["id"], "name": ex["name"], "slide": SLIDE_OF[("code", ex["id"])], "concepts": ex["concepts"],
                    "code": X.code_text(ex), "expected": out})
    names = [e["name"] for e in exs]
    assert len(set(names)) == len(names)
    man = {"chapter": 13, "version": "1.0.0", "kind": "chapter-code manifest",
           "deck": os.path.relpath(OUT_DECK, ROOT).replace(os.sep, "/"),
           "about": "The code examples of the chapter 13 deck, in the order of the code slides (order = slide order). The page's chapter-code tab lists them by NAME in this order; "
                    "every example is self-contained and runs with the Python standard library only, opens no network connection, reads no keyboard and opens no window (a page in a "
                    "browser cannot reach the web, so the third-party tools of the chapter are shown by their standard-library counterparts and by small stand-ins, named beside the code); "
                    "the expected result is what the run printed in a fresh interpreter. The code is exactly the text of the deck's code boxes with the '>>> ' prompt removed.",
           "executed_with": {"python": PYV, "library": None},
           "library": {"file": None, "sha256": None, "data_files": [],
                       "page_requirements": ["The Python standard library modules urllib.parse, urllib.request (urlopen on a data: address), urllib.robotparser, http, webbrowser (the examples replace webbrowser.open by a recording function), json, os, tempfile, re, contextlib, html.parser and xml.etree.ElementTree (imported by the examples that use them), and the built-ins print, repr, len, range, enumerate, round, type, list, dict",
                                             "A writable temporary folder (the chunks example); no subprocess, no keyboard input, no network, no browser window"]},
           "prelude": "",
           "examples": exs}
    os.makedirs(os.path.dirname(OUT_MAN), exist_ok=True)
    with open(OUT_MAN, "w", encoding="utf-8") as f:
        json.dump(man, f, ensure_ascii=False, indent=1)
        f.write("\n")


if __name__ == "__main__":
    n = build()
    write_manifest()
    print("deck %s | %d slides | %d code examples | manifest %s" % (os.path.relpath(OUT_DECK, ROOT), n, len(X.EXAMPLES), os.path.relpath(OUT_MAN, ROOT)))
