#!/usr/bin/env python3
"""Makes the fixture decks of sen0414_deck_teaching_check_v1_0_0.py into fixtures/deck_teaching_v1_0_0/: a clean deck, one deck per defect (each breaks exactly one element
T1..T6) and a deck without code, with expected.json (the known answer of every fixture). The decks hold synthetic text only. Usage: python3 sen0414_deck_teaching_fixtures_make_v1_0_0.py"""
__version__ = "1.0.0"
import json, os
from pptx import Presentation
from pptx.util import Inches, Pt

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "deck_teaching_v1_0_0")
NOTE45 = ("Tell the class what the slide shows in two short sentences. Give one example from the code on the slide and run it. "
          "Then ask the class one question about the result and let two students answer before you go on.")
NOTE30 = "This slide opens a part of the chapter. Name the concepts of the part and say which example comes first. Ask the class what they already know."
GOOD = "def total(items):\n    s = 0\n    for x in items:\n        s = s + x\n    return s\nprint(total([1, 2, 3]))"


def box(s, x, y, w, h, text, mono=False, size=14):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    for k, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.name = "Courier New" if mono else "Calibri"
    return tb


def deck(name, code=GOOD, text="A function is a recipe that returns a value.", note_content=NOTE45, title_note=NOTE30, extra_text=None, extra_note=None, with_code=True, code2=None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    lay = prs.slide_layouts[6]
    s = prs.slides.add_slide(lay); box(s, 0.5, 1.5, 9, 1, "Fixture title"); s.notes_slide.notes_text_frame.text = title_note
    s = prs.slides.add_slide(lay); box(s, 0.5, 0.2, 9, 0.7, "Part 1: Fixture divider"); s.notes_slide.notes_text_frame.text = title_note
    for k in range(2):
        s = prs.slides.add_slide(lay)
        box(s, 0.4, 0.28, 0.6, 0.36, ">>>", mono=True, size=13)
        box(s, 1.15, 0.2, 8, 0.5, "Content slide %d" % (k + 1))
        box(s, 0.5, 4.3, 8, 0.6, text if k == 0 else (extra_text or text))
        if with_code:
            box(s, 0.5, 1.0, 6, 3, code if k == 0 else (code2 or GOOD), mono=True, size=12)
        s.notes_slide.notes_text_frame.text = note_content if k == 0 else (extra_note or NOTE45)
    prs.save(os.path.join(OUT, name))


os.makedirs(OUT, exist_ok=True)
deck("clean.pptx")
deck("defect_t1_indent.pptx", code="if True:\n  print(1)", code2="for i in range(2): print(i)")
deck("defect_t2_lambda.pptx", code="double = lambda x: x * 2\nprint(double(4))")
deck("defect_t3_labels.pptx", text="WHO Ken Thompson and Rob Pike designed it.", extra_note=NOTE45 + " When: 1992 in a diner.")
deck("defect_t4_prompts.pptx", text=">>> x = 1\n>>> print(x)", extra_note=NOTE45)
deck("defect_t5_notes.pptx", note_content="Too short to teach anything.")
deck("defect_t6_compile.pptx", code="print('a'\nx = 1")
deck("nocode.pptx", with_code=False)
exp = {"clean.pptx": {}, "defect_t1_indent.pptx": {"T1": "FAIL"}, "defect_t2_lambda.pptx": {"T2": "FAIL"}, "defect_t3_labels.pptx": {"T3": "FAIL"},
       "defect_t4_prompts.pptx": {"T4": "FAIL"}, "defect_t5_notes.pptx": {"T5": "FAIL"}, "defect_t6_compile.pptx": {"T6": "FAIL"}, "nocode.pptx": {"T1": "NOT_MEASURED", "T6": "NOT_MEASURED"}}
full = {k: dict({t: "PASS" for t in ("T1", "T2", "T3", "T4", "T5", "T6")}, **v) for k, v in exp.items()}
json.dump(full, open(os.path.join(OUT, "expected.json"), "w", encoding="utf-8"), indent=1)
print("fixtures written:", OUT, len(full))
