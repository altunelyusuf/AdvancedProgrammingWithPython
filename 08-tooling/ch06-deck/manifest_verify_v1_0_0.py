#!/usr/bin/env python3
"""manifest_verify v1.0.0 - the builder's own consistency check of the chapter 6 deck against its chapter-code manifest (NOT the independent deck check).
Reads the deck, takes the code boxes (paragraphs starting '>>> ') slide by slide, and compares them with chapter_code_ch06_v1_0_0.json: order, slide number,
code text (prompt removed) and, by running every example under the given interpreters, the expected result.
Usage: manifest_verify_v1_0_0.py [python ...]     default interpreters: python3.14 (the recorded one), python3.13, python3.12"""
import json, os, subprocess, sys, tempfile, shutil
from pptx import Presentation
R = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
man = json.load(open(os.path.join(R, "03-materials/ch06/page/chapter_code_ch06_v1_0_0.json"), encoding="utf-8"))
deck = Presentation(os.path.join(R, man["deck"]))
bad, got = [], []
for i, s in enumerate(deck.slides, 1):
    lines = []
    for sh in s.shapes:
        if sh.has_text_frame and any(p.text.startswith(">>> ") for p in sh.text_frame.paragraphs):
            assert all(p.text.startswith(">>> ") for p in sh.text_frame.paragraphs), "a code line without the prompt on slide %d" % i
            lines += [p.text[4:] for p in sh.text_frame.paragraphs]
    if lines:
        got.append((i, "\n".join(lines)))
ex = man["examples"]
if [e["order"] for e in ex] != list(range(1, len(ex) + 1)): bad.append("order numbers")
if [e["slide"] for e in ex] != [g[0] for g in got]: bad.append("slides %s vs %s" % ([e["slide"] for e in ex], [g[0] for g in got]))
for e, g in zip(ex, got):
    if e["code"] != g[1]: bad.append("code differs on slide %d (example %s)" % (g[0], e["id"]))
names = [e["name"] for e in ex]
if len(set(names)) != len(names) or not all(names): bad.append("names")
print("manifest vs deck: %d examples, %d code slides: %s" % (len(ex), len(got), "FAIL " + "; ".join(bad) if bad else "PASS (order, slide numbers, names unique, code text equal line for line)"))
pys = sys.argv[1:] or [shutil.which("python3.14") or "/root/.local/bin/python3.14", shutil.which("python3.13") or "/root/.local/bin/python3.13", shutil.which("python3.12") or "/root/.local/bin/python3.12"]
fail = bool(bad)
for py in pys:
    mism = []
    ver = subprocess.run([py, "--version"], capture_output=True, text=True).stdout.strip()
    for e in ex:
        with tempfile.TemporaryDirectory() as d:
            f = os.path.join(d, "x.py"); open(f, "w", encoding="utf-8").write(e["code"] + "\n")
            r = subprocess.run([py, "-I", f], capture_output=True, text=True, cwd=d, timeout=120)
        if r.returncode != 0 or r.stdout != e["expected"]:
            mism.append(e["id"])
    print("run under %s: %d of %d examples print exactly the expected result%s" % (ver, len(ex) - len(mism), len(ex), "" if not mism else " - DIFFER: " + ",".join(mism)))
    fail = fail or bool(mism)
print("RESULT", "FAIL" if fail else "PASS")
sys.exit(1 if fail else 0)
