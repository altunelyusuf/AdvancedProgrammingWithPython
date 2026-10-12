"""Chapter 17 deck examples: every console session, whole program (or excerpt), chart and table the lecture deck shows. Nothing here
carries a result; examples_run_v1_0_0.py runs it all under the interpreter given (the uv Python 3.14.6 of the chapter 17 research
environment, with pypdf 6.20.1, python-docx 1.2.0, pdfminer.six 20260107, Pillow 12.3.0 and cryptography 50.0.2) and records the
results. A console group is one session: its rows run in order in one namespace (the deck check re-runs them per slide, one namespace
per slide). No console row touches the network or the file system: each uses only a document made in memory, so each can be typed into a
Python in a browser that offers pypdf, python-docx, Pillow and re (not tried in a browser by this build). A row that would print an
object whose repr holds a memory address is written as an assignment to _ or reads an attribute instead.

Programs come from sen0414_ch17_corpus_base_v1_0_0.py, all twelve, run and recorded under their own names. A corpus line is often wider than
a slide can show without wrapping, so the slides show deck-written copies (the corpus program cut into pieces that fit a slide, with
comments); the first lines of every program are the corpus helper make_pdf() and are not shown. SUBSET lists the copies whose printed
lines must all be lines the corpus original printed, in the same order; EQUIV lists copies whose transcript must equal the original's.
Project 12 has no corpus program; combinePdfsCut is written for the deck and executed like any other.
KIND says what a program's output depends on, so that honesty_check can require the matching words on the slide that shows it:
  live  the program reads a file the book ships (Recursion_Chapter1.pdf or demo.docx), fetched from nostarch.com at run time and
        checked against its sha256 before use
  none  the output depends on nothing outside the program: documents it builds itself in a temporary folder"""
__version__ = "1.0.0"
import importlib.util
import os

HERE = os.path.dirname(os.path.abspath(__file__))
_s = importlib.util.spec_from_file_location("base17", os.path.join(HERE, "..", "sen0414_ch17_corpus_base_v1_0_0.py"))
B = importlib.util.module_from_spec(_s)
_s.loader.exec_module(B)
_HEAD, _REAL = B.HEAD, B.REAL

CONSOLES = {
    "title": ["import platform, pypdf, docx", "platform.python_version()", "pypdf.__version__", "docx.__version__"],
    "names": ["import pypdf", "hasattr(pypdf, 'PdfReader')", "hasattr(pypdf, 'PdfWriter')",
              "hasattr(pypdf, 'PdfFileReader')", "hasattr(pypdf, 'PdfMerger')"],
    "hyphen": ["import re", "t = 'follow-\\ning the execution'", "t", "re.sub(r'-\\n', '', t)", "re.sub(r'-\\n', '', 'well-\\nknown')"],
    "scan": ["import io, pypdf", "from PIL import Image", "buf = io.BytesIO()", "Image.new('RGB', (40, 40), 'white').save(buf, 'PDF')",
             "page = pypdf.PdfReader(buf).pages[0]", "page.extract_text()", "[(i.name, i.data[:2]) for i in page.images]",
             "pypdf.PdfReader(io.BytesIO(b''))"],
    "zip": ["import io, zipfile, docx", "d = docx.Document()", "_ = d.add_paragraph('Hello, world!')", "buf = io.BytesIO()",
            "d.save(buf)", "buf.getvalue()[:2]", "z = zipfile.ZipFile(buf)", "len(z.namelist())", "'word/document.xml' in z.namelist()",
            "z.read('word/document.xml').decode().split('<w:body>')[1][:46]"],
}

# ---- programs written for the deck ---------------------------------------------------------------------------------------------
STRUCT = _HEAD + '''# The slide shows this program from line 16: make_pdf() above builds a small valid PDF of text pages.
data = make_pdf('Page one', 'Page two')  # two pages
print(data[:8], data.rstrip().endswith(b'%%EOF'))  # header, last line
start = int(re.search(rb'startxref\\n(\\d+)', data).group(1))
print('startxref points at', data[start:start + 4])  # the table
rows = re.findall(rb'(\\d{10}) (\\d{5}) ([nf]) \\n', data[start:])
print(len(rows), 'table entries; object 0 is',
      'free' if rows[0][2] == b'f' else 'in use')
print([data[int(o):int(o) + 7] for o, g, t in rows[1:4]])
reader = pypdf.PdfReader(io.BytesIO(data))
print('content stream:', reader.pages[0].get_contents().get_data())
c = ('BT /F1 12 Tf 200 700 Td (world) Tj ET '
     'BT /F1 12 Tf 72 700 Td (Hello) Tj ET')  # world first
page = pypdf.PdfReader(io.BytesIO(make_pdf(c))).pages[0]
print(repr(page.extract_text()),
      repr(page.extract_text(extraction_mode='layout').strip()))
s = ('BT /F1 18 Tf 72 720 Td (Salary: 120000 USD) Tj ET '
     '0 0 0 rg 60 710 300 30 re f')  # then a black box
secret = pypdf.PdfReader(io.BytesIO(make_pdf(s))).pages[0]
print('text under the black box:', secret.extract_text())
'''

EXTRACT = _HEAD + '''# Text extraction with pypdf and with pdfminer.six: pages, layout and a table-like page.
from pdfminer.high_level import extract_text
six = make_pdf(*['Page %d' % i for i in range(1, 7)])
reader = pypdf.PdfReader(io.BytesIO(six))
text = ''
for page in reader.pages:
    text += page.extract_text()
print('pypdf:', repr(text))
miner = extract_text(io.BytesIO(six))
print('pdfminer.six:', repr(miner[:16]), miner.count('\\x0c'))
c = ('BT /F1 12 Tf 72 700 Td (Name) Tj 200 0 Td (Price) Tj ET '
     'BT /F1 12 Tf 72 686 Td (Tea) Tj 200 0 Td (3) Tj ET')
page = pypdf.PdfReader(io.BytesIO(make_pdf(c))).pages[0]
print(repr(page.extract_text()))
print(page.extract_text(extraction_mode='layout').strip())
'''

SHIPPED_PDF = _HEAD + _REAL + '''# The sample chapter the book ships: the pages, the text of the first page and every image the chapter's loop writes.
name = real_file('Recursion_Chapter1.pdf')  # sha256 checked
reader = pypdf.PdfReader(name)
print(len(reader.pages), 'pages; encrypted', reader.is_encrypted)
print(repr(reader.pages[0].extract_text()[:36]))
print('images per page:', *[len(p.images) for p in reader.pages])
n = 0
for number, page in enumerate(reader.pages, 1):
    for image in page.images:
        n += 1
        with open('%d_%d_%s' % (n, number, image.name), 'wb') as f:
            f.write(image.data)
print(n, 'image files written')
'''

PAGES = _HEAD + '''# Copying pages into a new file: a tuple is a range, a list is exact, merge() inserts at a position.
six = make_pdf(*['Page %d' % i for i in range(1, 7)])
def texts(writer):                                   # the page texts of a writer
    buf = io.BytesIO(); writer.write(buf)
    return [p.extract_text() for p in pypdf.PdfReader(buf).pages]
for label, pages in (('(0, 5)', (0, 5)), ('(0, 5, 2)', (0, 5, 2)),
                     ('[0, 5]', [0, 5])):
    w = pypdf.PdfWriter(); w.append(io.BytesIO(six), pages)
    print('append', label, '->', texts(w))
three = make_pdf('A1', 'A2', 'A3')
w = pypdf.PdfWriter(); w.append(io.BytesIO(three))
w.merge(1, io.BytesIO(six), pages=(0, 2))
print('merge(1, file, pages=(0, 2)) ->', texts(w))
w = pypdf.PdfWriter(); w.append(io.BytesIO(three))
w.merge(1, io.BytesIO(six), (0, 2))                  # as the chapter writes it
print('merge(1, file, (0, 2)) ->', len(w.pages), 'pages; bookmark titles',
      [o['/Title'] for o in w.outline])
'''

ROTATE = _HEAD + '''# Rotating, inserting blank pages and laying one page over another.
def one(text): return pypdf.PdfReader(io.BytesIO(make_pdf(text))).pages[0]
w = pypdf.PdfWriter(); w.append(io.BytesIO(make_pdf('x')))
page, seen = w.pages[0], []
for angle in (90, 270, -90, 180):
    page.rotate(angle); seen.append(page.rotation)
print('rotation after 90, 270, -90, 180:', seen)
w = pypdf.PdfWriter()
for size in ((300, 400), (100, 100)): w.add_blank_page(*size)
w.insert_blank_page(index=1)
print([(float(p.mediabox.width), float(p.mediabox.height))
       for p in w.pages])
print('insert_blank_page(2) is', float(w.insert_blank_page(2).mediabox.width),
      'points wide')
def drawn(over):                       # in stream order
    base = one('BASE'); base.merge_page(one('MARK'), over=over)
    return re.findall(rb'\\((\\w+)\\) Tj', base.get_contents().get_data())
print('merge_page(over=False) draws', drawn(False))
print('merge_page(over=True) draws', drawn(True))
'''

DECRYPT = _HEAD + '''# AES-256 encryption, a reader that refuses until decrypt() matches, and what decrypt() returns.
w = pypdf.PdfWriter(); w.add_blank_page(100, 100)
w.encrypt('userpw', 'ownerpw', algorithm='AES-256')
b = io.BytesIO(); w.write(b)
def reader():
    return pypdf.PdfReader(io.BytesIO(b.getvalue()))
try:
    reader().pages[0]
except Exception as e:
    print(type(e).__name__, '-', e)
for pw in ('wrong', 'Userpw', 'userpw', 'ownerpw'):
    print(repr(pw), '->', reader().decrypt(pw).name)
r = reader(); r.decrypt('userpw')
print(len(r.pages), 'page after decrypt')
'''

COMBINE = _HEAD + '''# Project 12: combine every page after the first of all the PDFs of a folder (the chapter's steps, pdf_files corrected).
for name in ('b.pdf', 'A.pdf', 'c.PDF'):
    save(name, make_pdf(name + ' p1', name + ' p2'))
def find_pdfs():                  # step 1
    pdf_filenames = []
    for filename in os.listdir('.'):
        if filename.lower().endswith('.pdf'):  # c.PDF
            pdf_filenames.append(filename)
    pdf_filenames.sort(key=str.lower)
    return pdf_filenames
def combine():                    # steps 2 and 3
    writer = pypdf.PdfWriter()
    for pdf_filename in find_pdfs():
        reader = pypdf.PdfReader(pdf_filename)
        writer.append(pdf_filename, (1, len(reader.pages)))
    with open('combined.pdf', 'wb') as file:
        writer.write(file)
    return len(pypdf.PdfReader('combined.pdf').pages)
print(find_pdfs(), combine(), 'pages')
print(find_pdfs(), combine(), 'pages')  # again
'''

PARANOIA = _HEAD + '''# PDF Paranoia: encrypt every PDF of a folder tree, test the copy, delete the original; then decrypt.
for rel in ('a.pdf', 'sub/b.pdf', 'sub/deeper/C.PDF', 'notes.txt'):
    path = os.path.join('tree', rel); os.makedirs(os.path.dirname(path), exist_ok=True)
    save(path, make_pdf(rel) if rel.lower().endswith('.pdf') else b'plain text')
def pdfs(folder, suffix):
    for dirpath, dirnames, filenames in os.walk(folder):
        for filename in sorted(filenames):
            if filename.lower().endswith(suffix):
                yield os.path.join(dirpath, filename)
def encrypt_tree(folder, password):
    for path in list(pdfs(folder, '.pdf')):
        if '_encrypted' in path: continue
        reader = pypdf.PdfReader(path)
        writer = pypdf.PdfWriter(); writer.append(reader)
        writer.encrypt(password, algorithm='AES-256')
        new = path[:-4] + '_encrypted' + path[-4:]
        with open(new, 'wb') as f: writer.write(f)
        check = pypdf.PdfReader(new)    # test first
        if (check.decrypt(password).name != 'NOT_DECRYPTED'
                and len(check.pages) == len(reader.pages)):
            os.remove(path); print('Removed', path)
def decrypt_tree(folder, password):
    for path in pdfs(folder, '_encrypted.pdf'):
        reader = pypdf.PdfReader(path)
        if reader.decrypt(password).name == 'NOT_DECRYPTED':
            print('Wrong password for', path); continue
        writer = pypdf.PdfWriter(); writer.append(reader)
        out = path.replace('_encrypted', '_decrypted')
        with open(out, 'wb') as f: writer.write(f)
        print('Decrypted copy:', out)
encrypt_tree('tree', 'swordfish')
decrypt_tree('tree', 'wrong'); decrypt_tree('tree', 'swordfish')
'''

BREAKER = _HEAD + '''# PDF Password Breaker on a small dictionary: both cases of every word, until decrypt() succeeds.
import time
words = ['apple', 'banana', 'cherry', 'dragonfruit']
w = pypdf.PdfWriter(); w.append(io.BytesIO(make_pdf('Locked')))
w.encrypt('CHERRY', algorithm='AES-256')
b = io.BytesIO(); w.write(b); data = b.getvalue()
reader = pypdf.PdfReader(io.BytesIO(data))
attempts, found = 0, None
for word in words:
    for guess in (word, word.upper()):
        attempts += 1
        if reader.decrypt(guess).name != 'NOT_DECRYPTED':
            found = guess; break
    if found: break
print('found', found, 'after', attempts, 'attempts')
print('a dictionary of n words costs up to', 2 * len(words), 'attempts')
print(repr(reader.pages[0].extract_text()))
'''

DOCXREAD = _HEAD + '''# Reading a Word document: paragraphs, runs and what paragraphs leave out.
import docx
from docx.oxml import parse_xml
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
LINK = ('<w:hyperlink xmlns:w="%s"><w:r><w:t>linked</w:t></w:r>'
        '</w:hyperlink>' % W)
doc = docx.Document(); p = doc.add_paragraph('Some ')
p.add_run('bold').bold = True
p.add_run(' text')
t = doc.add_table(rows=1, cols=1); t.cell(0, 0).text = 'in a table'
doc.sections[0].header.paragraphs[0].text = 'IN THE HEADER'
q = doc.add_paragraph('before '); q._p.append(parse_xml(LINK))
print([x.text for x in doc.paragraphs])
print([(r.text, r.bold) for r in p.runs])
print('table:', doc.tables[0].cell(0, 0).text)
print('header:', doc.sections[0].header.paragraphs[0].text)
print('hyperlink paragraph: text', repr(q.text),
      'runs', [r.text for r in q.runs])
'''

SHIPPED_DOCX = _HEAD + _REAL + '''# The Word file the book ships: its paragraphs, its runs and the restyling of the chapter.
import docx
name = real_file('demo.docx')  # sha256 checked
doc = docx.Document(name)
for p in doc.paragraphs:
    print(p.style.name, '|', p.text)
for r in doc.paragraphs[1].runs:
    print('   ', repr(r.text), r.bold, r.italic)
doc.paragraphs[0].style = 'Normal'
doc.paragraphs[1].runs[0].style = 'Quote Char'
doc.paragraphs[1].runs[1].underline = True
doc.save('restyled.docx')
d = docx.Document('restyled.docx')
print(d.paragraphs[0].style.name)
for r in d.paragraphs[1].runs:
    print('   ', repr(r.text), r.style.name, r.underline)
'''

STYLES = _HEAD + '''# Styles: a paragraph takes a paragraph style, a run a character style, and a name must match exactly.
import docx
doc = docx.Document()
p = doc.add_paragraph('A plain paragraph')
r = p.add_run(' with a run')
for obj, name in ((p, 'Quote Char'), (r, 'Quote')):
    try:
        obj.style = name
    except ValueError as e:
        print(name, '->'); print('   ', e)
for n in ('List Number ', 'List Number'):
    try:
        doc.add_paragraph('x', n); print(repr(n), 'ok')
    except KeyError as e:
        print(repr(n), 'KeyError', e)
print(len(doc.styles), 'styles in a new document')
'''

WRITE_A = _HEAD + '''# Writing: add_paragraph() returns a Paragraph, add_run() a Run, and bold has three states.
import docx
doc = docx.Document()
p = doc.add_paragraph('Hello world!')
r = p.add_run(' A second run.')
print(type(p).__name__, type(r).__name__, len(p.runs))
tag = lambda: re.findall(r'<w:b[ /][^>]*>', r._r.xml)
for value in (True, False, None):   # three states
    r.bold = value; print(value, '->', r.bold, tag())
r.strike = True; print('run.strike =', r.font.strike)
r.font.strike = True; print('font.strike =', r.font.strike)
'''

WRITE_B = _HEAD + '''# Headings of levels 0 to 9, a page break and two pictures.
import docx
from PIL import Image
from docx.shared import Inches
from docx.enum.text import WD_BREAK
doc = docx.Document()
for level in (0, 1, 9, 10):
    try:
        print(level, doc.add_heading('Header %d' % level, level).style.name)
    except ValueError as e:
        print(level, 'ValueError:', e)
doc.add_paragraph('Hello').add_run().add_break(WD_BREAK.PAGE)
print(re.findall(r'<w:br[^>]*>', doc.element.xml))
img = io.BytesIO(); Image.new('RGB', (60, 80), 'red').save(img, 'PNG')
for label, kw in (('no width', {}), ('width=1 inch', {'width': Inches(1)})):
    shape = doc.add_picture(io.BytesIO(img.getvalue()), **kw)
    print(label, '->', int(shape.width), int(shape.height))
'''

INVITE = _HEAD + '''# Custom Invitations: styles made in the program, one invitation per page.
import docx
from docx.shared import Pt
from docx.enum.style import WD_STYLE_TYPE
guests = ['Prof. Plum', 'Miss Scarlet', 'Col. Mustard', 'Al Sweigart', 'Robocop']
doc = docx.Document()
for name, size, bold in (('Invite Lead', 20, False), ('Invite Name', 32, True),
                         ('Invite Detail', 14, False)):
    s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = doc.styles['Normal']; s.font.size = Pt(size); s.font.bold = bold
card = [('It would be a pleasure to have the company of', 'Invite Lead'),
        (None, 'Invite Name'),
        ('at 11010 Memory Lane on the Evening of', 'Invite Lead'),
        ('April 1st', 'Invite Detail'), ("at 7 o'clock", 'Invite Detail')]
for i, guest in enumerate(guests):
    for text, style in card:
        doc.add_paragraph(text or guest, style)
    if i < len(guests) - 1: doc.add_page_break()
doc.save('invitations.docx')
d = docx.Document('invitations.docx')
pages = d.element.xml.count('w:type="page"')
print(len(d.paragraphs), 'paragraphs;', pages, 'page breaks for', len(guests))
print([p.text for p in d.paragraphs if p.style.name == 'Invite Name'])
'''

PROGRAMS = dict(B.PROGRAMS)
for _n, _c, _live in (("structureCut", STRUCT, False), ("extractTextCut", EXTRACT, False), ("shippedChapterPdfCut", SHIPPED_PDF, True),
                      ("pageSelectionCut", PAGES, False), ("rotateBlankWatermarkCut", ROTATE, False), ("decryptCut", DECRYPT, False),
                      ("combinePdfsCut", COMBINE, False), ("pdfParanoiaCut", PARANOIA, False), ("passwordBreakerCut", BREAKER, False),
                      ("docxReadCut", DOCXREAD, False), ("shippedDemoDocxCut", SHIPPED_DOCX, True), ("stylesCut", STYLES, False),
                      ("writeRunsCut", WRITE_A, False), ("breaksCut", WRITE_B, False), ("customInvitationsCut", INVITE, False)):
    PROGRAMS[_n + "_v1_0_0.py"] = (_c, (), _live)

# copy -> original: every printed line of the copy must be a line the original printed, in the same order (SUBSET); a copy that prints its
# own lines (another way of printing, shorter lines) is not listed and is checked only by being run again. EQUIV (a copy whose transcript
# must equal the original's) is empty for this chapter
SUBSET = {"structureCut_v1_0_0.py": "pdfStructure_v1_0_0.py", "pageSelectionCut_v1_0_0.py": "pageSelection_v1_0_0.py"}
EQUIV = {}

KIND = {n: "none" for n in PROGRAMS}
# the words a slide must carry when it shows a program that reads a file the book ships
LIVEKEY = {"shippedChapterPdf_v1_0_0.py": ("Recursion_Chapter1.pdf", "sha256 d943896e0bd2"),
           "shippedChapterPdfCut_v1_0_0.py": ("Recursion_Chapter1.pdf", "sha256 d943896e0bd2"),
           "shippedDemoDocx_v1_0_0.py": ("demo.docx", "sha256 b43b21137d7d"),
           "shippedDemoDocxCut_v1_0_0.py": ("demo.docx", "sha256 b43b21137d7d")}
for _n in LIVEKEY:
    KIND[_n] = "live"

# ---- charts and tables: each a setup and one expression --------------------------------------------------------------------------
_ROT = ("import io, pypdf\n"
        "def rotations():\n"
        "    w = pypdf.PdfWriter(); w.add_blank_page(100, 100); page = w.pages[0]; out = []\n"
        "    for a in (90, 270, -90, 180):\n"
        "        page.rotate(a); out.append(['rotate(%d)' % a, page.rotation])\n"
        "    return out\n")
CHARTS = {"rotationSum": {"setup": _ROT, "expr": "rotations()"}}
_ENC = ("import io, pypdf\n"
        "def row(label, **kw):\n"
        "    w = pypdf.PdfWriter(); w.add_blank_page(100, 100); w.encrypt('swordfish', **kw)\n"
        "    b = io.BytesIO(); w.write(b); e = pypdf.PdfReader(b).trailer['/Encrypt']\n"
        "    return [label, str(e['/V']), str(e['/R']), str(e['/Length'])]\n")
TABLES = {"encryptAlgorithms": {
    "setup": _ENC,
    "expr": "[['encrypt() call', 'V', 'R', 'Length']] + [row(\"encrypt('swordfish')\"), "
            "row(\"algorithm='AES-128'\", algorithm='AES-128'), row(\"algorithm='AES-256'\", algorithm='AES-256')]"}}

# ---- runs under other environments of the research (recorded, then used in a card) ---------------------------------------------------
BARE_MAKE = ("import io, sys\nfrom PIL import Image\nbuf = io.BytesIO(); Image.new('RGB', (40, 40), 'white').save(buf, 'PDF')\n"
             "open(sys.argv[1], 'wb').write(buf.getvalue())\n")
BARE_CODE = ("import io, sys, pypdf\n"
             "page = pypdf.PdfReader(sys.argv[1]).pages[0]\n"
             "try:\n"
             "    print(page.images[0].name)\n"
             "except Exception as e:\n"
             "    print(type(e).__name__, '-', str(e).split('.')[0])\n"
             "w = pypdf.PdfWriter(); w.add_blank_page(100, 100)\n"
             "try:\n"
             "    w.encrypt('swordfish', algorithm='AES-256'); print('AES-256 written')\n"
             "except Exception as e:\n"
             "    print(type(e).__name__, '-', e)\n"
             "w2 = pypdf.PdfWriter(); w2.add_blank_page(100, 100); w2.encrypt('swordfish'); print('RC4 written')\n")
DOCX02_CODE = "import docx"
