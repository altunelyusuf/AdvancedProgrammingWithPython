#!/usr/bin/env python3
"""SEN0414 chapter 14 (Excel Spreadsheets, 3rd edition) - the shared base of the corpus, version 1.0.0.

Holds the chapter's whole programs as text, the single source the deck and the page's Lab tab both read, in the shape of
sen0414_ch13_corpus_base_v1_0_0.py. Each program is a complete script that exits 0 and prints what it printed on 2026-10-12 under
Python 3.14.6 (uv) with openpyxl 3.1.5, except where a program is marked LIVE.

The chapter's files are not in the repository, so every program starts with the same few lines (PRE) that make stand-ins of
them in a temporary folder and change into it: example3.xlsx (the seven rows of Table 14-1 on Sheet1, with Sheet2 and Sheet3),
censuspopdata.xlsx (a short sheet with the columns of the real one: tract, state, county, population) and produceSales3.xlsx
(ten rows with the TOTAL formula =ROUND(B2*C2,2)). The chapter's code is kept as printed apart from those lines and from the
changes a comment names. The two projects were also run, whole, on the real files of the book's online materials; that
record is in ch14-page/live_observations_v1_0_0.json.

* PROJECT9 is Project 9 as the chapter's three steps give it, assembled (Step 1's markers removed, Step 2's replaced by one
  space), run on the stand-in census sheet, followed by the shell session that imports census2010.
* PROJECT10 is Project 10 whole on the stand-in produce sheet, with a look at the rows before and after.
* MULTIPLICATION and BLANKROWS are the two practice programs, with sys.argv set in the first line instead of the command line.
* CELLS is the reading session of the chapter: workbook, sheet names, cells, column letters, slices.
* WRITING is the writing session: Workbook(), title, save, create_sheet, del, values.
* FORMULAS writes the chapter's formula, reads it back with and without data_only, and shows what a stored result is.
* LAYOUT sets fonts, row heights, column widths, merges and freezes panes, and reads back what the file records.
* BARCHART is the chapter's bar chart, the parts of the zip it becomes, and a load and a save.
* LIVE_EXAMPLE downloads example3.xlsx from automatetheboringstuff.com and reads Table 14-1. LIVE; needs requests.
* RECALC runs LibreOffice headless on the produce stand-in and compares its totals with Python's round(). Needs soffice.
* SAFELOAD builds a small file with an entity bomb in it and loads it with openpyxl, with defusedxml on the path.

Sources read for this chapter, on 2026-10-12: the book chapter (automatetheboringstuff.com/3e/chapter14.html, sha256 6965d2d5...be5a
as the book's ontology records), the documentation of openpyxl, Python 3.14, LibreOffice, XlsxWriter and Expat, and the Microsoft
pages on Excel limits, ROUND and freeze panes.
"""

PRE = r"""import os, tempfile, datetime, openpyxl
os.chdir(tempfile.mkdtemp())            # stand-ins for the chapter's files live in a temporary folder
def make_stand_ins():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sheet1'
    for r in [(datetime.datetime(2035, 4, 5, 13, 34, 2), 'Apples', 73), (datetime.datetime(2035, 4, 5, 3, 41, 23), 'Cherries', 85), (datetime.datetime(2035, 4, 6, 12, 46, 51), 'Pears', 14), (datetime.datetime(2035, 4, 8, 8, 59, 43), 'Oranges', 52), (datetime.datetime(2035, 4, 10, 2, 7, 0), 'Apples', 152), (datetime.datetime(2035, 4, 10, 18, 10, 37), 'Bananas', 23), (datetime.datetime(2035, 4, 10, 2, 40, 46), 'Strawberries', 98)]: ws.append(r)
    wb.create_sheet('Sheet2'); wb.create_sheet('Sheet3'); wb.save('example3.xlsx')
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Population by Census Tract'; ws.append(['CensusTract', 'State', 'County', 'POP2010'])
    for t in [('02020000100', 'AK', 'Anchorage', 4115), ('02020000200', 'AK', 'Anchorage', 3040), ('02020000300', 'AK', 'Anchorage', 5360), ('02013000100', 'AK', 'Aleutians East', 3141), ('01001020100', 'AL', 'Autauga', 1912), ('01001020200', 'AL', 'Autauga', 2170), ('01003010100', 'AL', 'Baldwin', 3000), ('01003010200', 'AL', 'Baldwin', 4100), ('06075010100', 'CA', 'San Francisco', 4000)]: ws.append(list(t))
    wb.save('censuspopdata.xlsx')
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sheet'; ws.append(['PRODUCE', 'COST PER POUND', 'POUNDS SOLD', 'TOTAL'])
    for i, (p, c, q) in enumerate([('Potatoes', 0.86, 21.6), ('Garlic', 1.19, 10.0), ('Celery', 3.07, 4.5), ('Lemon', 1.29, 7.0), ('Okra', 2.26, 38.6), ('Garlic', 1.19, 2.5), ('Celery', 3.07, 0.5), ('Parsnips', 2.27, 14.5), ('Lemon', 1.29, 3.0), ('Garlic', 1.19, 6.0)], 2): ws.append([p, c, q, '=ROUND(B%d*C%d,2)' % (i, i)])
    wb.save('produceSales3.xlsx')
make_stand_ins()
"""

PROJECT9 = PRE + r"""# readCensusExcel.py - Project 9, Steps 1 to 3 assembled
import openpyxl, pprint
print('Opening workbook...')
wb = openpyxl.load_workbook('censuspopdata.xlsx')
sheet = wb['Population by Census Tract']
county_data = {}
print('Reading rows...')
for row in range(2, sheet.max_row + 1):
    state  = sheet['B' + str(row)].value
    county = sheet['C' + str(row)].value
    pop    = sheet['D' + str(row)].value
    county_data.setdefault(state, {})
    county_data[state].setdefault(county, {'tracts': 0, 'pop': 0})
    county_data[state][county]['tracts'] += 1
    county_data[state][county]['pop'] += int(pop)
print('Writing results...')
result_file = open('census2010.py', 'w')
result_file.write('allData = ' + pprint.pformat(county_data))
result_file.close()
print('Done.')
# the shell session of the chapter: the current folder is the script's folder, so the import finds the file
import census2010
print(census2010.allData['AK']['Anchorage'])
anchorage_pop = census2010.allData['AK']['Anchorage']['pop']
print('The 2010 population of Anchorage was ' + str(anchorage_pop))
print(open('census2010.py').read())
"""

PROJECT10 = PRE + r"""# updateProduce.py - Project 10, Steps 1 and 2 assembled
import openpyxl
wb = openpyxl.load_workbook('produceSales3.xlsx')
sheet = wb['Sheet']
# The produce types and their updated prices:
PRICE_UPDATES = {'Garlic': 3.07,
                 'Celery': 1.19,
                 'Lemon': 1.27}
# Loop through the rows and update the prices.
for row_num in range(2, sheet.max_row + 1):  # Skip the first row.
    produce_name = sheet.cell(row=row_num, column=1).value
    if produce_name in PRICE_UPDATES:
        sheet.cell(row=row_num, column=2).value = PRICE_UPDATES[produce_name]
wb.save('updatedProduceSales3.xlsx')
old = openpyxl.load_workbook('produceSales3.xlsx')['Sheet']
new = openpyxl.load_workbook('updatedProduceSales3.xlsx')['Sheet']
for r in range(1, old.max_row + 1):
    print(old.cell(row=r, column=1).value, old.cell(row=r, column=2).value, '->', new.cell(row=r, column=2).value, new.cell(row=r, column=4).value)
# openpyxl saved the formulas without the results Excel had stored for them:
print(openpyxl.load_workbook('updatedProduceSales3.xlsx', data_only=True)['Sheet']['D2'].value)
"""

MULTIPLICATION = PRE + r"""# multiplicationTable.py - Practice Program: Multiplication Table Maker
import sys, openpyxl
from openpyxl.styles import Font
sys.argv = ['multiplicationTable.py', '6']     # stands for: python multiplicationTable.py 6
n = int(sys.argv[1])
wb = openpyxl.Workbook()
sheet = wb.active
bold = Font(bold=True)
for i in range(1, n + 1):
    sheet.cell(row=1, column=i + 1, value=i).font = bold      # labels along row 1, from B1
    sheet.cell(row=i + 1, column=1, value=i).font = bold      # labels down column A, from A2
    for j in range(1, n + 1):
        sheet.cell(row=i + 1, column=j + 1, value=i * j)
wb.save('multiplicationTable.xlsx')
check = openpyxl.load_workbook('multiplicationTable.xlsx').active
print(check.dimensions, check['D5'].value, check['G7'].value, check['A1'].value)
for row in check.iter_rows(min_row=1, max_row=3, values_only=True): print(row)
"""

BLANKROWS = PRE + r"""# blankRowInserter.py - Practice Program: Blank Row Inserter
import sys, openpyxl
sys.argv = ['blankRowInserter.py', '3', '2', 'myProduce.xlsx']   # stands for: python blankRowInserter.py 3 2 myProduce.xlsx
n, m, filename = int(sys.argv[1]), int(sys.argv[2]), sys.argv[3]
wb = openpyxl.Workbook(); sh = wb.active
for r, v in enumerate(['Potatoes', 'Okra', 'Fava beans', 'Watermelon', 'Garlic', 'Parsnips'], 1): sh.cell(row=r, column=1, value=v)
sh['C6'] = '=SUM(B1:B5)'
wb.save(filename)
wb = openpyxl.load_workbook(filename)
sheet = wb.active
sheet.insert_rows(n, m)                         # m blank rows before row n
wb.save('blankRowsInserted.xlsx')
out = openpyxl.load_workbook('blankRowsInserted.xlsx').active
print([out.cell(row=r, column=1).value for r in range(1, 9)])
print(out['C8'].value, out['C6'].value)         # the formula moved, and its text did not change
"""

CELLS = PRE + r"""import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
wb = openpyxl.load_workbook('example3.xlsx')
print(type(wb))
print(wb.sheetnames)
sheet = wb['Sheet3']
print(sheet, type(sheet), sheet.title)
print(wb.active)
sheet = wb['Sheet1']
print(sheet['A1'], sheet['A1'].value)
c = sheet['B1']
print(c.value, 'Row %s, Column %s is %s' % (c.row, c.column, c.value), 'Cell %s is %s' % (c.coordinate, c.value))
print(sheet.cell(row=1, column=2))
for i in range(1, 8, 2): print(i, sheet.cell(row=i, column=2).value)
print(sheet.max_row, sheet.max_column)
print(get_column_letter(1), get_column_letter(27), get_column_letter(900), column_index_from_string('AA'))
for row_of_cell_objects in sheet['A1':'C3']:
    for cell_obj in row_of_cell_objects: print(cell_obj.coordinate, cell_obj.value)
    print('--- END OF ROW ---')
print(list(sheet.columns)[1][0].value, len(list(sheet.rows)), len(list(sheet.columns)))
"""

WRITING = PRE + r"""import openpyxl
wb = openpyxl.Workbook()
print(wb.sheetnames)
sheet = wb.active
sheet.title = 'Spam Bacon Eggs Sheet'
print(wb.sheetnames)
wb.save('example3_copy.xlsx')
print(openpyxl.load_workbook('example3_copy.xlsx').sheetnames)
wb = openpyxl.Workbook()
print(wb.create_sheet())
print(wb.sheetnames)
wb.create_sheet(index=0, title='First Sheet')
wb.create_sheet(index=2, title='Middle Sheet')
print(wb.sheetnames)
del wb['Middle Sheet']
del wb['Sheet1']
print(wb.sheetnames)
sheet = wb['Sheet']
sheet['A1'] = 'Hello, world!'
print(sheet['A1'].value)
a = wb.create_sheet('Dup'); b = wb.create_sheet('Dup'); c = wb.create_sheet('dup')
print(a.title, b.title, c.title)
"""

FORMULAS = PRE + r"""import openpyxl, io, zipfile
wb = openpyxl.Workbook()
sheet = wb.active
sheet['A1'] = 200
sheet['A2'] = 300
sheet['A3'] = '=SUM(A1:A2)'
wb.save('writeFormula3.xlsx')
print(sheet['A3'].data_type, sheet['A3'].value)
wb = openpyxl.load_workbook('writeFormula3.xlsx')
print(wb.active['A3'].value)
wb_data_only = openpyxl.load_workbook('writeFormula3.xlsx', data_only=True)
print(wb_data_only.active['A3'].value)          # None: nobody has calculated the formula
# What Excel would have stored: put the result into the file as the spreadsheet application does
z = zipfile.ZipFile('writeFormula3.xlsx'); out = io.BytesIO(); zo = zipfile.ZipFile(out, 'w')
for n in z.namelist():
    t = z.read(n)
    zo.writestr(n, t.replace(b'<f>SUM(A1:A2)</f><v />', b'<f>SUM(A1:A2)</f><v>500</v>') if n.startswith('xl/worksheets/') else t)
zo.close(); open('stored500.xlsx', 'wb').write(out.getvalue())
print(openpyxl.load_workbook('stored500.xlsx', data_only=True).active['A3'].value)
# openpyxl loads and saves it again: the formula stays, the stored result is gone
openpyxl.load_workbook('stored500.xlsx').save('resaved.xlsx')
print(openpyxl.load_workbook('resaved.xlsx').active['A3'].value, openpyxl.load_workbook('resaved.xlsx', data_only=True).active['A3'].value)
sheet['A4'] = '=NOSUCHFUNC(1'                   # openpyxl does not check a formula
print(sheet['A4'].data_type, sheet['A4'].value)
"""

LAYOUT = PRE + r"""import openpyxl, io, zipfile, re
from openpyxl.styles import Font
wb = openpyxl.Workbook()
sheet = wb.active
italic_24_font = Font(size=24, italic=True)
sheet['A1'].font = italic_24_font
sheet['A1'] = 'Hello, world!'
bold_font = Font(name='Times New Roman', bold=True)
sheet['B3'].font = bold_font
sheet['B3'] = 'Bold Times New Roman'
print(sheet['A1'].font.name, sheet['A1'].font.sz, sheet['B3'].font.name, sheet['B3'].font.sz, sheet['C5'].font.name, sheet['C5'].font.sz)
try: sheet['A1'].font.bold = False
except AttributeError as e: print('AttributeError', str(e)[:36])
sheet['A5'] = 'Tall row'; sheet['B5'] = 'Wide column'
sheet.row_dimensions[5].height = 70
sheet.column_dimensions['B'].width = 20
sheet.column_dimensions['D'].hidden = True
sheet.merge_cells('F1:I3'); sheet['F1'] = 'Twelve cells merged together.'
sheet.merge_cells('F5:G5'); sheet['F5'] = 'Two merged cells.'
print(sorted(map(str, sheet.merged_cells.ranges)), type(sheet['G2']).__name__)
try: sheet['G2'] = 'x'
except AttributeError as e: print('AttributeError', e)
sheet.freeze_panes = 'A2'
wb.save('layout.xlsx')
xml = zipfile.ZipFile('layout.xlsx').read('xl/worksheets/sheet1.xml').decode()
print(re.findall(r'<pane[^>]*>', xml)); print(re.findall(r'<col [^>]*>', xml)); print(re.findall(r'<row r="5"[^>]*>', xml))
sheet.unmerge_cells('F1:I3')
print(sorted(map(str, sheet.merged_cells.ranges)), type(sheet['G2']).__name__)
for setting in ('A2', 'B1', 'C1', 'C2'):
    ws = openpyxl.Workbook().active; ws.freeze_panes = setting; p = ws.sheet_view.pane; print(setting, p.xSplit, p.ySplit)
ws = openpyxl.Workbook().active
before = (len(ws.row_dimensions), len(ws.column_dimensions))
ws.row_dimensions[3].height; ws.column_dimensions['A'].width       # only reading
print(before, (len(ws.row_dimensions), len(ws.column_dimensions)))
"""

BARCHART = PRE + r"""import openpyxl, io, zipfile
wb = openpyxl.Workbook()
sheet = wb.active
for i in range(1, 11):  # Create some data in column A.
    sheet['A' + str(i)] = i * i
ref_obj = openpyxl.chart.Reference(sheet, 1, 1, 1, 10)       # min_col, min_row, max_col, max_row
print(ref_obj)
series_obj = openpyxl.chart.Series(ref_obj, title='First series')
chart_obj = openpyxl.chart.BarChart()
chart_obj.title = 'My Chart'
chart_obj.append(series_obj)
sheet.add_chart(chart_obj, 'C5')
wb.save('sampleChart3.xlsx')
print([n for n in zipfile.ZipFile('sampleChart3.xlsx').namelist() if 'chart' in n or 'drawing' in n])
again = openpyxl.load_workbook('sampleChart3.xlsx')
print(len(again.active._charts))
again.save('sampleChart3_again.xlsx')
print(len(openpyxl.load_workbook('sampleChart3_again.xlsx').active._charts))
print(openpyxl.chart.Reference(sheet, 2, 1, 2, 10), openpyxl.chart.Reference(sheet, 1, 1, 1))   # column first; a missing max_row gives one cell
"""

LIVE_EXAMPLE = r"""import io, requests, openpyxl
res = requests.get('https://automatetheboringstuff.com/files/example3.xlsx', timeout=60)
print(res.status_code, len(res.content))
wb = openpyxl.load_workbook(io.BytesIO(res.content))
print(wb.sheetnames, wb.active.title, wb.active.max_row, wb.active.max_column)
sheet = wb['Sheet1']
for row in sheet.iter_rows(min_row=1, max_row=3, values_only=True): print(row)
"""

RECALC = PRE + r"""import subprocess, shutil, openpyxl
soffice = shutil.which('soffice')
print('LibreOffice found:', soffice is not None)
if soffice:
    import os
    os.mkdir('out')
    r = subprocess.run([soffice, '--headless', '--convert-to', 'xlsx', '--outdir', 'out', 'produceSales3.xlsx'], capture_output=True, text=True, timeout=180)
    print('exit', r.returncode)
    src = openpyxl.load_workbook('produceSales3.xlsx')['Sheet']
    lo = openpyxl.load_workbook('out/produceSales3.xlsx', data_only=True)['Sheet']
    differ = 0
    for row in range(2, src.max_row + 1):
        b, c, d = src.cell(row, 2).value, src.cell(row, 3).value, lo.cell(row, 4).value
        if round(b * c, 2) != d: differ += 1
    print('rows', src.max_row - 1, 'LibreOffice total differs from round(b * c, 2):', differ)
    print(lo['D3'].value, openpyxl.load_workbook('out/produceSales3.xlsx')['Sheet']['D3'].value)
"""

SAFELOAD = r"""import io, sys, zipfile, openpyxl, openpyxl.xml, pyexpat
def bomb(depth):
    ents = ''.join('<!ENTITY e%d "%s">' % (i, ('&e%d;' % (i - 1)) * 10) for i in range(1, depth + 1))
    b = io.BytesIO(); src = openpyxl.Workbook(); src.active['A1'] = 'seed'; src.save(b)
    zin = zipfile.ZipFile(io.BytesIO(b.getvalue())); out = io.BytesIO(); zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
    for n in zin.namelist():
        data = zin.read(n)
        if n == 'xl/worksheets/sheet1.xml':
            t = data.decode(); i = t.index('<worksheet')
            t = '<?xml version="1.0"?><!DOCTYPE worksheet [<!ENTITY e0 "aaaaaaaaaa">%s]>' % ents + t[i:]
            data = t.replace('>seed<', '>&e%d;<' % depth).encode()
        zout.writestr(n, data)
    zout.close(); return out.getvalue()
print('expat', pyexpat.EXPAT_VERSION, '| defusedxml in use:', openpyxl.xml.DEFUSEDXML)
for depth in (4, 6):
    data = bomb(depth)
    try: print(depth, 'file under 6000 bytes:', len(data) < 6000, '->', len(openpyxl.load_workbook(io.BytesIO(data)).active['A1'].value), 'characters in A1')
    except Exception as e: print(depth, 'file under 6000 bytes:', len(data) < 6000, '->', type(e).__name__, str(e)[:60])
"""


# name: (program text, packages it needs beyond openpyxl, live?)
PROGRAMS = {
    "project9CensusTally_v1_0_0.py": (PROJECT9, (), False),
    "project10UpdateProduce_v1_0_0.py": (PROJECT10, (), False),
    "multiplicationTable_v1_0_0.py": (MULTIPLICATION, (), False),
    "blankRowInserter_v1_0_0.py": (BLANKROWS, (), False),
    "readingCells_v1_0_0.py": (CELLS, (), False),
    "writingWorkbooks_v1_0_0.py": (WRITING, (), False),
    "formulasAndResults_v1_0_0.py": (FORMULAS, (), False),
    "layoutFontMergeFreeze_v1_0_0.py": (LAYOUT, (), False),
    "barChart_v1_0_0.py": (BARCHART, (), False),
    "liveExample3_v1_0_0.py": (LIVE_EXAMPLE, ("requests",), True),
    "recalcWithLibreOffice_v1_0_0.py": (RECALC, ("soffice",), False),
    "safeLoadEntityBomb_v1_0_0.py": (SAFELOAD, ("defusedxml",), False),
}

if __name__ == "__main__":
    # Usage: sen0414_ch14_corpus_base_v1_0_0.py   (run it with the interpreter that has openpyxl; CH14_EXTRA = folders, joined by the
    # path separator, that hold requests and defusedxml, added to PYTHONPATH only for the programs that name them)
    import os, subprocess, sys
    for name, (code, needs, live) in PROGRAMS.items():
        env = dict(os.environ)
        if any(n in ("requests", "defusedxml") for n in needs) and os.environ.get("CH14_EXTRA"): env["PYTHONPATH"] = os.environ["CH14_EXTRA"]
        r = subprocess.run([sys.executable, "-W", "ignore", "-c", code], capture_output=True, text=True, env=env, timeout=300)
        print("=== %s (exit %d)%s" % (name, r.returncode, " LIVE" if live else "")); print(r.stdout.rstrip()); print(r.stderr[-300:].rstrip())
