#!/usr/bin/env python3
"""Chapter 14 (Excel Spreadsheets): runs the chapter's openpyxl programs, whole, under the given interpreter, against workbooks that
the programs build themselves in a temporary folder, and writes or checks their recorded output.
The chapter's own spreadsheets (example3.xlsx, censuspopdata.xlsx, produceSales3.xlsx) belong to the book's online materials; this
file uses small stand-ins with the same sheets, columns and formulas so that the record can be re-checked anywhere, and the real files
are run in sen0414_ch14_live_v1_0_0.py. Every program below is either a listing of the chapter (an interactive session written as a
script, or Project 9 and Project 10 whole) or a short program that tests one statement of the chapter.
openpyxl is installed in a virtual environment of the interpreter (uv venv; uv pip install openpyxl). Four optional packages sit in
extra site directories given on the command line, each used only by the programs that name it, because their presence changes what
openpyxl does: Pillow (pictures), defusedxml (XML attacks), xlsxwriter and xlrd (other libraries). A program that needs one is given its
directory through the environment variable CH14_SITES, a list of name=directory pairs.
Usage: sen0414_ch14_thirdparty_v1_0_0.py write|check <python> [pillow=<dir> defusedxml=<dir> other=<dir>]
Record: 08-tooling/ch14-page/third_party_record_v1_0_0.json"""
__version__ = "1.0.0"
import json, os, subprocess, sys
MODE, PY = sys.argv[1], sys.argv[2]
SITES = dict(a.split("=", 1) for a in sys.argv[3:])
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REC = os.path.join(REPO, "08-tooling", "ch14-page", "third_party_record_v1_0_0.json")
FIX = r'''
import os, tempfile, datetime, openpyxl
_d = tempfile.mkdtemp(); os.chdir(_d)
def _fixtures():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sheet1'
    for r in [(datetime.datetime(2035, 4, 5, 13, 34, 2), 'Apples', 73), (datetime.datetime(2035, 4, 5, 3, 41, 23), 'Cherries', 85), (datetime.datetime(2035, 4, 6, 12, 46, 51), 'Pears', 14), (datetime.datetime(2035, 4, 8, 8, 59, 43), 'Oranges', 52), (datetime.datetime(2035, 4, 10, 2, 7, 0), 'Apples', 152), (datetime.datetime(2035, 4, 10, 18, 10, 37), 'Bananas', 23), (datetime.datetime(2035, 4, 10, 2, 40, 46), 'Strawberries', 98)]: ws.append(r)
    wb.create_sheet('Sheet2'); wb.create_sheet('Sheet3'); wb.save('example3.xlsx')
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Population by Census Tract'; ws.append(['CensusTract', 'State', 'County', 'POP2010'])
    for t in [('02020000100', 'AK', 'Anchorage', 4115), ('02020000200', 'AK', 'Anchorage', 3040), ('02020000300', 'AK', 'Anchorage', 5360), ('02013000100', 'AK', 'Aleutians East', 3141), ('01001020100', 'AL', 'Autauga', 1912), ('01001020200', 'AL', 'Autauga', 2170), ('01003010100', 'AL', 'Baldwin', 3000), ('01003010200', 'AL', 'Baldwin', '4100'), ('06075010100', 'CA', 'San Francisco', 4000)]: ws.append(list(t))
    wb.save('censuspopdata.xlsx')
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sheet'; ws.append(['PRODUCE', 'COST PER POUND', 'POUNDS SOLD', 'TOTAL'])
    for i, (p, c, q) in enumerate([('Potatoes', 0.86, 21.6), ('Garlic', 1.19, 10.0), ('Celery', 3.07, 4.5), ('Lemon', 1.29, 7.0), ('Okra', 2.26, 38.6), ('Garlic', 1.19, 2.5), ('Celery', 3.07, 0.5), ('Parsnips', 2.27, 14.5), ('Lemon', 1.29, 3.0), ('Garlic', 1.19, 6.0)], 2): ws.append([p, c, q, '=ROUND(B%d*C%d,2)' % (i, i)])
    wb.save('produceSales3.xlsx')
_fixtures()
'''
PROGS = [
("versions: the interpreter, openpyxl and what it installs with", "import sys, openpyxl, importlib.metadata as m\nprint(sys.version.split()[0]); print('openpyxl', openpyxl.__version__)\nprint(sorted(r.split()[0] for r in m.requires('openpyxl')))\nprint(m.metadata('openpyxl')['Requires-Python'], m.metadata('openpyxl')['License-Expression'] if m.metadata('openpyxl')['License-Expression'] else m.metadata('openpyxl')['License'])\nimport openpyxl.xml as x; print('defusedxml in use:', x.DEFUSEDXML, '| lxml in use:', x.LXML)"),
("Opening a Workbook and Getting Sheets from the Workbook, as typed in the chapter's shell", FIX + r'''
import openpyxl
wb = openpyxl.load_workbook('example3.xlsx')
print(type(wb))
print(wb.sheetnames)
sheet = wb['Sheet3']
print(sheet); print(type(sheet)); print(sheet.title)
another_sheet = wb.active
print(another_sheet)
'''),
("Getting Cells from the Sheets: A1, B1, cell(), the loop over every other row, max_row and max_column", FIX + r'''
import openpyxl
wb = openpyxl.load_workbook('example3.xlsx'); sheet = wb['Sheet1']
print(sheet['A1']); print(repr(sheet['A1'].value))
c = sheet['B1']; print(repr(c.value)); print(f'Row {c.row}, Column {c.column} is {c.value}'); print(f'Cell {c.coordinate} is {c.value}'); print(sheet['C1'].value)
print(sheet.cell(row=1, column=2)); print(sheet.cell(row=1, column=2).value)
for i in range(1, 8, 2): print(i, sheet.cell(row=i, column=2).value)
print(sheet.max_row, sheet.max_column)
'''),
("Converting Between Column Letters and Numbers", FIX + r'''
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
print([get_column_letter(n) for n in (1, 2, 27, 900)])
wb = openpyxl.load_workbook('example3.xlsx'); sheet = wb['Sheet1']
print(get_column_letter(sheet.max_column)); print(column_index_from_string('A'), column_index_from_string('AA'))
print(get_column_letter(18278), column_index_from_string('ZZZ'), column_index_from_string('XFD'))
for bad in (0, 18279):
    try: get_column_letter(bad)
    except ValueError as e: print(bad, 'ValueError', e)
try: column_index_from_string('A1')
except ValueError as e: print('ValueError', e)
print(column_index_from_string('aa'))
'''),
("Getting Rows and Columns: the slice, the nested loops, the columns list", FIX + r'''
import openpyxl
wb = openpyxl.load_workbook('example3.xlsx'); sheet = wb['Sheet1']
print(sheet['A1':'C3'])
for row_of_cell_objects in sheet['A1':'C3']:
    for cell_obj in row_of_cell_objects:
        print(cell_obj.coordinate, cell_obj.value)
    print('--- END OF ROW ---')
print(list(sheet.columns)[1])
for cell_obj in list(sheet.columns)[1]: print(cell_obj.value)
print(len(list(sheet.rows)), len(list(sheet.columns)), type(sheet.rows).__name__)
try: sheet.columns[1]
except TypeError as e: print('TypeError', e)
'''),
("Project 9, Steps 1 to 3 assembled whole (stand-in census sheet), then the shell session that imports census2010", FIX + r'''
import openpyxl, pprint
print('Opening workbook...')
wb = openpyxl.load_workbook('censuspopdata.xlsx')
sheet = wb['Population by Census Tract']
county_data = {}
print('Reading rows...')
for row in range(2, sheet.max_row + 1):
    state  = sheet['B' + str(row)].value
    county = sheet['C' + str(row)].value
    pop    = sheet['D' + str(row)].value
    county_data.setdefault(state, {})
    county_data[state].setdefault(county, {'tracts': 0, 'pop': 0})
    county_data[state][county]['tracts'] += 1
    county_data[state][county]['pop'] += int(pop)
print('Writing results...')
result_file = open('census2010.py', 'w')
result_file.write('allData = ' + pprint.pformat(county_data))
result_file.close()
print('Done.')
import census2010
print(census2010.allData['AK']['Anchorage'])
anchorage_pop = census2010.allData['AK']['Anchorage']['pop']
print('The 2010 population of Anchorage was ' + str(anchorage_pop))
print(open('census2010.py').read())
'''),
("Project 9 as the book's online materials ship it: wb.get_sheet_by_name and sheet.get_highest_row", FIX + r'''
import warnings, openpyxl
wb = openpyxl.load_workbook('censuspopdata.xlsx')
with warnings.catch_warnings(record=True) as w:
    warnings.simplefilter('always'); sheet = wb.get_sheet_by_name('Population by Census Tract')
print(sheet, w[0].category.__name__, str(w[0].message))
try: sheet.get_highest_row()
except AttributeError as e: print('AttributeError', e)
print([n for n in dir(sheet) if 'highest' in n], sheet.max_row)
'''),
("Writing: Workbook(), the sheet title, save() to a copy, and the sheets after a reload", FIX + r'''
import openpyxl
wb = openpyxl.Workbook()
print(wb.sheetnames); sheet = wb.active; print(sheet.title)
sheet.title = 'Spam Bacon Eggs Sheet'; print(wb.sheetnames)
wb = openpyxl.load_workbook('example3.xlsx'); sheet = wb['Sheet1']; sheet.title = 'Spam Spam Spam'
wb.save('example3_copy.xlsx')
print(openpyxl.load_workbook('example3_copy.xlsx').sheetnames, openpyxl.load_workbook('example3.xlsx').sheetnames)
'''),
("Creating and Removing Sheets", FIX + r'''
import openpyxl
wb = openpyxl.Workbook(); print(wb.sheetnames)
print(wb.create_sheet()); print(wb.sheetnames)
print(wb.create_sheet(index=0, title='First Sheet')); print(wb.sheetnames)
print(wb.create_sheet(index=2, title='Middle Sheet')); print(wb.sheetnames)
del wb['Middle Sheet']; del wb['Sheet1']; print(wb.sheetnames)
print(wb.create_sheet('Last Before', -1).title, wb.sheetnames)
a = wb.create_sheet('Dup'); b = wb.create_sheet('Dup'); c = wb.create_sheet('dup'); print(a.title, b.title, c.title)
for bad in ('a/b', 'a[b', '*', '?'):
    try: wb.create_sheet(bad)
    except ValueError as e: print(repr(bad), 'ValueError', e)
import warnings
with warnings.catch_warnings(record=True) as w:
    warnings.simplefilter('always'); wb.create_sheet('x' * 32)
print(w[0].category.__name__, str(w[0].message))
try: wb['Nope']
except KeyError as e: print('KeyError', e)
'''),
("Writing Values to Cells, and the type each value gets", FIX + r'''
import openpyxl, datetime
wb = openpyxl.Workbook(); sheet = wb['Sheet']
sheet['A1'] = 'Hello, world!'; print(sheet['A1'].value)
for k, v in enumerate([1, 1.5, 'text', True, datetime.datetime(2035, 4, 5, 13, 34, 2), datetime.date(2035, 4, 5), '=1+1', '123'], 2):
    sheet.cell(row=k, column=1).value = v
    c = sheet.cell(row=k, column=1); print(repr(v), '->', c.data_type, c.number_format)
for bad in ([1, 2], {'a': 1}, 1 + 2j):
    try: sheet['B1'] = bad
    except ValueError as e: print('ValueError', e)
wb.save('types.xlsx'); back = openpyxl.load_workbook('types.xlsx')['Sheet']
print([type(c.value).__name__ for c in back['A']])
'''),
("Project 10 whole: updateProduce.py on the stand-in produce sheet, with before and after", FIX + r'''
import openpyxl
wb = openpyxl.load_workbook('produceSales3.xlsx')
sheet = wb['Sheet']
PRICE_UPDATES = {'Garlic': 3.07,
                 'Celery': 1.19,
                 'Lemon': 1.27}
for row_num in range(2, sheet.max_row + 1):  # Skip the first row.
    produce_name = sheet.cell(row=row_num, column=1).value
    if produce_name in PRICE_UPDATES:
        sheet.cell(row=row_num, column=2).value = PRICE_UPDATES[produce_name]
wb.save('updatedProduceSales3.xlsx')
old = openpyxl.load_workbook('produceSales3.xlsx')['Sheet']; new = openpyxl.load_workbook('updatedProduceSales3.xlsx')['Sheet']
for r in range(1, old.max_row + 1): print([c.value for c in old[r]][:3], '->', new.cell(row=r, column=2).value, new.cell(row=r, column=4).value)
'''),
("Project 10 as the book's online materials ship it: get_highest_row() restored as max_row shows the loop stops one row early", FIX + r'''
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
Worksheet.get_highest_row = lambda self: self.max_row
wb = openpyxl.load_workbook('produceSales3.xlsx')
sheet = wb.get_sheet_by_name('Sheet')
PRICE_UPDATES = {'Garlic': 3.07, 'Celery': 1.19, 'Lemon': 1.27}
for row_num in range(2, sheet.get_highest_row()):
    produce_name = sheet.cell(row=row_num, column=1).value
    if produce_name in PRICE_UPDATES:
        sheet.cell(row=row_num, column=2).value = PRICE_UPDATES[produce_name]
print(sheet.max_row, [(sheet.cell(row=r, column=1).value, sheet.cell(row=r, column=2).value) for r in (sheet.max_row - 1, sheet.max_row)])
'''),
("The cached results of formulas: what a loaded and re-saved workbook keeps", FIX + r'''
import openpyxl
f = openpyxl.Workbook(); s = f.active; s['A1'] = 200; s['A2'] = 300; s['A3'] = '=SUM(A1:A2)'; f.save('writeFormula3.xlsx')
wb = openpyxl.load_workbook('writeFormula3.xlsx'); print(repr(wb.active['A3'].value), wb.active['A3'].data_type)
wb = openpyxl.load_workbook('writeFormula3.xlsx', data_only=True); print(repr(wb.active['A3'].value))
p = openpyxl.load_workbook('produceSales3.xlsx', data_only=True)['Sheet']; print('stand-in produce sheet, saved by openpyxl:', p['D2'].value)
'''),
("Setting the Font Style of Cells: both interactive examples, and what the file records", FIX + r'''
import openpyxl, zipfile, re
from openpyxl.styles import Font
wb = openpyxl.Workbook(); sheet = wb['Sheet']
italic_24_font = Font(size=24, italic=True)
sheet['A1'].font = italic_24_font
sheet['A1'] = 'Hello, world!'
print(sheet['A1'].font.size, sheet['A1'].font.i, sheet['A1'].value)
wb = openpyxl.Workbook(); sheet = wb['Sheet']
bold_font = Font(name='Times New Roman', bold=True)
sheet['A1'].font = bold_font
sheet['A1'] = 'Bold Times New Roman'
italic_font = Font(size=24, italic=True)
sheet['B3'].font = italic_font
sheet['B3'] = '24 pt Italic'
wb.save('styles3.xlsx')
back = openpyxl.load_workbook('styles3.xlsx')['Sheet']
for ref in ('A1', 'B3', 'C5'): f = back[ref].font; print(ref, f.name, f.sz, f.b, f.i)
print(re.findall(r'<font>.*?</font>', zipfile.ZipFile('styles3.xlsx').read('xl/styles.xml').decode()))
try: back['A1'].font.bold = False
except AttributeError as e: print('AttributeError', e)
for kw in ({'size': 'big'}, {'bold': 'yes'}):
    try: f = Font(**kw); print(kw, 'accepted', f.sz, f.b)
    except TypeError as e: print(kw, 'TypeError', e)
'''),
("Adjusting Rows and Columns: height, width, hiding, and what reading a dimension does to the file", FIX + r'''
import openpyxl, zipfile, re
wb = openpyxl.Workbook(); sheet = wb['Sheet']
sheet['A1'] = 'Tall row'; sheet['B2'] = 'Wide column'
sheet.row_dimensions[1].height = 70
sheet.column_dimensions['B'].width = 20
print(type(sheet.row_dimensions).__name__, type(sheet.row_dimensions[1]).__name__, type(sheet.column_dimensions['B']).__name__)
sheet.column_dimensions['C'].hidden = True; sheet.row_dimensions[5].height = 100
wb.save('dimensions3.xlsx')
x = zipfile.ZipFile('dimensions3.xlsx').read('xl/worksheets/sheet1.xml').decode()
print(re.findall(r'<cols>.*?</cols>', x), re.findall(r'<row [^>]*>', x))
fresh = openpyxl.Workbook().active
print(len(fresh.column_dimensions), fresh.column_dimensions['A'].width, len(fresh.column_dimensions), fresh.row_dimensions[3].height, len(fresh.row_dimensions))
fresh['A1'] = 'x'; fresh.parent.save('touched.xlsx'); y = zipfile.ZipFile('touched.xlsx').read('xl/worksheets/sheet1.xml').decode()
print(re.findall(r'<cols>.*?</cols>', y), re.findall(r'<row [^>]*>', y))
'''),
("Merging and Unmerging Cells", FIX + r'''
import openpyxl
wb = openpyxl.Workbook(); sheet = wb['Sheet']
sheet.merge_cells('A1:D3'); sheet['A1'] = 'Twelve cells merged together.'
sheet.merge_cells('C5:D5'); sheet['C5'] = 'Two merged cells.'
wb.save('merged3.xlsx')
print(sorted(map(str, sheet.merged_cells.ranges)), type(sheet['B2']).__name__, type(sheet['A1']).__name__)
try: sheet['B2'] = 'x'
except AttributeError as e: print('AttributeError', e)
wb = openpyxl.load_workbook('merged3.xlsx'); sheet = wb['Sheet']
sheet.unmerge_cells('A1:D3'); sheet.unmerge_cells('C5:D5'); wb.save('unmerged3.xlsx')
print(openpyxl.load_workbook('unmerged3.xlsx')['Sheet'].merged_cells.ranges, type(openpyxl.load_workbook('unmerged3.xlsx')['Sheet']['B2']).__name__)
w = openpyxl.Workbook(); s = w.active; s['A1'] = 'keep'; s['B1'] = 'lost'; s['A2'] = 'lost too'; s.merge_cells('A1:B2'); print(s['A1'].value, s['B1'].value, s['A2'].value)
'''),
("Freezing Panes: Table 14-3 as the file records it, and the same file's saved scroll position", FIX + r'''
import openpyxl, zipfile, re
wb = openpyxl.load_workbook('produceSales3.xlsx'); sheet = wb.active
sheet.freeze_panes = 'A2'; wb.save('freezeExample3.xlsx')
print(re.findall(r'<pane [^>]*>', zipfile.ZipFile('freezeExample3.xlsx').read('xl/worksheets/sheet1.xml').decode()))
for setting in ('A2', 'B1', 'C1', 'C2', 'A1', None):
    s = openpyxl.Workbook().active; s.freeze_panes = setting; p = s.sheet_view.pane
    print(setting, '->', s.freeze_panes, (p.xSplit, p.ySplit) if p else 'no pane')
s = openpyxl.Workbook().active; s.freeze_panes = 'C2'; s.freeze_panes = None; print(s.freeze_panes, s.sheet_view.pane)
'''),
("Charts: the chapter's bar chart, the XML it becomes, and the Reference arguments", FIX + r'''
import openpyxl, zipfile, re
wb = openpyxl.Workbook(); sheet = wb.active
for i in range(1, 11): sheet['A' + str(i)] = i * i
ref_obj = openpyxl.chart.Reference(sheet, 1, 1, 1, 10)
series_obj = openpyxl.chart.Series(ref_obj, title='First series')
chart_obj = openpyxl.chart.BarChart()
chart_obj.title = 'My Chart'
chart_obj.append(series_obj)
sheet.add_chart(chart_obj, 'C5')
wb.save('sampleChart3.xlsx')
z = zipfile.ZipFile('sampleChart3.xlsx'); c = z.read('xl/charts/chart1.xml').decode()
print(ref_obj, [n for n in z.namelist() if 'chart' in n or 'drawing' in n])
print(re.findall(r'<barDir[^>]*>|<grouping[^>]*>|<f>[^<]*</f>|<v>[^<]*</v>', c))
print(re.findall(r'<(?:xdr:)?from>.*?</(?:xdr:)?from>', z.read('xl/drawings/drawing1.xml').decode()))
r = openpyxl.chart.Reference(sheet, 2, 1, 2, 10); print(r, '|', openpyxl.chart.Reference(sheet, min_col=1, min_row=3, max_col=1, max_row=5))
print(openpyxl.chart.Reference(sheet, 1, 1, 1))
'''),
("Charts: the four kinds the chapter names, and the scatter chart's x values", FIX + r'''
import openpyxl, zipfile, re
from openpyxl.chart import BarChart, LineChart, ScatterChart, PieChart, Reference, Series
wb = openpyxl.Workbook(); sheet = wb.active
for i in range(1, 11): sheet['A%d' % i] = i * i; sheet['B%d' % i] = i
y = Reference(sheet, 1, 1, 1, 10); x = Reference(sheet, 2, 1, 2, 10)
for n, cls in enumerate((BarChart, LineChart, ScatterChart, PieChart)):
    c = cls(); c.title = cls.__name__
    c.append(Series(y, xvalues=x, title='squares')) if cls is ScatterChart else c.append(Series(y, title='squares'))
    sheet.add_chart(c, 'D%d' % (1 + 16 * n))
wb.save('four.xlsx'); z = zipfile.ZipFile('four.xlsx')
print([re.findall(r'<(barChart|lineChart|scatterChart|pieChart)>', z.read('xl/charts/chart%d.xml' % k).decode())[0] for k in (1, 2, 3, 4)])
print(re.findall(r'<(xVal|yVal)>', z.read('xl/charts/chart3.xml').decode()), re.findall(r'<(cat|val)>', z.read('xl/charts/chart1.xml').decode()))
'''),
("Charts and pictures through a load and a save, and copy_worksheet", FIX + r'''
import openpyxl, zipfile
wb = openpyxl.Workbook(); sheet = wb.active
for i in range(1, 6): sheet['A' + str(i)] = i
c = openpyxl.chart.BarChart(); c.append(openpyxl.chart.Series(openpyxl.chart.Reference(sheet, 1, 1, 1, 5))); sheet.add_chart(c, 'C2'); wb.save('chart.xlsx')
again = openpyxl.load_workbook('chart.xlsx'); print(len(again.active._charts)); again.save('chart2.xlsx')
print([n for n in zipfile.ZipFile('chart2.xlsx').namelist() if n.startswith('xl/charts/')])
copy = again.copy_worksheet(again.active); print(copy.title, len(copy._charts), copy['A3'].value)
'''),
("Insert rows as Blank Row Inserter does: what moves and what does not", FIX + r'''
import openpyxl
wb = openpyxl.Workbook(); s = wb.active
for r in range(1, 7): s.cell(row=r, column=1, value='row%d' % r); s.cell(row=r, column=2, value=r * 10)
s['C6'] = '=SUM(B1:B5)'; s.merge_cells('A5:B5'); s.row_dimensions[4].height = 50; s['A4'].font = openpyxl.styles.Font(bold=True)
s.insert_rows(3, 2)
print([s.cell(row=r, column=1).value for r in range(1, 9)])
print(s['C8'].value, s['C6'].value, sorted(map(str, s.merged_cells.ranges)), s.row_dimensions[4].height, s.row_dimensions[6].height, s['A6'].font.b, s['A4'].font.b)
'''),
("Practice Programs: multiplicationTable.py and blankRowInserter.py run as scripts with command line arguments", FIX + r'''
import subprocess, sys, openpyxl
MULT = """import openpyxl, sys
from openpyxl.styles import Font
n = int(sys.argv[1])
wb = openpyxl.Workbook(); sheet = wb.active; bold = Font(bold=True)
for i in range(1, n + 1):
    sheet.cell(row=1, column=i + 1, value=i).font = bold
    sheet.cell(row=i + 1, column=1, value=i).font = bold
    for j in range(1, n + 1):
        sheet.cell(row=i + 1, column=j + 1, value=i * j)
wb.save('multiplicationTable.xlsx')
"""
INS = """import openpyxl, sys
n, m, filename = int(sys.argv[1]), int(sys.argv[2]), sys.argv[3]
wb = openpyxl.load_workbook(filename); sheet = wb.active
sheet.insert_rows(n, m)
wb.save('inserted_' + filename)
"""
open('multiplicationTable.py', 'w').write(MULT); open('blankRowInserter.py', 'w').write(INS)
print(subprocess.run([sys.executable, 'multiplicationTable.py', '6']).returncode)
s = openpyxl.load_workbook('multiplicationTable.xlsx').active
print(s.dimensions, s['A1'].value, [c.value for c in s[1]], [c.value for c in s['A']], s['D5'].value, s['G7'].value, s['B1'].font.b, s['A7'].font.b, s['B2'].font.b)
w = openpyxl.Workbook(); sh = w.active
for r, v in enumerate(['Potatoes', 'Okra', 'Fava beans', 'Watermelon', 'Garlic', 'Parsnips'], 1): sh.cell(row=r, column=1, value=v)
w.save('myProduce.xlsx')
print(subprocess.run([sys.executable, 'blankRowInserter.py', '3', '2', 'myProduce.xlsx']).returncode)
print([c.value for c in openpyxl.load_workbook('inserted_myProduce.xlsx').active['A']])
'''),
("Practice Questions 1 to 18 answered with code", FIX + r'''
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
wb = openpyxl.load_workbook('example3.xlsx'); sheet = wb['Sheet1']
print(1, type(wb).__name__); print(2, wb.sheetnames); print(3, wb['Sheet1']); print(4, wb.active)
print(5, sheet['C5'].value); sheet['C5'] = 'Hello'; print(6, sheet['C5'].value)
print(7, sheet['C5'].row, sheet['C5'].column); print(8, sheet.max_column, sheet.max_row, type(sheet.max_column).__name__)
print(9, column_index_from_string('M')); print(10, get_column_letter(14), 'rows have no letters:', [n for n in dir(openpyxl.utils) if 'row' in n.lower()])
print(11, len(sheet['A1':'F1']), len(sheet['A1':'F1'][0]), sheet['A1':'F1'][0][5].coordinate)
wb.save('example3.xlsx'); print(12, 'example3.xlsx' in __import__('os').listdir('.'))
sheet['A9'] = '=SUM(C1:C7)'; print(13, sheet['A9'].value)
sheet.row_dimensions[5].height = 100; print(15, sheet.row_dimensions[5].height)
sheet.column_dimensions['C'].hidden = True; print(16, sheet.column_dimensions['C'].hidden)
sheet.freeze_panes = 'A2'; print(17, sheet.freeze_panes)
'''),
("Printed listings as typeset: the call-out markers sit inside the code", r'''
import re
step2 = """for row in range(2, sheet.max_row + 1):
    state  = sheet['B' + str(row)].value
    county = sheet['C' + str(row)].value
    pop    = sheet['D' + str(row)].value
  ❶ county_data.setdefault(state, {})
    county_data[state].setdefault(county, {'tracts': 0, 'pop': 0})
"""
step1 = """❶ import openpyxl, pprint
wb = openpyxl.load_workbook('censuspopdata.xlsx')
"""
for name, text in (('Project 9 Step 1', step1), ('Project 9 Step 2', step2)):
    for how, f in (('as printed', lambda t: t), ('marker removed', lambda t: re.sub('[❶-❿]', '', t)), ('marker and one space removed', lambda t: re.sub('[❶-❿] ', '', t)), ('marker replaced by one space', lambda t: re.sub('[❶-❿]', ' ', t))):
        try: compile(f(text), 'listing', 'exec'); print(name, how, 'compiles')
        except (SyntaxError, IndentationError) as e: print(name, how, type(e).__name__, e.msg)
print(repr(step2.splitlines()[4][:6]))
'''),
("Reading more than the book does: read-only mode, iter_rows, and a touched cell that makes the sheet look longer", FIX + r'''
import openpyxl
ro = openpyxl.load_workbook('censuspopdata.xlsx', read_only=True); s = ro['Population by Census Tract']
print(type(s).__name__, s.max_row, [c for c in next(s.iter_rows(min_row=2, max_row=2, values_only=True))])
try: list(s.columns)
except AttributeError as e: print('AttributeError', e)
try: s['A1'].value = 1
except AttributeError as e: print('AttributeError', e)
try: ro.save('x.xlsx')
except TypeError as e: print('TypeError', e)
ro.close()
wb = openpyxl.load_workbook('censuspopdata.xlsx'); s = wb.active
print(s.max_row, [sum(1 for _ in s.iter_rows(min_row=2, values_only=True))])
s['D40']; print(s.max_row)
try:
    for row in range(2, s.max_row + 1): int(s['D' + str(row)].value)
except TypeError as e: print('TypeError', e)
'''),
("Files that are not what their name says: .xls, .txt, .xlsm and a wrong extension", FIX + r'''
import openpyxl, shutil
wb = openpyxl.Workbook(); wb.save('named.xls')
for name in ('named.xls',):
    try: openpyxl.load_workbook(name)
    except Exception as e: print(type(e).__name__, str(e)[:70])
for ext in ('txt', 'xlsm', 'xlsx'):
    shutil.copy('named.xls', 'named.' + ext)
    try: print(ext, openpyxl.load_workbook('named.' + ext).sheetnames)
    except Exception as e: print(ext, type(e).__name__, str(e)[:60])
open('notzip.xlsx', 'w').write('hello')
try: openpyxl.load_workbook('notzip.xlsx')
except Exception as e: print(type(e).__name__, e)
try: openpyxl.load_workbook('missing.xlsx')
except Exception as e: print(type(e).__name__)
'''),
("The active sheet: a workbook opens on one sheet, and setting wb.active leaves the first sheet selected", FIX + r'''
import openpyxl, zipfile, re
x = openpyxl.load_workbook('example3.xlsx'); x.active.sheet_view.tabSelected = True; x.save('example3_selected.xlsx')
wb = openpyxl.load_workbook('example3_selected.xlsx'); print(wb.active.title)
wb.active = 2; print(wb.active.title); wb.save('act.xlsx')
z = zipfile.ZipFile('act.xlsx')
print(re.findall(r'activeTab="\d"', z.read('xl/workbook.xml').decode()), [bool(re.search(r'tabSelected="1"', z.read('xl/worksheets/sheet%d.xml' % n).decode())) for n in (1, 2, 3)])
'''),
("Untrusted files: an XML entity bomb of 4.9 KB read by openpyxl alone (its page warns it does not guard against it)", FIX + r'''
import io, sys, zipfile, openpyxl, openpyxl.xml, pyexpat
def bomb(depth):
    ents = ''.join('<!ENTITY e%d "%s">' % (i, ('&e%d;' % (i - 1)) * 10) for i in range(1, depth + 1))
    b = io.BytesIO(); src = openpyxl.Workbook(); src.active['A1'] = 'seed'; src.save(b); zin = zipfile.ZipFile(io.BytesIO(b.getvalue())); out = io.BytesIO(); zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
    for n in zin.namelist():
        data = zin.read(n)
        if n == 'xl/worksheets/sheet1.xml':
            t = data.decode(); i = t.index('<worksheet'); t = '<?xml version="1.0"?><!DOCTYPE worksheet [<!ENTITY e0 "aaaaaaaaaa">%s]>' % ents + t[i:]; data = t.replace('>seed<', '>&e%d;<' % depth).encode()
        zout.writestr(n, data)
    zout.close(); return out.getvalue()
print('expat', pyexpat.EXPAT_VERSION, '| defusedxml in use:', openpyxl.xml.DEFUSEDXML)
for depth in (4, 6):
    data = bomb(depth)
    try: print(depth, 'file under 6000 bytes:', len(data) < 6000, '->', len(openpyxl.load_workbook(io.BytesIO(data)).active['A1'].value), 'characters in A1')
    except Exception as e: print(depth, 'file under 6000 bytes:', len(data) < 6000, '->', type(e).__name__, str(e)[:60])
'''),
("Pictures without Pillow: add_image on an environment that has openpyxl alone", FIX + r'''
import openpyxl, zipfile, importlib.util
print('Pillow importable:', importlib.util.find_spec('PIL') is not None)
if importlib.util.find_spec('PIL'):
    from PIL import Image
    Image.new('RGB', (40, 20), (200, 30, 30)).save('red.png')
try:
    wb = openpyxl.Workbook(); wb.active.add_image(openpyxl.drawing.image.Image('red.png'), 'B2'); wb.save('img.xlsx')
    again = openpyxl.load_workbook('img.xlsx'); print('images kept after load:', len(again.active._images)); again.save('img2.xlsx'); print([n for n in zipfile.ZipFile('img2.xlsx').namelist() if 'media' in n])
except Exception as e: print(type(e).__name__, e)
'''),
]
EXTRA = [
("Untrusted files: the same XML entity bomb read with defusedxml on the path", "defusedxml", FIX + r'''
import io, sys, zipfile, openpyxl, openpyxl.xml, pyexpat
def bomb(depth):
    ents = ''.join('<!ENTITY e%d "%s">' % (i, ('&e%d;' % (i - 1)) * 10) for i in range(1, depth + 1))
    b = io.BytesIO(); src = openpyxl.Workbook(); src.active['A1'] = 'seed'; src.save(b); zin = zipfile.ZipFile(io.BytesIO(b.getvalue())); out = io.BytesIO(); zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
    for n in zin.namelist():
        data = zin.read(n)
        if n == 'xl/worksheets/sheet1.xml':
            t = data.decode(); i = t.index('<worksheet'); t = '<?xml version="1.0"?><!DOCTYPE worksheet [<!ENTITY e0 "aaaaaaaaaa">%s]>' % ents + t[i:]; data = t.replace('>seed<', '>&e%d;<' % depth).encode()
        zout.writestr(n, data)
    zout.close(); return out.getvalue()
print('expat', pyexpat.EXPAT_VERSION, '| defusedxml in use:', openpyxl.xml.DEFUSEDXML)
for depth in (4, 6):
    data = bomb(depth)
    try: print(depth, 'file under 6000 bytes:', len(data) < 6000, '->', len(openpyxl.load_workbook(io.BytesIO(data)).active['A1'].value), 'characters in A1')
    except Exception as e: print(depth, 'file under 6000 bytes:', len(data) < 6000, '->', type(e).__name__, str(e)[:60])
'''),
("Pictures need Pillow: add_image without it, and with it the picture survives a load and a save", "pillow", FIX + r'''
import openpyxl, zipfile, importlib.util
print('Pillow importable:', importlib.util.find_spec('PIL') is not None)
if importlib.util.find_spec('PIL'):
    from PIL import Image
    Image.new('RGB', (40, 20), (200, 30, 30)).save('red.png')
try:
    wb = openpyxl.Workbook(); wb.active.add_image(openpyxl.drawing.image.Image('red.png'), 'B2'); wb.save('img.xlsx')
    again = openpyxl.load_workbook('img.xlsx'); print('images kept after load:', len(again.active._images)); again.save('img2.xlsx'); print([n for n in zipfile.ZipFile('img2.xlsx').namelist() if 'media' in n])
except Exception as e: print(type(e).__name__, e)
'''),
("Other libraries: xlsxwriter writes but cannot open, xlrd 2 reads only .xls, and the cached value xlsxwriter stores for a formula", "other", FIX + r'''
import io, openpyxl, xlrd, xlsxwriter
print('xlsxwriter', xlsxwriter.__version__, '| xlrd', xlrd.__version__)
try: xlrd.open_workbook('example3.xlsx')
except Exception as e: print(type(e).__name__, e)
b = io.BytesIO(); wb = xlsxwriter.Workbook(b, {'in_memory': True}); ws = wb.add_worksheet(); ws.write('A1', 200); ws.write('A2', 300); ws.write_formula('A3', '=SUM(A1:A2)'); wb.close()
f = openpyxl.load_workbook(io.BytesIO(b.getvalue())); v = openpyxl.load_workbook(io.BytesIO(b.getvalue()), data_only=True)
print(f.active['A3'].value, v.active['A3'].value)
b = io.BytesIO(); wb = xlsxwriter.Workbook(b, {'in_memory': True}); ws = wb.add_worksheet(); ws.write_formula('A1', '=1+1', None, 2); wb.close()
print(openpyxl.load_workbook(io.BytesIO(b.getvalue()), data_only=True).active['A1'].value, [hasattr(xlsxwriter, n) for n in ('load_workbook', 'open_workbook')])
'''),
]
def run(code, need=None):
    env = dict(os.environ, PYTHONDONTWRITEBYTECODE="1")
    env.pop("PYTHONPATH", None)
    if need: env["PYTHONPATH"] = SITES[need]
    r = subprocess.run([PY, "-W", "ignore", "-c", code], capture_output=True, text=True, env=env, timeout=600)
    return (r.stdout if r.returncode == 0 else r.stdout + "FAILED: " + r.stderr.strip().splitlines()[-1]).rstrip("\n")
out = [{"what": w, "code": c, "output": run(c)} for w, c in PROGS] + [{"what": w, "code": c, "output": run(c, n), "needs": n} for w, n, c in EXTRA]
if MODE == "write":
    json.dump({"_version": "1.0.0", "_note": "Recorded by sen0414_ch14_thirdparty_v1_0_0.py write; every output was printed by the program beside it under the interpreter named in the field python, with workbooks the programs build in a temporary folder. Programs with a field needs ran with that one optional package on the path, the others with openpyxl alone.", "python": subprocess.run([PY, "--version"], capture_output=True, text=True).stdout.strip(), "items": out}, open(REC, "w"), indent=1)
    print("recorded", len(out), "programs to", REC)
else:
    old = json.load(open(REC))["items"]; bad = [a["what"] for a, b in zip(old, out) if a["output"] != b["output"] or a["code"] != b["code"]]
    print("%d programs; %d differ" % (len(out), len(bad)))
    for b in bad: print("  DIFFERS", b)
    sys.exit(1 if bad else 0)
